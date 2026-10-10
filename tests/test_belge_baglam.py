import io
import os
import zipfile

from openpyxl import Workbook

from koordmes.services import aktarim, baglanti, okuyucu, sgk_karsilastirma, sigorta
from koordmes.services.okuyucu import Tablo

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


def _xlsx(yol, satirlar):
    wb = Workbook()
    ws = wb.active
    for r in satirlar:
        ws.append(r)
    wb.save(yol)
    return yol


# ------------------------------------------------------------------ bağlam çözümleme

def test_baglam_coz_kaliplari():
    assert aktarim.baglam_coz(["Sınıfı : 11-A   Alanı : Bilişim Teknolojileri Alanı"]) == {
        "sinif": "11-A", "alan": "Bilişim Teknolojileri"}
    assert aktarim.baglam_coz(["Sınıf: 11  Şube: B"]) == {"sinif": "11-B"}
    assert aktarim.baglam_coz(["Sınıfı", "10/C"]) == {"sinif": "10-C"}
    assert aktarim.baglam_coz(["12. Sınıf / A Şubesi Öğrenci Listesi"]) == {"sinif": "12-A"}
    assert aktarim.baglam_coz(["11-B"]) == {"sinif": "11-B"}
    assert aktarim.baglam_coz(["EKİM 2026 DÖNEMİ DEVAMSIZLIK LİSTESİ"]) == {"ay": "2026-10"}
    assert aktarim.baglam_coz(["Dönem: 2026/10", "Belge Türü: 07 - Çırak"]) == {"ay": "2026-10", "belge_turu": "7"}
    assert aktarim.baglam_coz(["Ders Yılı / Dönem: 2026-2027 / 1"]) == {}
    assert aktarim.baglam_coz(["Ahmet Yılmaz", "Matematik"]) == {}
    assert aktarim.baglam_coz(["9. SINIF ÖĞRENCİ LİSTESİ"]) == {}


def test_birlesik_iki_satirli_baslik():
    rows = [["Sıra", "Öğrenci", None, "İşletme", None, "İşe Başlama Tarihi"],
            [None, "Adı Soyadı", "T.C. Kimlik No", "Adı", "Vergi No", None],
            [1, "Ali Can", "10000000078", "Usta Oto", "1234567890", "15.09.2026"]]
    t = Tablo("x.xlsx", "Sayfa1", rows)
    tah = aktarim.tahmin_et(t)[0]
    assert tah.anahtar == "yerlestirmeler" and tah.bas == 1
    assert tah.eslesme["isletme"] == 3 and tah.eslesme["ad_soyad"] == 1 and tah.eslesme["tc"] == 2


# ------------------------------------------------------------------ dosyalarda bağlam

def test_excel_ara_baslikli_coklu_sinif_ve_imza(db, tmp_path):
    yol = _xlsx(tmp_path / "liste.xlsx", [
        ["T.C. ÖRNEK MESLEKİ VE TEKNİK ANADOLU LİSESİ"],
        ["Okul No", "Adı Soyadı", "Cinsiyeti"],
        ["11. Sınıf / A Şubesi"],
        ["101", "AYŞE KAYA", "K"],
        ["102", "Mehmet Demir", "E"],
        ["11-B SINIFI"],
        ["201", "Zeynep Ak", "K"],
        [None, "Okul Müdürü", None],
        ["Genel Toplam: 3 öğrenci"],
    ])
    sonuclar, hatalar = aktarim.otomatik_aktar(db, [yol])
    assert not hatalar and sonuclar[0].anahtar == "ogrenciler" and sonuclar[0].eklenen == 3
    rows = {r["ad"]: r["sinif"] for r in db.q("SELECT o.ad, s.ad AS sinif FROM ogrenciler o "
                                                "JOIN siniflar s ON s.id = o.sinif_id")}
    assert rows == {"Ayşe": "11-A", "Mehmet": "11-A", "Zeynep": "11-B"}


def test_html_tablo_disi_baslik_ve_bilgi_tablosu(db, tmp_path):
    html = ("<html><body><h3>Sınıfı : 10-B</h3><table><tr><td>Alanı</td><td>Elektrik-Elektronik</td></tr></table>"
            "<table><tr><th>Öğrenci No</th><th>Adı</th><th>Soyadı</th></tr>"
            "<tr><td>301</td><td>Can</td><td>Er</td></tr></table></body></html>")
    yol = tmp_path / "OOG01001R020.xls"
    yol.write_text(html, encoding="utf-8")
    t = okuyucu.oku(yol)[0]
    assert "10-B" in aktarim.baglam_ozeti(t)
    sonuclar, _ = aktarim.otomatik_aktar(db, [yol])
    assert sonuclar[0].eklenen == 1
    r = db.q1("SELECT s.ad FROM ogrenciler o JOIN siniflar s ON s.id = o.sinif_id")
    assert r["ad"] == "10-B"


def test_docx_paragraf_baglami_ve_zip(db, tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("Sınıfı: 9-C")
    t = d.add_table(rows=2, cols=2)
    for j, v in enumerate(["Öğrenci No", "Adı Soyadı"]):
        t.cell(0, j).text = v
    t.cell(1, 0).text, t.cell(1, 1).text = "401", "Ela Su"
    buf = io.BytesIO()
    d.save(buf)
    arsiv = tmp_path / "eokul.zip"
    with zipfile.ZipFile(arsiv, "w") as z:
        z.writestr("liste.docx", buf.getvalue())
        z.writestr("benioku.bin", b"\x00\x01")
    tablolar = okuyucu.oku(arsiv)
    assert tablolar[0].kaynak.startswith("eokul.zip")
    sonuclar, _ = aktarim.otomatik_aktar(db, [arsiv])
    assert sonuclar[0].eklenen == 1
    assert db.q1("SELECT s.ad FROM ogrenciler o JOIN siniflar s ON s.id = o.sinif_id")["ad"] == "9-C"


def test_pdf_sayfa_basligi_sinif_bilgisi(db, tmp_path):
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from koordmes.ui import ortak
    tablo = lambda satirlar: ("<table border='1'><tr><th>Okul No</th><th>Adı Soyadı</th></tr>"
                              + "".join(f"<tr><td>{a}</td><td>{b}</td></tr>" for a, b in satirlar) + "</table>")
    html = ("<p>Sınıfı : 12-A</p>" + tablo([("501", "Ali Veli"), ("502", "Can Ece")])
            + "<p>Sınıfı : 12-B</p>" + tablo([("601", "Deniz Kar")]))
    yol = tmp_path / "liste.pdf"
    ortak.html_pdf(html, yol)
    sonuclar, _ = aktarim.otomatik_aktar(db, [yol])
    assert sum(s.eklenen for s in sonuclar) == 3
    rows = {r["okul_no"]: r["sinif"] for r in db.q("SELECT o.okul_no, s.ad AS sinif FROM ogrenciler o "
                                                     "JOIN siniflar s ON s.id = o.sinif_id")}
    assert rows == {"501": "12-A", "502": "12-A", "601": "12-B"}


def test_toplu_giriste_baglam_kapali(db):
    from koordmes.services import toplu_giris
    tur = toplu_giris.GIRIS_SOZLUK["ogrenciler"]
    satirlar = [[""] * len(tur.sutunlar) for _ in range(2)]
    j = [s.anahtar for s in tur.sutunlar].index("ad_soyad")
    satirlar[0][j], satirlar[1][j] = "Okul Müdürü", "11-A"
    sonuclar, _ = toplu_giris.kaydet(db, tur, satirlar)
    assert sonuclar[0].eklenen == 2  # elle girilen satırlar alt bilgi/ara başlık sayılmaz


def test_izlenecek_dosyalar(tmp_path):
    assert baglanti.izlenecek_mi(tmp_path / "liste.xlsx")
    assert not baglanti.izlenecek_mi(tmp_path / "liste.pdf")
    assert baglanti.izlenecek_mi(tmp_path / "liste.pdf", belgeler=True)
    assert not baglanti.izlenecek_mi(tmp_path / "liste.pdf.crdownload", belgeler=True)


# ------------------------------------------------------------------ SGK karşılaştırma

def test_sgk_aphb_karsilastirma(ornek_db, tmp_path):
    db = ornek_db
    okul = sigorta.aylik_liste(db, 2026, 10)
    assert len(okul) >= 5
    eksik, gun_farkli, belge_farkli, ad_farkli = okul[0], okul[1], okul[2], okul[3]
    satirlar = [["Dönem: 2026/10"], ["T.C. Kimlik No", "Adı", "Soyadı", "Belge Türü", "Prim Ödeme Günü",
                                     "Prime Esas Kazanç", "Eksik Gün Sayısı"]]
    for s in okul[1:]:
        gun = s.gun - 5 if s is gun_farkli else s.gun
        belge = "42" if s is belge_farkli and s.belge_turu == "7" else "7" if s is belge_farkli else s.belge_turu
        soyad = s.soyad + "oğlu" if s is ad_farkli else s.soyad
        satirlar.append([s.tc, s.ad.upper(), soyad, belge, gun, s.pek, s.eksik_gun])
    satirlar.append(["10000000078", "Yabancı", "Kişi", "7", 30, 16515, 0])
    yol = _xlsx(tmp_path / "aphb.xlsx", satirlar)
    liste = sgk_karsilastirma.listeyi_oku([yol])
    assert liste.aphb and liste.donem == (2026, 10) and len(liste.kayitlar) == len(okul)
    sonuc = sgk_karsilastirma.karsilastir(db, liste)
    turler = {(f.tur, f.tc) for f in sonuc.farklar}
    assert ("SGK listesinde yok", eksik.tc) in turler
    assert ("Prim günü farklı", gun_farkli.tc) in turler
    assert ("Belge türü farklı", belge_farkli.tc) in turler
    assert ("Ad/soyad farklı", ad_farkli.tc) in turler
    assert ("Okul kayıtlarında yok", "10000000078") in turler
    assert sonuc.tahmini_ceza > 0 and "eşleşti" in sonuc.ozet()

    ad = next(f for f in sonuc.farklar if f.tur == "Ad/soyad farklı")
    sgk_karsilastirma.duzelt(db, ad)
    assert db.q1("SELECT soyad FROM ogrenciler WHERE id = ?", (ad.ogrenci_id,))["soyad"].endswith("oğlu")
    assert sgk_karsilastirma.excel_yaz(sonuc, tmp_path / "rapor.xlsx").exists()


def test_sgk_tc_farkli_ad_ile_eslesir(ornek_db, tmp_path):
    db = ornek_db
    s = sigorta.aylik_liste(db, 2026, 10)[0]
    db.update("ogrenciler", s.ogrenci_id, {"tc": "123"})
    yol = _xlsx(tmp_path / "aphb.xlsx", [["Sigortalının T.C. Kimlik No", "Adı Soyadı", "Belge Türü", "Gün"],
                                         [s.tc, f"{s.ad} {s.soyad}", s.belge_turu, s.gun]])
    sonuc = sgk_karsilastirma.karsilastir(db, sgk_karsilastirma.listeyi_oku([yol]), 2026, 10)
    f = next(f for f in sonuc.farklar if f.tur == "T.C. farklı")
    assert f.eylem == "tc_duzelt" and "geçersiz" in f.oneri
    sgk_karsilastirma.duzelt(db, f)
    assert db.q1("SELECT tc FROM ogrenciler WHERE id = ?", (s.ogrenci_id,))["tc"] == s.tc


def test_sgk_listesi_indirmede_taninir(db, tmp_path):
    yol = _xlsx(tmp_path / "hizmet.xlsx", [["T.C. Kimlik No", "Ad Soyad", "Prim Gün", "PEK"],
                                           ["10000000078", "Ali Veli", 30, 16515]])
    s = baglanti.indirileni_isle(db, yol, True)
    assert s.sgk_listesi and "SGK" in s.ozet
    assert sgk_karsilastirma.son_liste()[0] == yol
