import io
import os
import zipfile
from datetime import date

import pytest

from koordmes.services import aktarim, okuyucu
from koordmes.services.okuyucu import Tablo

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

BASLIK = ["Öğrenci No", "Adı Soyadı", "Sınıfı", "İşletme Adı", "İşe Başlama Tarihi"]
VERI = [["501", "ÇAĞLA ÖZTÜRK", "11. Sınıf / A Şubesi", "Işık Bilişim Ltd. Şti.", "15.09.2026"],
        ["502", "Şükrü Güneş", "11-A", "Işık Bilişim Ltd. Şti.", "15.09.2026"]]


def _en_iyi(yol):
    tablolar = okuyucu.oku(yol)
    assert tablolar
    return tablolar, aktarim.tahmin_et(tablolar[0])


# ------------------------------------------------------------------ okuyucular

def test_docx_tablosu(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("2026-2027 İşletmede Mesleki Eğitim Öğrenci Listesi")
    t = d.add_table(rows=1 + len(VERI), cols=len(BASLIK))
    for i, row in enumerate([BASLIK] + VERI):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    yol = tmp_path / "liste.docx"
    d.save(yol)
    tablolar, tahmin = _en_iyi(yol)
    assert tablolar[0].satirlar[0] == BASLIK and len(tablolar[0].satirlar) == 3
    assert tahmin[0].anahtar == "yerlestirmeler"


def test_eokul_html_xls_ve_birlesik_baslik(tmp_path):
    html = ("<html><body><table><tr><td colspan='5'>ATATÜRK MTAL ÖĞRENCİ LİSTESİ</td></tr>"
            "<tr>" + "".join(f"<th>{b}</th>" for b in BASLIK[:3]) + "</tr>"
            + "".join("<tr>" + "".join(f"<td>{v}</td>" for v in r[:3]) + "</tr>" for r in VERI)
            + "</table></body></html>")
    yol = tmp_path / "OOG01001R020.xls"
    yol.write_text(html, encoding="utf-8")
    tablolar, tahmin = _en_iyi(yol)
    assert tahmin[0].anahtar == "ogrenciler" and tahmin[0].bas == 1


def test_csv_cp1254(tmp_path):
    yol = tmp_path / "ogrenci.csv"
    yol.write_bytes("Okul No;Adı;Soyadı;Sınıfı\n601;Gülşen;Işıkçı;10/B\n".encode("cp1254"))
    t = okuyucu.oku(yol)[0]
    assert t.satirlar[1][1:3] == ["Gülşen", "Işıkçı"]
    assert aktarim.tahmin_et(t)[0].anahtar == "ogrenciler"


def test_ods(tmp_path):
    ns = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
          'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
          'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"')
    hucre = lambda v: f"<table:table-cell><text:p>{v}</text:p></table:table-cell>"
    satirlar = "".join("<table:table-row>" + "".join(hucre(v) for v in r) + "</table:table-row>"
                       for r in [["Tatil Adı", "Başlangıç", "Bitiş"], ["Yarıyıl", "19.01.2027", "30.01.2027"]])
    xml = (f'<?xml version="1.0"?><office:document-content {ns}><office:body><office:spreadsheet>'
           f'<table:table table:name="Tatiller">{satirlar}</table:table></office:spreadsheet></office:body>'
           f'</office:document-content>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        z.writestr("content.xml", xml)
    yol = tmp_path / "takvim.ods"
    yol.write_bytes(buf.getvalue())
    tablolar, tahmin = _en_iyi(yol)
    assert tablolar[0].ad == "Tatiller" and tahmin[0].anahtar == "tatiller"


def test_pdf_tablosu(tmp_path):
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from koordmes.ui import ortak
    yol = ortak.html_pdf(ortak.html_tablo("Öğrenci Listesi", BASLIK, VERI), tmp_path / "liste.pdf")
    t = okuyucu.oku(yol)[0]
    assert t.satirlar[1][:2] == ["501", "ÇAĞLA ÖZTÜRK"] and "Bilişim" in t.satirlar[1][3]
    assert aktarim.tahmin_et(t)[0].anahtar == "yerlestirmeler"


def test_etiket_deger_formu(tmp_path):
    yol = tmp_path / "isletme.txt"
    yol.write_text("İşletme Adı: Usta Mobilya\nVergi No: 1234567890\nYetkili: Ali Kaya\nTelefon: 0312 111 22 33\n",
                   encoding="utf-8")
    t = okuyucu.oku(yol)[0]
    assert t.form and len(t.satirlar) == 2
    assert aktarim.tahmin_et(t)[0].anahtar == "isletmeler"


def test_desteklenmeyen_dosya(tmp_path):
    yol = tmp_path / "x.bin"
    yol.write_bytes(b"\x00\x01\x02" * 50)
    with pytest.raises(okuyucu.OkumaHatasi):
        okuyucu.oku(yol)


# ------------------------------------------------------------------ dönüştürücüler

def test_donusturuculer():
    assert aktarim.sinif_adi("11. Sınıf / A Şubesi") == "11-A"
    assert aktarim.sinif_adi("10/b") == "10-B"
    assert aktarim.sinif_adi("9 T A") == "9 T A" and aktarim.sinif_adi("10EE A") == "10EE A"
    assert aktarim.sinif_adi("9-A Sınıfı Haftalık Ders Programı") == "9-A"
    assert aktarim.tr_ad("ÇAĞLA İLKİM") == "Çağla İlkim"
    assert aktarim.sayi("1.234,50") == 1234.5
    assert aktarim.tarih("1 Eylül 2026") == date(2026, 9, 1)
    assert aktarim.tarih(46266) == date(2026, 9, 1)
    assert aktarim.isletme_anahtari("Işık Bilişim Ltd. Şti.") == aktarim.isletme_anahtari("IŞIK BİLİŞİM LİMİTED ŞİRKETİ")


# ------------------------------------------------------------------ aktarım motoru

def test_deneme_kaydetmez_gercek_aktarim_iliskili_kayitlari_olusturur(db):
    t = Tablo("liste.xlsx", "Sayfa1", [BASLIK] + VERI)
    s = aktarim.aktar(db, t, "yerlestirmeler", deneme=True)
    assert s.deneme and s.eklenen == 2
    assert db.q1("SELECT COUNT(*) FROM yerlestirmeler")[0] == 0
    assert db.q1("SELECT COUNT(*) FROM ogrenciler")[0] == 0
    s = aktarim.aktar(db, t, "yerlestirmeler")
    assert s.eklenen == 2 and len(s.yeni["yerlestirmeler"]) == 2
    assert db.q1("SELECT COUNT(*) FROM isletmeler")[0] == 1
    assert db.q1("SELECT COUNT(*) FROM siniflar WHERE ad = '11-A'")[0] == 1
    assert db.q1("SELECT ad FROM ogrenciler WHERE okul_no = '501'")[0] == "Çağla"


def test_isletme_degisince_eski_yerlestirme_kapanir(db):
    aktarim.aktar(db, Tablo("a", "a", [BASLIK] + VERI[:1]), "yerlestirmeler")
    yeni = [BASLIK, ["501", "Çağla Öztürk", "11-A", "Güneş Elektrik", "01.02.2027"]]
    s = aktarim.aktar(db, Tablo("b", "b", yeni), "yerlestirmeler")
    assert s.iliskili["kapatılan eski yerleştirme"] == 1
    rows = db.q("SELECT durum, bitis FROM yerlestirmeler ORDER BY baslangic")
    assert [r["durum"] for r in rows] == ["Sona Erdi", "Aktif"] and rows[0]["bitis"] == "2027-01-31"


def test_eokul_toplam_devamsizlik_oncekinin_yerine_gecer(db):
    aktarim.aktar(db, Tablo("a", "a", [BASLIK] + VERI), "yerlestirmeler")
    bas = ["Öğrenci No", "Adı Soyadı", "Özürsüz", "Özürlü"]
    aktarim.aktar(db, Tablo("d", "d", [bas, ["501", "Çağla Öztürk", 3, 1]]), "devamsizlik")
    s = aktarim.aktar(db, Tablo("d", "d", [bas, ["501", "Çağla Öztürk", 5, 1]]), "devamsizlik")
    assert s.guncellenen == 1
    r = db.q("SELECT tur, gun FROM devamsizlik ORDER BY tur")
    assert [(x["tur"], x["gun"]) for x in r] == [("Özürlü", 1), ("Özürsüz", 5)]


def test_dekont_eksik_odeme_ve_sgk_isaretleme(db):
    aktarim.aktar(db, Tablo("a", "a", [BASLIK] + VERI), "yerlestirmeler")
    dek = [["Okul No", "Adı Soyadı", "Dönem", "Ödenen Tutar", "Dekont Tarihi"],
           ["501", "Çağla Öztürk", "Ekim 2026", "100,00", "05.11.2026"]]
    s = aktarim.aktar(db, Tablo("dekont", "d", dek), "dekontlar")
    assert s.eklenen == 1 and "EKSİK ÖDEME" in s.satirlar[0].mesaj
    r = db.q1("SELECT * FROM aylik_odemeler WHERE donem = '2026-10'")
    assert r["dekont_teslim"] == 1 and r["odenen_ucret"] == 100
    ogr = db.q1("SELECT tc FROM ogrenciler WHERE okul_no = '502'")
    db.execute("UPDATE ogrenciler SET tc = '10000000146' WHERE okul_no = '502'")
    sgk = [["T.C. Kimlik No", "Ad Soyad", "İşe Giriş Tarihi", "Bildirim Tarihi"],
           ["10000000146", "Şükrü Güneş", "15.09.2026", "14.09.2026"]]
    s = aktarim.aktar(db, Tablo("e-Bildirge", "s", sgk), "sgk_bildirimleri")
    assert s.guncellenen == 1, s.hatalar
    assert db.q1("""SELECT y.sigorta_giris_bildirildi FROM yerlestirmeler y JOIN ogrenciler o ON o.id = y.ogrenci_id
                    WHERE o.okul_no = '502'""")[0] == 1
    assert ogr is not None


def test_zorunlu_sutun_yoksa_hata(db):
    with pytest.raises(ValueError):
        aktarim.aktar(db, Tablo("x", "x", [["Renk", "Boy"], ["mavi", 3]]), "ogrenciler")


def test_otomatik_aktar_bagimlilik_sirasi(db, tmp_path):
    from koordmes.services.excel_io import ExcelYazici
    x = ExcelYazici()
    x.sayfa("Yerleştirmeler", BASLIK, VERI)
    x.sayfa("Öğretmenler", ["Adı Soyadı", "Branşı", "Kadro Türü"], [["Mehmet Usta", "Bilişim", "Kadrolu"]])
    yol = x.kaydet(tmp_path / "toplu.xlsx")
    sonuclar, hatalar = aktarim.otomatik_aktar(db, [yol])
    assert not hatalar
    assert [s.anahtar for s in sonuclar] == ["ogretmenler", "yerlestirmeler"]


def test_sablon_tum_turler(tmp_path):
    import openpyxl
    yol = aktarim.sablon_olustur(None, tmp_path / "s.xlsx")
    wb = openpyxl.load_workbook(yol)
    assert len(wb.sheetnames) == len(aktarim.TANIMLAR)
    ad_anahtar = {t.etiket[:31].replace("/", "-"): k for k, t in aktarim.TANIMLAR.items()}
    for ws in wb.worksheets:
        basliklar = [c.value for c in ws[1] if c.value]
        tah = aktarim.tahmin_et(Tablo("şablon.xlsx", ws.title, [basliklar, ["x"] * len(basliklar)]))
        assert tah and tah[0].anahtar == ad_anahtar[ws.title], (ws.title, [x.anahtar for x in tah[:3]])
