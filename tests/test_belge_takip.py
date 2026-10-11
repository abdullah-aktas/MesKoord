import struct
import zlib
from datetime import date, timedelta

import pytest
from docx import Document

from koordmes.services import barkod as bk, belge_takip as bt


def test_code128_desenleri_ve_cozme():
    assert len(bk.DESENLER) == 107 and len(set(bk.DESENLER)) == 107
    assert all(sum(map(int, d)) == 11 for d in bk.DESENLER[:106]) and sum(map(int, bk.DESENLER[106])) == 13
    kod = "71202610050000300042"
    assert bk.coz(bk.moduller(kod)) == kod
    # bilinen örnek: "00" -> başlangıç C, 0, sağlama (105 % 103 = 2), bitiş
    assert bk.degerler("00") == [105, 0, 2, 106]
    with pytest.raises(ValueError):
        bk.degerler("123")
    veri = bk.png(kod, 2, 10)
    assert veri.startswith(b"\x89PNG\r\n\x1a\n")
    w, h = struct.unpack(">II", veri[16:24])
    assert (w, h) == (len(bk.moduller(kod)) * 2, 10)
    idat_boy = struct.unpack(">I", veri[33:37])[0]
    ham = zlib.decompress(veri[41:41 + idat_boy])
    assert len(ham) == h * (w + 1)


def test_barkod_metni_ve_cozumu():
    k = bt.barkod_metni("G", date(2026, 10, 5), 3, 42)
    assert k == "71202610050000300042" and len(k) == 20
    b = bt.barkod_coz(" 7120261005-00003-00042\n")
    assert (b.tur, b.tarih, b.ogretmen_id, b.isletme_id) == ("G", date(2026, 10, 5), 3, 42)
    assert bt.barkod_coz(bt.barkod_metni("A", date(2026, 10, 1), 7, None)).isletme_id is None
    assert bt.barkod_coz("123") is None and bt.barkod_coz("71202613050000300042") is None


def _kur(db):
    t1 = db.insert("ogretmenler", {"ad": "Ali", "soyad": "Can"})
    t2 = db.insert("ogretmenler", {"ad": "Ece", "soyad": "Su"})
    i1 = db.insert("isletmeler", {"ad": "Usta Oto", "yetkili": "Hasan Usta"})
    i2 = db.insert("isletmeler", {"ad": "Işık Bilişim"})
    db.insert("koordinator_gorevleri", {"ogretmen_id": t1, "isletme_id": i1, "gun": 1, "haftalik_saat": 4})
    db.insert("koordinator_gorevleri", {"ogretmen_id": t1, "isletme_id": i2, "gun": 3, "haftalik_saat": 2})
    db.insert("koordinator_gorevleri", {"ogretmen_id": t2, "isletme_id": i2, "gun": 1, "haftalik_saat": 2,
                                        "bitis": "2026-10-10"})
    db.insert("tatiller", {"ad": "Cumhuriyet", "baslangic": "2026-10-28", "bitis": "2026-10-29"})
    s = db.insert("siniflar", {"ad": "12-A", "seviye": 12})
    o = db.insert("ogrenciler", {"ad": "Veli", "soyad": "Tan", "okul_no": "12", "sinif_id": s})
    db.insert("yerlestirmeler", {"ogrenci_id": o, "isletme_id": i1, "tur": "İşletmede Mesleki Eğitim",
                                 "baslangic": "2026-09-15", "durum": "Aktif"})
    return t1, t2, i1, i2


def test_formlar_uretim_ve_teslim(db, tmp_path):
    t1, t2, i1, i2 = _kur(db)
    ekim = (date(2026, 10, 1), date(2026, 10, 31))
    gunluk = bt.formlar(db, ["G"], *ekim)
    # Ali: 4 pazartesi (5,12,19,26) + 3 çarşamba (7,14,21; 28 tatil); Ece: yalnız 5 Ekim (görev 10'unda bitiyor)
    assert len([f for f in gunluk if f.ogretmen_id == t1]) == 7
    assert [f.tarih for f in gunluk if f.ogretmen_id == t2] == [date(2026, 10, 5)]
    aylik = bt.formlar(db, ["R", "A"], *ekim)
    assert sorted((f.tur, f.ogretmen_id, f.isletme_id) for f in aylik) == sorted(
        [("R", t1, i1), ("R", t1, i2), ("R", t2, i2), ("A", t1, None), ("A", t2, None)])
    a = next(f for f in aylik if f.tur == "A" and f.ogretmen_id == t1)
    assert len(a.gunler) == 7 and sum(g[2] for g in a.gunler) == 4 * 4 + 3 * 2

    yollar = bt.formlari_uret(db, gunluk + aylik, tmp_path)
    assert len(yollar) == 2 and all(y.exists() for y in yollar)
    doc = Document(str(next(y for y in yollar if "Ali" in y.name)))
    metin = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "GÜNLÜK REHBERLİK GÖREV FORMU" in metin and "Veli Tan" in metin and "Hasan Usta" in metin
    assert len(doc.inline_shapes) == 7 + 3  # her formda bir barkod
    assert db.q1("SELECT COUNT(*) c FROM belge_takibi")["c"] == 13
    assert bt.kaydet(db, gunluk) == 0  # ikinci kez eklenmez

    f = next(x for x in gunluk if x.ogretmen_id == t1 and x.tarih == date(2026, 10, 5))
    s = bt.teslim_al(db, f.barkod, date(2026, 10, 6))
    assert s.tamam and "Usta Oto" in s.mesaj
    z = db.q1("SELECT durum FROM ziyaretler WHERE ogretmen_id = ? AND tarih = '2026-10-05'", (t1,))
    assert z["durum"] == "Yapıldı"
    tekrar = bt.teslim_al(db, f.barkod)
    assert not tekrar.tamam and "Zaten" in tekrar.mesaj
    assert not bt.teslim_al(db, "999").tamam

    liste = bt.takip_listesi(db, *ekim, tur="G", ogretmen_id=t1)
    assert len(liste) == 7
    teslim = next(x for x in liste if x.teslim)
    assert teslim.durum == "Teslim alındı" and teslim.son_teslim == date(2026, 10, 10)
    assert len(bt.takip_listesi(db, *ekim, sadece_bekleyen=True)) == 12
    ozet = {r[0]: r[1:] for r in bt.ozet(bt.takip_listesi(db, *ekim))}
    assert ozet["Ali Can"][:2] == [10, 1]
    bt.teslimi_geri_al(db, teslim.id)
    assert not bt.takip_listesi(db, *ekim, tur="G", ogretmen_id=t1)[0].teslim


def test_listede_olmayan_formun_barkodu_da_teslim_alinir(db):
    t1, _, i1, _ = _kur(db)
    kod = bt.barkod_metni("R", date(2026, 9, 1), t1, i1)
    s = bt.teslim_al(db, kod, date(2026, 10, 2))
    assert s.tamam
    [satir] = bt.takip_listesi(db, date(2026, 9, 1), date(2026, 9, 30))
    assert satir.tur == "R" and satir.son_teslim == date(2026, 9, 30) + timedelta(days=5)
    assert satir.durum == "Teslim alındı"
    assert not bt.teslim_al(db, bt.barkod_metni("G", date(2026, 9, 1), 999, i1)).tamam
