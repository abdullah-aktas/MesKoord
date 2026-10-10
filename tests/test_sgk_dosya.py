from datetime import date

import pytest
from lxml import etree
from openpyxl import load_workbook

from koordmes.services import sgk_dosya
from koordmes.services.sgk_risk import aphb_anahtari

SICIL = "2 8510 01 01 1234567 063 01 12 000"


def _kurum(db, sicil=SICIL):
    db.update("kurum", 1, {"ad": "Şehit Ömer Halisdemir Mesleki ve Teknik Anadolu Lisesi", "il": "Ankara",
                           "ilce": "Çankaya", "sgk_isyeri_sicil": sicil, "adres": ""})


def _ogrenci(db, ad, soyad, tc, program="MTAL-AMP", statu="Öğrenci", yukumlu=1):
    s = db.q1("SELECT id FROM siniflar WHERE program = ?", (program,))
    sid = s["id"] if s else db.insert("siniflar", {"ad": f"12-{program}", "seviye": 12, "program": program})
    return db.insert("ogrenciler", {"ad": ad, "soyad": soyad, "tc": tc, "sinif_id": sid, "statu": statu,
                                    "bakmakla_yukumlu": yukumlu})


def _yerlestir(db, ogr, bas, bit=None, tur="İşletmede Mesleki Eğitim"):
    isl = db.q1("SELECT id FROM isletmeler") or {"id": db.insert("isletmeler", {"ad": "Usta Elektrik"})}
    return db.insert("yerlestirmeler", {"ogrenci_id": ogr, "isletme_id": isl["id"], "tur": tur, "baslangic": bas,
                                        "bitis": bit})


@pytest.fixture
def veri(db):
    _kurum(db)
    a = _ogrenci(db, "Ayşe", "Işıkçı", "10000000078")
    b = _ogrenci(db, "İsmail", "Öztürk", "10000000146", yukumlu=0)
    c = _ogrenci(db, "Hatalı", "Kimlik", "123")
    d = _ogrenci(db, "Mehmet", "Çırak", "10000000214", program="MEM", statu="Çırak")
    ids = {"a": _yerlestir(db, a, "2026-09-14"),
           "b": _yerlestir(db, b, "2026-10-05", "2026-10-20"),
           "c": _yerlestir(db, c, "2026-10-01"),
           "d": _yerlestir(db, d, "2026-10-01", tur="Çıraklık"),
           "okul": _yerlestir(db, a, "2026-10-01", "2026-10-31", tur="Okulda Alan Eğitimi")}
    return db, ids


def test_sicil_cozumleme():
    assert sgk_dosya.sicil_coz(SICIL) == ("285100101123456706301", "12", "000")
    assert sgk_dosya.sicil_coz("12345") is None
    assert sgk_dosya.tr_buyuk("ayşe ılgın") == "AYŞE ILGIN"


def test_sicil_yoksa_acik_hata(db):
    _kurum(db, "")
    with pytest.raises(sgk_dosya.SgkDosyaHatasi, match="26 hane"):
        sgk_dosya.isyeri(db)


def test_aphb_xml_sgk_semasina_uyar(veri, tmp_path):
    db, _ = veri
    sonuc = sgk_dosya.aphb_dosyalari(db, 2026, 10, tmp_path, "xml", bugun=date(2026, 11, 3))
    [yol] = sonuc.yollar
    assert sgk_dosya.xsd_dogrula(yol) == []
    assert not any("şemasına uymuyor" in u for u in sonuc.uyarilar)
    assert any("Hatalı" in u for u in sonuc.uyarilar)
    ham = yol.read_bytes()
    assert ham.startswith(b'<?xml version="1.0" encoding="UTF-8"?>\r\n')
    kok = etree.fromstring(ham)
    isy = kok.find("ISYERI")
    assert (isy.get("ISYERISICIL"), isy.get("KONTROLNO"), isy.get("ISYERIARACINO")) == \
           ("285100101123456706301", "12", "000")
    assert len(isy.get("ISYERIUNVAN")) <= 50 and isy.get("ISYERIADRES") == "Çankaya Ankara"
    bordro = kok.find("BORDRO")
    assert (bordro.get("DONEMAY"), bordro.get("DONEMYIL"), bordro.get("TAHAKKUKNEDEN"),
            bordro.get("BELGEMAHIYET")) == ("10", "2026", "A", "A")
    belgeler = {b.get("BELGETURU"): b for b in kok.iter("BILDIRGELER")}
    assert set(belgeler) >= {"07", "42", "49"} and all(b.get("KANUN") == "00000" for b in belgeler.values())
    ayse = belgeler["07"].find("SIGORTALILAR/SIGORTALI[@TCKNO='10000000078']")
    assert ayse.get("AD") == "AYŞE" and ayse.get("SOYAD") == "IŞIKÇI" and ayse.get("SIRA") == "1"
    assert ayse.get("GUN") == "30" and ayse.get("MESLEKKOD") == "9901.02" and ayse.get("GIRISGUN") is None
    ismail = belgeler["42"].find("SIGORTALILAR/SIGORTALI")
    assert ismail.get("GIRISGUN") == "0510" and ismail.get("CIKISGUN") == "2010"
    assert ismail.get("ISTENCIKISNEDENI") == "22" and float(ismail.get("PEK")) > 0
    assert sonuc.kisi == len(list(kok.iter("SIGORTALI")))


def test_aphb_gec_ve_ek_belge_uyarilari(veri, tmp_path):
    db, _ = veri
    db.ayar_yaz(aphb_anahtari(2026, 10), "2026-11-20")
    sonuc = sgk_dosya.aphb_dosyalari(db, 2026, 10, tmp_path, "xml", bugun=date(2026, 12, 15))
    kok = etree.parse(str(sonuc.yollar[0])).getroot()
    assert kok.find("BORDRO").get("TAHAKKUKNEDEN") == "B"
    assert any("yasal süresi" in u for u in sonuc.uyarilar) and any("'Ek'" in u for u in sonuc.uyarilar)
    ek = sgk_dosya.aphb_dosyalari(db, 2026, 10, tmp_path, "xml", mahiyet="E", tahakkuk="A")
    assert etree.parse(str(ek.yollar[0])).getroot().find("BORDRO").get("BELGEMAHIYET") == "E"
    assert sgk_dosya.xsd_dogrula(ek.yollar[0]) == []
    with pytest.raises(sgk_dosya.SgkDosyaHatasi):
        sgk_dosya.aphb_dosyalari(db, 2026, 10, tmp_path, mahiyet="X")
    with pytest.raises(sgk_dosya.SgkDosyaHatasi, match="sigortalı öğrenci yok"):
        sgk_dosya.aphb_dosyalari(db, 2020, 1, tmp_path)


def test_okul_icin_ayri_isyeri_ayri_dosya(veri, tmp_path):
    db, _ = veri
    db.ayar_yaz("sgk_isyeri_sicil_okul", "28510010199999990630112000")
    sonuc = sgk_dosya.aphb_dosyalari(db, 2026, 10, tmp_path, "xml", bugun=date(2026, 11, 3))
    assert len(sonuc.yollar) == 2
    okul = next(y for y in sonuc.yollar if "okul" in y.name)
    kok = etree.parse(str(okul)).getroot()
    assert kok.find("ISYERI").get("ISYERISICIL") == "285100101999999906301"
    assert {b.get("BELGETURU") for b in kok.iter("BILDIRGELER")} == {"49"}
    assert all(sgk_dosya.xsd_dogrula(y) == [] for y in sonuc.yollar)


def test_aphb_excel_sgk_sablonu_duzeni(veri, tmp_path):
    db, _ = veri
    sonuc = sgk_dosya.aphb_dosyalari(db, 2026, 10, tmp_path, "excel", bugun=date(2026, 11, 3))
    ws = load_workbook(sonuc.yollar[0])["Bildirge"]
    assert (ws["E2"].value, ws["H2"].value, ws["E3"].value) == ("285100101123456706301", "12", "000")
    assert (ws["E8"].value, ws["E9"].value, ws["E12"].value, ws["E13"].value) == ("2026", "10", "A", "A")
    satirlar = [[c.value for c in r] for r in ws.iter_rows(min_row=17, min_col=2, max_col=19) if r[0].value]
    assert len(satirlar) == sonuc.kisi
    ilk = satirlar[0]
    assert ilk[0] == 1 and ilk[1] == "07" and ilk[2] == "00000" and ilk[4] == "10000000078" and ilk[5] == "AYŞE"
    assert ilk[10] == 30 and ilk[16] == "9901.02"


def test_ise_giris_xml(veri, tmp_path):
    db, ids = veri
    sonuc = sgk_dosya.bildirge_dosyalari(db, [ids["a"], ids["d"], ids["c"]], "Giriş", tmp_path)
    [yol] = sonuc.yollar
    ham = yol.read_bytes()
    assert ham.startswith(b'<?xml version="1.0" encoding="ISO-8859-9"?>\r\n') and "AYŞE".encode("cp1254") in ham
    kok = etree.fromstring(ham)
    assert kok.tag == "SGK4AISEGIRIS"
    isy = kok.find("ISYERI")
    assert isy.get("ISYERISICIL") == "285100101123456706301" and isy.get("NAKILGELDIGIISYERISICIL") == ""
    kisiler = {s.get("TCKNO"): s for s in kok.iter("SIGORTALI")}
    assert set(kisiler) == {"10000000078", "10000000214"}
    ayse = kisiler["10000000078"]
    assert ayse.get("ISEGIRISTARIHI") == "2026-09-14" and ayse.get("SIGORTAKOLU") == "19"
    assert (ayse.get("GOREVKODU"), ayse.get("MESLEKKODU"), ayse.get("CSGBISKOLU")) == ("05", "9901.02", "10")
    assert kisiler["10000000214"].get("SIGORTAKOLU") == "7"
    assert any("Hatalı" in u for u in sonuc.uyarilar) and any("Bildirildi" in u for u in sonuc.uyarilar)


def test_isten_cikis_xml_donem_bilgileri(veri, tmp_path):
    db, ids = veri
    sonuc = sgk_dosya.bildirge_dosyalari(db, [ids["b"], ids["okul"]], "Çıkış", tmp_path)
    assert len(sonuc.yollar) == 1
    kok = etree.parse(str(sonuc.yollar[0])).getroot()
    assert kok.tag == "SGK4AISTENCIKIS" and kok.find("ISYERI").get("NAKILGIDECEGIISYERISICIL") == ""
    kisiler = {s.get("TCKNO"): s for s in kok.iter("SIGORTALI")}
    ismail = kisiler["10000000146"]
    assert ismail.get("ISTENCIKISTARIHI") == "2026-10-20" and ismail.get("ISTENAYRILISNEDENI") == "22"
    simdi = ismail.find("BULUNDUGUMUZDONEM")
    assert simdi.get("BELGETURU") == "42" and float(simdi.get("HAKEDILENUCRET")) > 0
    assert ismail.find("ONCEKIDONEM").get("BELGETURU") == "0"
    assert kisiler["10000000078"].find("BULUNDUGUMUZDONEM").get("BELGETURU") == "49"


def test_cikista_onceki_donem_verilmemisse_doldurulur(db, tmp_path):
    _kurum(db)
    o = _ogrenci(db, "Ali", "Kaya", "10000000078")
    y = _yerlestir(db, o, "2026-09-14", "2026-10-02")
    r = sgk_dosya._bildirge_satirlari(db, [y], "Çıkış", [])[0]
    simdi, once = sgk_dosya.cikis_donemleri(db, r)
    assert simdi.belge == "7" and once.belge == "7" and once.hakedilen > simdi.hakedilen
    db.ayar_yaz(aphb_anahtari(2026, 9), "2026-10-01")
    assert sgk_dosya.cikis_donemleri(db, r)[1].belge == "0"


def test_bildirge_excel(veri, tmp_path):
    db, ids = veri
    giris = sgk_dosya.bildirge_dosyalari(db, [ids["a"]], "Giriş", tmp_path, "excel")
    ws = load_workbook(giris.yollar[0])["4A İşe Giriş"]
    assert ws["E2"].value == "285100101123456706301" and ws["E3"].value == "000"
    assert [c.value for c in ws[14]][1:8] == [1, "10000000078", "AYŞE", "IŞIKÇI", "2026-09-14", "19", "H"]
    cikis = sgk_dosya.bildirge_dosyalari(db, [ids["b"]], "Çıkış", tmp_path, "excel")
    ws = load_workbook(cikis.yollar[0])["4A İşten Çıkış"]
    assert ws["K11"].value == "BULUNDUĞUMUZ DÖNEM" and ws["P11"].value == "ÖNCEKİ DÖNEM"
    satir = [c.value for c in ws[14]]
    assert satir[2] == "10000000146" and satir[5] == "2026-10-20" and satir[7] == "22" and satir[10] == "42"


def test_eylemler_asistan_ve_kapanis(ornek_db, tmp_path, monkeypatch):
    from koordmes import config
    from koordmes.services import asistan, surecler

    db = ornek_db
    monkeypatch.setattr(config, "output_dir", lambda: tmp_path)
    s = surecler.calistir(db, "sgk_xml", yil=2026, ay=10)
    assert s.dosyalar and s.dosyalar[0].suffix == ".xml" and sgk_dosya.xsd_dogrula(s.dosyalar[0]) == []
    assert surecler.calistir(db, "sgk_excel", yil=2026, ay=10).dosyalar[0].suffix == ".xlsx"
    y = asistan.yorumla(db, "ekim e-bildirge xml", date(2026, 11, 3))
    assert y.dosyalar and y.dosyalar[0].suffix == ".xml" and "Ekim 2026" in y.metin
    y = asistan.yorumla(db, "ekim sgk excel şablonu", date(2026, 11, 3))
    assert y.dosyalar and y.dosyalar[0].suffix == ".xlsx"
    k = surecler.aylik_kapanis(db, 2026, 10)
    sgk = tmp_path / "Aylık Kapanış" / "2026-10 Ekim" / "SGK e-Bildirge"
    assert sorted(p.suffix for p in sgk.iterdir()) == [".xlsx", ".xml"], k.mesaj
    ozet = (tmp_path / "Aylık Kapanış" / "2026-10 Ekim" / "0 Kapanış Özeti.txt").read_text("utf-8")
    assert "✓ SGK e-Bildirge V2 aylık prim ve hizmet belgesi (XML)" in ozet and "hatalı" not in k.mesaj


def test_ui_dugmeleri(ornek_db, tmp_path):
    import os
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from koordmes.ui.sayfa_sgk_ucret import SgkSayfasi

    db = ornek_db
    sayfa = SgkSayfasi(db)
    aphb = sayfa.widget(1)
    aphb.ay.ayarla(2026, 10)
    aphb.mahiyet.setCurrentIndex(aphb.mahiyet.findData("E"))
    aphb.tahakkuk.setCurrentIndex(aphb.tahakkuk.findData("B"))
    s = aphb.sgk_dosyasi("xml", tmp_path / "aphb", sessiz=True)
    kok = etree.parse(str(s.yollar[0])).getroot()
    assert (kok.find("BORDRO").get("BELGEMAHIYET"), kok.find("BORDRO").get("TAHAKKUKNEDEN")) == ("E", "B")
    assert aphb.sgk_dosyasi("excel", tmp_path / "aphb", sessiz=True).yollar[0].suffix == ".xlsx"
    aphb.ay.ayarla(2020, 1)
    assert aphb.sgk_dosyasi("xml", tmp_path / "aphb", sessiz=True) is None

    bildirim = sayfa.widget(0)
    bildirim.yenile()
    sonuc = bildirim.sgk_dosyasi("xml", tmp_path / "bildirge", sessiz=True)
    if bildirim.liste:
        assert sonuc is not None and sonuc.kisi > 0
        kokler = {etree.parse(str(p)).getroot().tag for p in sonuc.yollar}
        turler = {b.tur for b in bildirim.liste}
        assert kokler == {{"Giriş": "SGK4AISEGIRIS", "Çıkış": "SGK4AISTENCIKIS"}[t] for t in turler}
        assert bildirim.sgk_dosyasi("excel", tmp_path / "bildirge", sessiz=True).yollar

