"""SGK'ya yüklenecek dosyalar.

- e-Bildirge V2 aylık prim ve hizmet belgesi XML'i (SGK "EBildirgeV2 XML ŞEMA TANIMI" sürüm 3.1, veri/sgk).
- 4A toplu işe giriş ve işten ayrılış XML'leri (SGK "IsegirisExceldenXML" / "IstenCikisExceldenXML" şablonlarının
  ürettiği yapı).
- Aynı bilgiler, SGK'nın Excel şablonlarıyla aynı hücre düzeninde Excel dosyası olarak (şablona yapıştırılabilir).

İşyeri sicil numarası 26 hanedir: ilk 21 hane sicil, sonraki 2 hane kontrol no, son 3 hane aracı no.
Okulda alan eğitimi (belge 49/50) ayrı işyerinden bildiriliyorsa "sgk_isyeri_sicil_okul" parametresine yazılır;
her işyeri için ayrı dosya üretilir.
"""
from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from xml.sax.saxutils import quoteattr

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .. import config
from ..mevzuat import MEM_STATULERI as CIRAK_STATULERI, OKUL_ICI_TURLER, mem_mi
from . import sigorta
from .sgk_risk import aphb_son_gunu, aphb_verildi
from .utils import tarih, tc_gecerli

XSD_APHB = Path(__file__).resolve().parents[1] / "veri" / "sgk" / "EBildirgeV2.xsd"
MAHIYETLER = {"A": "Asıl", "E": "Ek", "I": "İptal"}
TAHAKKUK_NEDENLERI = {"A": "Yasal süresi içinde verilen", "B": "Yasal süresi dışında verilen"}
_TAB = "\t"
_SATIR = "\r\n"


class SgkDosyaHatasi(ValueError):
    pass


@dataclass
class Isyeri:
    sicil: str
    kontrol: str
    araci: str
    unvan: str
    adres: str
    okul_ici: bool = False

    @property
    def etiket(self) -> str:
        return f"{self.sicil}{self.kontrol}{self.araci}"


@dataclass
class DosyaSonucu:
    yollar: list[Path] = field(default_factory=list)
    kisi: int = 0
    uyarilar: list[str] = field(default_factory=list)

    def ozet(self) -> str:
        s = f"{self.kisi} kişi, {len(self.yollar)} dosya"
        return s + (f", {len(self.uyarilar)} uyarı" if self.uyarilar else "")


# ------------------------------------------------------------------ ortak

def sicil_coz(metin: str | None) -> tuple[str, str, str] | None:
    rakam = re.sub(r"\D", "", metin or "")
    if len(rakam) != 26:
        return None
    return rakam[:21], rakam[21:23], rakam[23:]


def isyeri(db, okul_ici: bool = False) -> Isyeri:
    k = db.kurum()
    ayri = str(db.ayar("sgk_isyeri_sicil_okul") or "").strip() if okul_ici else ""
    ham = ayri or k["sgk_isyeri_sicil"]
    parca = sicil_coz(ham)
    if parca is None:
        yer = ("Mevzuat Parametreleri > 'Okulda alan eğitimi için ayrı SGK işyeri sicil no'" if ayri else
               "Ayarlar > Kurum Bilgileri > SGK işyeri sicil no")
        raise SgkDosyaHatasi(f"SGK işyeri sicil numarası 26 hane olmalı ({yer}): "
                             f"'{ham or 'boş'}'. Örnek: 2 8510 01 01 1234567 063 01 12 000")
    adres = (k["adres"] or " ".join(x for x in (k["ilce"], k["il"]) if x) or k["ad"] or "-").strip()
    return Isyeri(*parca, unvan=(k["ad"] or "-").strip(), adres=adres, okul_ici=bool(ayri))


def _okul_ici_belge(db, belge: str) -> bool:
    return str(belge) in {str(db.ayar("belge_okul_yukumlu")), str(db.ayar("belge_okul_degil"))}


def _isyerleri(db, okul_ici_listesi: list[bool]) -> dict[bool, Isyeri]:
    """okul içi / işletme -> işyeri; ayrı okul sicili yoksa ikisi de kurum sicilidir."""
    out = {}
    for oi in sorted(set(okul_ici_listesi)):
        out[oi] = isyeri(db, oi)
    return out


def tr_buyuk(s: str | None) -> str:
    return (s or "").strip().replace("i", "İ").replace("ı", "I").upper()


def _kisalt(s: str, n: int, uyarilar: list[str], kim: str, alan: str) -> str:
    if len(s) > n:
        uyarilar.append(f"{kim}: {alan} {n} karakterden uzun, '{s[:n]}' olarak kısaltıldı; e-Bildirge'de kontrol edin.")
        return s[:n]
    return s


def _tutar(x: float) -> str:
    return f"{x:.2f}"


def _ggaa(d: date | None) -> str:
    return f"{d:%d%m}" if d else ""


def _attr(ad: str, deger) -> str:
    return f"{ad}={quoteattr(str(deger))}"


def _yaz(yol: Path, satirlar: list[str], kodlama: str) -> Path:
    yol.parent.mkdir(parents=True, exist_ok=True)
    yol.write_bytes(_SATIR.join(satirlar + [""]).encode(kodlama, errors="replace"))
    return yol


def _klasor(klasor: Path | str | None) -> Path:
    return Path(klasor) if klasor else config.output_dir()


def _dosya_adi(temel: str, isy: Isyeri, coklu: bool, uzanti: str) -> str:
    return f"{temel}{' - okul' if coklu and isy.okul_ici else ''}{' - işletme' if coklu and not isy.okul_ici else ''}" \
           f".{uzanti}"


# ------------------------------------------------------------------ aylık prim ve hizmet belgesi

@dataclass
class AphbSatiri:
    belge: str
    kanun: str
    tc: str
    ad: str
    soyad: str
    pek: float
    gun: int
    giris: date | None
    cikis: date | None
    eksik_gun: int
    eksik_neden: str
    cikis_nedeni: str
    meslek: str
    okul_ici: bool
    ogrenci_id: int | None = None


def tahakkuk_nedeni(db, yil: int, ay: int, bugun: date | None = None) -> str:
    return "A" if (bugun or date.today()) <= aphb_son_gunu(db, yil, ay) else "B"


def aphb_satirlari(db, yil: int, ay: int, uyarilar: list[str] | None = None) -> list[AphbSatiri]:
    uyarilar = [] if uyarilar is None else uyarilar
    kanun = str(db.ayar("sgk_kanun_no") or "00000").zfill(5)
    meslek = str(db.ayar("sgk_meslek_kodu") or "")
    cikis_nedeni = str(db.ayar("sgk_cikis_nedeni") or "")
    out = []
    for s in sigorta.aylik_liste(db, yil, ay):
        kim = f"{s.ad} {s.soyad}"
        if not tc_gecerli(s.tc):
            uyarilar.append(f"{kim}: T.C. kimlik no geçersiz ('{s.tc}'), dosyaya alınmadı.")
            continue
        ad = _kisalt(tr_buyuk(s.ad), 18, uyarilar, kim, "ad")
        soyad = _kisalt(tr_buyuk(s.soyad), 18, uyarilar, kim, "soyad")
        out.append(AphbSatiri(str(s.belge_turu).zfill(2), kanun, s.tc, ad, soyad, s.pek, s.gun, s.giris_tarihi,
                              s.cikis_tarihi, s.eksik_gun, s.eksik_neden if s.eksik_gun else "",
                              cikis_nedeni if s.cikis_tarihi else "", meslek, _okul_ici_belge(db, s.belge_turu),
                              s.ogrenci_id))
    out.sort(key=lambda r: (r.belge, r.kanun, r.ad, r.soyad))
    return out


def _aphb_xml(isy: Isyeri, yil: int, ay: int, mahiyet: str, tahakkuk: str, satirlar: list[AphbSatiri]) -> list[str]:
    x = ['<?xml version="1.0" encoding="UTF-8"?>', '<AYLIKBILDIRGELER VER="2">']
    x.append(_TAB + "<ISYERI " + " ".join([
        _attr("ISYERISICIL", isy.sicil), _attr("KONTROLNO", isy.kontrol), _attr("ISYERIARACINO", isy.araci),
        _attr("ISYERIUNVAN", isy.unvan[:50]), _attr("ISYERIADRES", isy.adres[:50])]) + "/>")
    x.append(_TAB + "<BORDRO " + " ".join([
        _attr("DONEMAY", ay), _attr("DONEMYIL", yil), _attr("HIZMETDONEMAY", ay), _attr("HIZMETDONEMYIL", yil),
        _attr("TAHAKKUKNEDEN", tahakkuk), _attr("BELGEMAHIYET", mahiyet)]) + "/>")
    gruplar: dict[tuple[str, str], list[AphbSatiri]] = defaultdict(list)
    for r in satirlar:
        gruplar[(r.belge, r.kanun)].append(r)
    for (belge, kanun), grup in gruplar.items():
        x.append(_TAB + f"<BILDIRGELER {_attr('BELGETURU', belge)} {_attr('KANUN', kanun)}>")
        x.append(_TAB * 2 + "<SIGORTALILAR>")
        for i, r in enumerate(grup, 1):
            a = [_attr("SIRA", i), _attr("TCKNO", r.tc), _attr("AD", r.ad), _attr("SOYAD", r.soyad),
                 _attr("PEK", _tutar(r.pek)), _attr("GUN", r.gun)]
            if r.giris:
                a.append(_attr("GIRISGUN", _ggaa(r.giris)))
            if r.cikis:
                a.append(_attr("CIKISGUN", _ggaa(r.cikis)))
            if r.eksik_gun:
                a.append(_attr("EKSIKGUNSAYISI", r.eksik_gun))
                if r.eksik_neden:
                    a.append(_attr("EKSIKGUNNEDENI", r.eksik_neden))
            if r.cikis_nedeni:
                a.append(_attr("ISTENCIKISNEDENI", r.cikis_nedeni))
            if r.meslek:
                a.append(_attr("MESLEKKOD", r.meslek))
            x.append(_TAB * 3 + "<SIGORTALI " + " ".join(a) + "/>")
        x.append(_TAB * 2 + "</SIGORTALILAR>")
        x.append(_TAB + "</BILDIRGELER>")
    x.append("</AYLIKBILDIRGELER>")
    return x


def xsd_dogrula(yol: Path | str, xsd: Path = XSD_APHB) -> list[str]:
    from lxml import etree
    sema = etree.XMLSchema(etree.parse(str(xsd)))
    belge = etree.parse(str(yol))
    if sema.validate(belge):
        return []
    return [f"satır {e.line}: {e.message}" for e in sema.error_log][:20]


# ------------------------------------------------------------------ Excel (SGK şablonu düzeni)

_BASLIK = Font(bold=True)
_DOLGU = PatternFill("solid", fgColor="DDEBF7")


def _hucreler(ws, degerler: dict[str, object]) -> None:
    for adres, v in degerler.items():
        ws[adres] = v
        ws[adres].number_format = "@"


def _baslik_satirlari(ws, satirlar: dict[int, dict[str, str]]) -> None:
    for satir, hucre in satirlar.items():
        for sutun, v in hucre.items():
            c = ws[f"{sutun}{satir}"]
            c.value = v
            c.font = _BASLIK
            c.fill = _DOLGU
            c.alignment = Alignment(horizontal="center", wrap_text=True)


def _veri(ws, ilk_satir: int, ilk_sutun: int, satirlar: list[list[object]]) -> None:
    for i, r in enumerate(satirlar):
        for j, v in enumerate(r):
            c = ws.cell(row=ilk_satir + i, column=ilk_sutun + j, value=v)
            if isinstance(v, str):
                c.number_format = "@"
    for sutun in "BCDEFGHIJKLMNOPQRST":
        ws.column_dimensions[sutun].width = 14
    ws.column_dimensions["B"].width = 30


def _aphb_excel(isy: Isyeri, yil: int, ay: int, mahiyet: str, tahakkuk: str, satirlar: list[AphbSatiri],
                yol: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bildirge"
    _baslik_satirlari(ws, {1: {"B": "İŞYERİ"}, 7: {"B": "BORDRO"}})
    for satir, etiket in [(2, "Sicil No"), (3, "Aracı No"), (4, "Ünvan"), (5, "Adres"), (6, "Vergi No"),
                          (8, "Yıl (yyyy)"), (9, "Ay (aa)"), (10, "Hizmet Yıl (yyyy)"), (11, "Hizmet Ay (aa)"),
                          (12, "Mahiyet"), (13, "Tahakkuk Nedeni")]:
        ws[f"B{satir}"] = etiket
    ws["G2"] = "Kontrol No"
    ws["F12"] = "(Asıl için A, Ek için E, İptal için I)"
    _hucreler(ws, {"E2": isy.sicil, "H2": isy.kontrol, "E3": isy.araci, "E4": isy.unvan[:50], "E5": isy.adres[:50],
                   "E8": str(yil), "E9": f"{ay:02d}", "E10": str(yil), "E11": f"{ay:02d}", "E12": mahiyet,
                   "E13": tahakkuk})
    _baslik_satirlari(ws, {15: dict(zip("BCDEFGHIJKLMNOPQRS", [
        "Sıra No", "Belge Türü", "Kanun No (5 hane)", "Sigortalı Sicil Numarası", "SGNO (T.C. Kimlik) 11 hane",
        "Sigortalı Adı", "Sigortalı Soyadı", "Sigortalı İlk Soyadı", "ÜCRET TL", "İKRAMİYE TL", "Prim Öd. Günü",
        "İşe Giriş G. (GGAA)", "İşten Çıkış G. (GGAA)", "Eksik Gün Sayısı", "Eksik Gün Nedeni", "İşten Çıkış Nedeni",
        "Meslek Kod", "İstirahat Sürelerinde Çalışmamıştır"]))})
    _veri(ws, 17, 2, [[i, r.belge, r.kanun, "", r.tc, r.ad, r.soyad, "", round(r.pek, 2), 0, r.gun,
                       _ggaa(r.giris), _ggaa(r.cikis), r.eksik_gun or "", r.eksik_neden, r.cikis_nedeni, r.meslek, ""]
                      for i, r in enumerate(satirlar, 1)])
    yol.parent.mkdir(parents=True, exist_ok=True)
    wb.save(yol)
    return yol


def aphb_yaz(db, yil: int, ay: int, satirlar: list[AphbSatiri], mahiyet: str, tahakkuk: str,
             klasor: Path | str | None = None, bicim: str = "xml", temel: str | None = None,
             sonuc: DosyaSonucu | None = None) -> DosyaSonucu:
    """Verilen satırlarla APHB dosyalarını yazar (her işyeri için ayrı dosya); sonuc verilirse ona eklenir."""
    sonuc = sonuc if sonuc is not None else DosyaSonucu()
    isyerleri = _isyerleri(db, [r.okul_ici for r in satirlar])
    gruplar: dict[str, tuple[Isyeri, list[AphbSatiri]]] = {}
    for r in satirlar:
        isy = isyerleri[r.okul_ici]
        gruplar.setdefault(isy.etiket, (isy, []))[1].append(r)
    coklu = len(gruplar) > 1
    hedef = _klasor(klasor)
    temel = temel or f"e-Bildirge APHB {yil}-{ay:02d} {MAHIYETLER[mahiyet]}"
    for isy, grup in gruplar.values():
        if bicim == "excel":
            yol = _aphb_excel(isy, yil, ay, mahiyet, tahakkuk, grup, hedef / _dosya_adi(temel, isy, coklu, "xlsx"))
        else:
            yol = _yaz(hedef / _dosya_adi(temel, isy, coklu, "xml"), _aphb_xml(isy, yil, ay, mahiyet, tahakkuk, grup),
                       "utf-8")
            hatalar = xsd_dogrula(yol)
            if hatalar:
                sonuc.uyarilar[:0] = [f"{yol.name} SGK şemasına uymuyor: {h}" for h in hatalar]
        sonuc.yollar.append(yol)
    sonuc.kisi += len(satirlar)
    if any(len(isy.unvan) > 50 or len(isy.adres) > 50 for isy, _ in gruplar.values()):
        uyari = "e-Bildirge işyeri ünvanı/adresi en çok 50 karakter alır; dosyada kısaltıldı."
        if uyari not in sonuc.uyarilar:
            sonuc.uyarilar.append(uyari)
    return sonuc


def aphb_dosyalari(db, yil: int, ay: int, klasor: Path | str | None = None, bicim: str = "xml",
                   mahiyet: str = "A", tahakkuk: str | None = None, bugun: date | None = None) -> DosyaSonucu:
    """bicim: 'xml' | 'excel'. Her işyeri için ayrı dosya. İptal belgesi, kayıtlı beyan varsa beyan edilen
    satırlarla (SGK'daki değerlerle) hazırlanır."""
    from .sgk_beyan import beyan_edilen
    if mahiyet not in MAHIYETLER:
        raise SgkDosyaHatasi(f"Belge mahiyeti A, E veya I olmalı: {mahiyet}")
    tahakkuk = tahakkuk or tahakkuk_nedeni(db, yil, ay, bugun)
    sonuc = DosyaSonucu()
    beyan = list(beyan_edilen(db, yil, ay).values()) if mahiyet == "I" else []
    if beyan:
        satirlar = sorted(beyan, key=lambda r: (r.belge, r.kanun, r.ad, r.soyad))
        sonuc.uyarilar.append("İptal belgesi, programda kayıtlı beyandaki (SGK'ya verilen) değerlerle hazırlandı. "
                              "Yalnızca değişen kişiler için SGK > Aylık Beyan sekmesindeki düzeltmeyi kullanın.")
    else:
        satirlar = aphb_satirlari(db, yil, ay, sonuc.uyarilar)
    if not satirlar:
        raise SgkDosyaHatasi(f"{ay:02d}/{yil} döneminde bildirilecek sigortalı öğrenci yok.")
    aphb_yaz(db, yil, ay, satirlar, mahiyet, tahakkuk, klasor, bicim, sonuc=sonuc)
    if tahakkuk == "B":
        sonuc.uyarilar.append(f"{ay:02d}/{yil} belgesinin yasal süresi ({aphb_son_gunu(db, yil, ay):%d.%m.%Y}) "
                              "geçti; tahakkuk nedeni B (yasal süresi dışında) yazıldı, idari para cezası doğabilir.")
    if aphb_verildi(db, yil, ay) and mahiyet == "A":
        sonuc.uyarilar.append(f"{ay:02d}/{yil} belgesi verildi olarak işaretli; değişen ya da eksik kalan "
                              "sigortalılar için 'İptal' + 'Ek' mahiyetinde düzeltme gerekir (SGK > Aylık Beyan).")
    return sonuc


# ------------------------------------------------------------------ 4A işe giriş / işten ayrılış

@dataclass
class BildirgeSatiri:
    yerlestirme_id: int
    tc: str
    ad: str
    soyad: str
    tarih: date
    sigorta_kolu: str
    okul_ici: bool
    belge: str
    ogrenci_id: int


def sigorta_kolu(db, tur: str, statu: str | None, program: str | None) -> str:
    if mem_mi(program, statu, tur):
        return str(db.ayar("sgk_sigorta_kolu_mem"))
    if tur in OKUL_ICI_TURLER or tur == "Tamamlayıcı Eğitim":
        return str(db.ayar("sgk_sigorta_kolu_okul"))
    return str(db.ayar("sgk_sigorta_kolu_isletme"))


def _bildirge_satirlari(db, yerlestirme_idleri: list[int], tur: str, uyarilar: list[str]) -> list[BildirgeSatiri]:
    if not yerlestirme_idleri:
        return []
    soru = ",".join("?" * len(yerlestirme_idleri))
    rows = db.q(f"""SELECT y.id, y.tur, y.baslangic, y.bitis, y.ogrenci_id, o.tc, o.ad, o.soyad, o.statu,
                           o.bakmakla_yukumlu, o.baska_sigortali, o.dogum_tarihi, s.program
                    FROM yerlestirmeler y JOIN ogrenciler o ON o.id = y.ogrenci_id
                    LEFT JOIN siniflar s ON s.id = o.sinif_id WHERE y.id IN ({soru})
                    ORDER BY o.ad, o.soyad""", tuple(yerlestirme_idleri))
    out = []
    for r in rows:
        kim = f"{r['ad']} {r['soyad']}"
        t = tarih(r["baslangic"] if tur == "Giriş" else r["bitis"])
        if not tc_gecerli(r["tc"]):
            uyarilar.append(f"{kim}: T.C. kimlik no geçersiz ('{r['tc'] or ''}'), dosyaya alınmadı.")
            continue
        if t is None:
            uyarilar.append(f"{kim}: {'başlangıç' if tur == 'Giriş' else 'bitiş'} tarihi yok, dosyaya alınmadı.")
            continue
        out.append(BildirgeSatiri(r["id"], r["tc"], _kisalt(tr_buyuk(r["ad"]), 50, uyarilar, kim, "ad"),
                                  _kisalt(tr_buyuk(r["soyad"]), 50, uyarilar, kim, "soyad"), t,
                                  sigorta_kolu(db, r["tur"], r["statu"], r["program"]), r["tur"] in OKUL_ICI_TURLER,
                                  sigorta.belge_turu(db, r["tur"], sigorta.gss_primi_yok(db, r, t)), r["ogrenci_id"]))
    return out


def _isyeri_satiri(isy: Isyeri, nakil_attr: str) -> str:
    a = [_attr("ISYERISICIL", isy.sicil), _attr("ISYERIARACINO", isy.araci), _attr(nakil_attr, ""),
         _attr("ISYERIUNVAN", isy.unvan[:250]), _attr("ISYERIADRES", isy.adres[:200])]
    return _TAB + "<ISYERI " + " ".join(a) + " />"


def _giris_xml(db, isy: Isyeri, satirlar: list[BildirgeSatiri]) -> list[str]:
    x = ['<?xml version="1.0" encoding="ISO-8859-9"?>', "<SGK4AISEGIRIS>", _isyeri_satiri(isy, "NAKILGELDIGIISYERISICIL"),
         _TAB + "<SIGORTALILAR>"]
    for r in satirlar:
        x.append(_TAB * 2 + "<SIGORTALI " + " ".join([
            _attr("TCKNO", r.tc), _attr("AD", r.ad), _attr("SOYAD", r.soyad),
            _attr("ISEGIRISTARIHI", r.tarih.isoformat()), _attr("SIGORTAKOLU", r.sigorta_kolu),
            _attr("OZURLUKODU", "H"), _attr("KISMICALISMA", "H"), _attr("KISMICALISMAGUNSAYISI", 0),
            _attr("ESKIHUKUMLU", "H"), _attr("OGRENIMKODU", 0), _attr("MEZUNIYETYILI", 0),
            _attr("MEZUNIYETBOLUMU", ""), _attr("CSGBISKOLU", db.ayar("sgk_csgb_iskolu")),
            _attr("MESLEKKODU", db.ayar("sgk_meslek_kodu")), _attr("GOREVKODU", db.ayar("sgk_gorev_kodu"))]) + " />")
    x += [_TAB + "</SIGORTALILAR>", "</SGK4AISEGIRIS>"]
    return x


@dataclass
class DonemBilgisi:
    belge: str = "0"
    hakedilen: float = 0.0
    eksik_gun: int = 0
    eksik_neden: str = "0"


def _donem_bilgisi(liste: list[sigorta.SigortaSatiri], ogrenci_id: int, belge: str) -> DonemBilgisi | None:
    for s in liste:
        if s.ogrenci_id == ogrenci_id and str(s.belge_turu) == str(belge):
            return DonemBilgisi(str(s.belge_turu), s.pek, s.eksik_gun, s.eksik_neden or "0")
    return None


def cikis_donemleri(db, r: BildirgeSatiri) -> tuple[DonemBilgisi, DonemBilgisi]:
    """(bulunduğumuz dönem, önceki dönem). Önceki dönem yalnızca o ayın belgesi henüz verilmemişse doldurulur."""
    y, a = r.tarih.year, r.tarih.month
    simdi = _donem_bilgisi(sigorta.aylik_liste(db, y, a), r.ogrenci_id, r.belge) or DonemBilgisi(r.belge)
    oy, oa = (y, a - 1) if a > 1 else (y - 1, 12)
    once = DonemBilgisi()
    if r.tarih <= aphb_son_gunu(db, oy, oa) and not aphb_verildi(db, oy, oa):
        once = _donem_bilgisi(sigorta.aylik_liste(db, oy, oa), r.ogrenci_id, r.belge) or once
    return simdi, once


def _donem_attr(etiket: str, d: DonemBilgisi) -> str:
    return _TAB * 3 + f"<{etiket} " + " ".join([
        _attr("BELGETURU", int(d.belge or 0)), _attr("HAKEDILENUCRET", _tutar(d.hakedilen)), _attr("PRIMIKRAMIYE", 0),
        _attr("EKSIKGUNSAYISI", d.eksik_gun), _attr("EKSIKGUNNEDENI", d.eksik_neden or "0")]) + "/>"


def _cikis_xml(db, isy: Isyeri, satirlar: list[BildirgeSatiri]) -> list[str]:
    x = ['<?xml version="1.0" encoding="ISO-8859-9"?>', "<SGK4AISTENCIKIS>",
         _isyeri_satiri(isy, "NAKILGIDECEGIISYERISICIL"), _TAB + "<SIGORTALILAR>"]
    for r in satirlar:
        simdi, once = cikis_donemleri(db, r)
        x.append(_TAB * 2 + "<SIGORTALI " + " ".join([
            _attr("TCKNO", r.tc), _attr("AD", r.ad), _attr("SOYAD", r.soyad),
            _attr("ISTENCIKISTARIHI", r.tarih.isoformat()), _attr("MESLEKKODU", db.ayar("sgk_meslek_kodu")),
            _attr("ISTENAYRILISNEDENI", str(db.ayar("sgk_cikis_nedeni")).zfill(2)),
            _attr("CSGBISKOLU", db.ayar("sgk_csgb_iskolu")), _attr("UCRETYUZDEUSULU", "H")]) + " >")
        x += [_donem_attr("ONCEKIDONEM", once), _donem_attr("BULUNDUGUMUZDONEM", simdi), _TAB * 2 + "</SIGORTALI>"]
    x += [_TAB + "</SIGORTALILAR>", "</SGK4AISTENCIKIS>"]
    return x


def _giris_excel(db, isy: Isyeri, satirlar: list[BildirgeSatiri], yol: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "4A İşe Giriş"
    _baslik_satirlari(ws, {1: {"B": "İŞYERİ"}})
    for satir, etiket in [(2, "İşyeri Sicil No"), (3, "İşyeri Aracı No"), (4, "İşyeri Unvan"), (5, "İşyeri Adres"),
                          (6, "Nakil Geldiği İşyeri Sicil")]:
        ws[f"B{satir}"] = etiket
    _hucreler(ws, {"E2": isy.sicil, "E3": isy.araci, "E4": isy.unvan, "E5": isy.adres})
    _baslik_satirlari(ws, {12: dict(zip("BCDEFGHIJKLMNOPQ", [
        "Sıra No", "TCKimlik Numarası", "Sigortalı Adı", "Sigortalı Soyadı", "İşe Giriş Tarihi", "Sigorta Kolu",
        "Özürlü", "Eski Hükümlü", "Öğrenim Kodu", "Mezuniyet yılı", "Mezuniyet Bölümü", "CSGB ISKOLU", "Meslek Kodu",
        "Görev Kodu", "Kısmi Çalışma", "Kısmi Çalışma Gün Sayısı"]))})
    _veri(ws, 14, 2, [[i, r.tc, r.ad, r.soyad, r.tarih.isoformat(), r.sigorta_kolu, "H", "H", "0", "0", "",
                       str(db.ayar("sgk_csgb_iskolu")), str(db.ayar("sgk_meslek_kodu")), str(db.ayar("sgk_gorev_kodu")),
                       "H", "0"] for i, r in enumerate(satirlar, 1)])
    yol.parent.mkdir(parents=True, exist_ok=True)
    wb.save(yol)
    return yol


def _cikis_excel(db, isy: Isyeri, satirlar: list[BildirgeSatiri], yol: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "4A İşten Çıkış"
    _baslik_satirlari(ws, {1: {"B": "İŞYERİ"}, 11: {"K": "BULUNDUĞUMUZ DÖNEM", "P": "ÖNCEKİ DÖNEM"}})
    for satir, etiket in [(2, "İşyeri Sicil No"), (3, "İşyeri Aracı No"), (4, "İşyeri Unvan"), (5, "İşyeri Adres"),
                          (6, "Nakil Gideceği İşyeri Sicil")]:
        ws[f"B{satir}"] = etiket
    _hucreler(ws, {"E2": isy.sicil, "E3": isy.araci, "E4": isy.unvan, "E5": isy.adres})
    donem = ["Belge Türü", "Hakedilen Ücret", "Prim İkramiye", "Eksik Gün Sayısı", "Eksik Gün Nedeni"]
    _baslik_satirlari(ws, {12: dict(zip("BCDEFGHIJKLMNOPQRST", [
        "Sıra No", "TCKimlik Numarası", "Sigortalı Adı", "Sigortalı Soyadı", "İşten Çıkış Tarihi", "Meslek Kodu",
        "İşten Ayrılış Nedeni", "CSGB ISKOLU", "Ücret Yüzde Usulü"] + donem + donem))})
    rows = []
    for i, r in enumerate(satirlar, 1):
        simdi, once = cikis_donemleri(db, r)
        rows.append([i, r.tc, r.ad, r.soyad, r.tarih.isoformat(), str(db.ayar("sgk_meslek_kodu")),
                     str(db.ayar("sgk_cikis_nedeni")).zfill(2), str(db.ayar("sgk_csgb_iskolu")), "H",
                     simdi.belge, round(simdi.hakedilen, 2), 0, simdi.eksik_gun, simdi.eksik_neden or "0",
                     once.belge, round(once.hakedilen, 2), 0, once.eksik_gun, once.eksik_neden or "0"])
    _veri(ws, 14, 2, rows)
    yol.parent.mkdir(parents=True, exist_ok=True)
    wb.save(yol)
    return yol


def bildirge_dosyalari(db, yerlestirme_idleri: list[int], tur: str, klasor: Path | str | None = None,
                       bicim: str = "xml") -> DosyaSonucu:
    """tur: 'Giriş' | 'Çıkış'. Toplu işe giriş / işten ayrılış dosyası; her işyeri için ayrı dosya."""
    sonuc = DosyaSonucu()
    satirlar = _bildirge_satirlari(db, list(yerlestirme_idleri), tur, sonuc.uyarilar)
    if not satirlar:
        raise SgkDosyaHatasi("\n".join(sonuc.uyarilar) or f"Dosyaya alınacak {tur.lower()} bildirimi yok.")
    isyerleri = _isyerleri(db, [r.okul_ici for r in satirlar])
    gruplar: dict[str, tuple[Isyeri, list[BildirgeSatiri]]] = {}
    for r in satirlar:
        isy = isyerleri[r.okul_ici]
        gruplar.setdefault(isy.etiket, (isy, []))[1].append(r)
    coklu = len(gruplar) > 1
    hedef = _klasor(klasor)
    gunler = sorted({r.tarih for r in satirlar})
    temel = f"SGK {'İşe Giriş' if tur == 'Giriş' else 'İşten Ayrılış'} {gunler[0]:%Y-%m-%d}" + \
            (f" - {gunler[-1]:%Y-%m-%d}" if len(gunler) > 1 else "")
    for isy, grup in gruplar.values():
        if bicim == "excel":
            fn = _giris_excel if tur == "Giriş" else _cikis_excel
            yol = fn(db, isy, grup, hedef / _dosya_adi(temel, isy, coklu, "xlsx"))
        else:
            fn = _giris_xml if tur == "Giriş" else _cikis_xml
            yol = _yaz(hedef / _dosya_adi(temel, isy, coklu, "xml"), fn(db, isy, grup), "iso-8859-9")
        sonuc.yollar.append(yol)
    sonuc.kisi = len(satirlar)
    if tur == "Giriş" and any(r.tarih <= date.today() for r in satirlar):
        sonuc.uyarilar.append("Başlangıç tarihi bugün veya geçmiş öğrenciler var: işe giriş bildirgesi geç kalmış "
                              "olabilir (SGK > Risk ve Öneriler sekmesinde tahmini ceza ve indirim görünür).")
    sonuc.uyarilar.append("SGK sistemine yükleyip onayladıktan sonra Giriş/Çıkış Bildirimleri sekmesinde "
                          "'Bildirildi' işaretlemeyi unutmayın.")
    return sonuc
