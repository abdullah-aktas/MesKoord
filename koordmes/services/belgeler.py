"""Word (.docx) belge üretimi.

Şablonlar Belgelerim/MesKoord/Sablonlar klasöründedir. İlk kullanımda varsayılan şablonlar
oluşturulur; okul kendi resmî formatına göre Word'de düzenleyebilir. {{alan_adi}} biçimindeki
yer tutucular üretim sırasında doldurulur.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Callable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .. import config
from ..mevzuat import AYLAR, GUNLER
from . import ucret as ucret_srv
from .utils import gunler_listesi, para, tr

YER_TUTUCU = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")


# ------------------------------------------------------------------ doldurma

def _paragraf_doldur(p, baglam: dict) -> None:
    if "{{" not in p.text:
        return
    yeni = YER_TUTUCU.sub(lambda m: "" if baglam.get(m.group(1)) is None else str(baglam[m.group(1)]), p.text)
    if yeni == p.text:
        return
    runs = p.runs
    if not runs:
        p.add_run(yeni)
        return
    runs[0].text = yeni
    for r in runs[1:]:
        r.text = ""


def _tablolari_doldur(tablolar, baglam: dict) -> None:
    for t in tablolar:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _paragraf_doldur(p, baglam)
                _tablolari_doldur(cell.tables, baglam)


LOGO_YER_TUTUCU = "{{kurum_logosu}}"
LOGO_YUKSEKLIK = Cm(2.0)


def _logo_koy(doc, logo: Path | None) -> None:
    """{{kurum_logosu}} paragrafına kurum logosunu koyar; yer tutucusu olmayan eski şablonlarda logo
    'T.C.' başlığının üstüne eklenir. Logo yoksa yer tutucu paragrafı kaldırılır."""
    paragraflar = list(doc.paragraphs)
    for s in doc.sections:
        paragraflar += list(s.header.paragraphs)
    hedefler = [p for p in paragraflar if p.text.strip() == LOGO_YER_TUTUCU]
    if not hedefler and logo and doc.paragraphs and doc.paragraphs[0].text.strip() == "T.C.":
        hedefler = [doc.paragraphs[0].insert_paragraph_before()]
    for p in hedefler:
        if not logo:
            p._element.getparent().remove(p._element)
            continue
        for r in p.runs:
            r.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        (p.runs[0] if p.runs else p.add_run()).add_picture(str(logo), height=LOGO_YUKSEKLIK)


def doldur(sablon: Path, baglam: dict, cikti: Path) -> Path:
    doc = Document(str(sablon))
    _logo_koy(doc, baglam.get("_logo"))
    for p in doc.paragraphs:
        _paragraf_doldur(p, baglam)
    _tablolari_doldur(doc.tables, baglam)
    for s in doc.sections:
        for part in (s.header, s.footer):
            for p in part.paragraphs:
                _paragraf_doldur(p, baglam)
            _tablolari_doldur(part.tables, baglam)
    cikti.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(cikti))
    return cikti


# ------------------------------------------------------------------ varsayılan şablonlar

def _yeni_belge() -> Document:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2)
        s.left_margin = s.right_margin = Cm(2.2)
    return doc


def _ortali(doc, metin: str, kalin: bool = True, boyut: int = 12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(metin)
    r.bold = kalin
    r.font.size = Pt(boyut)
    return p


def _kurum_basligi(doc) -> None:
    _ortali(doc, LOGO_YER_TUTUCU, kalin=False)
    _ortali(doc, "T.C.")
    _ortali(doc, "{{kurum_ilce}} KAYMAKAMLIĞI")
    _ortali(doc, "{{kurum_adi}} Müdürlüğü")


def _tablo(doc, satirlar: list[tuple[str, str]], genislik=(5.5, 11)):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for a, b in satirlar:
        c = t.add_row().cells
        c[0].text, c[1].text = a, b
        c[0].paragraphs[0].runs[0].bold = True
        c[0].width, c[1].width = Cm(genislik[0]), Cm(genislik[1])
    return t


def _imza_tablosu(doc, unvanlar: list[tuple[str, str]]):
    doc.add_paragraph()
    t = doc.add_table(rows=3, cols=len(unvanlar))
    for i, (u, ad) in enumerate(unvanlar):
        t.cell(0, i).text = u
        t.cell(1, i).text = ad
        t.cell(2, i).text = "İmza:"
        for r in range(3):
            t.cell(r, i).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        t.cell(0, i).paragraphs[0].runs[0].bold = True


def _maddeler(doc, maddeler: list[tuple[str, str]]) -> None:
    for i, (baslik, metin) in enumerate(maddeler, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"MADDE {i} - {baslik}: ")
        r.bold = True
        p.add_run(metin)


def _sablon_isletme_sozlesmesi(yol: Path) -> None:
    doc = _yeni_belge()
    _kurum_basligi(doc)
    _ortali(doc, "İŞLETMELERDE MESLEKİ EĞİTİM SÖZLEŞMESİ", boyut=13)
    doc.add_paragraph("Sözleşme No: {{sozlesme_no}}        Sözleşme Tarihi: {{sozlesme_tarihi}}")
    doc.add_paragraph().add_run("A) ÖĞRENCİ BİLGİLERİ").bold = True
    _tablo(doc, [("Adı Soyadı", "{{ogrenci_ad_soyad}}"), ("T.C. Kimlik No", "{{ogrenci_tc}}"),
                 ("Okul No / Sınıf", "{{ogrenci_no}} / {{sinif}}"), ("Alan / Dal", "{{alan}} / {{dal}}"),
                 ("Doğum Tarihi", "{{ogrenci_dogum}}"), ("Adres / Telefon", "{{ogrenci_adres}} / {{ogrenci_tel}}"),
                 ("Veli Adı Soyadı / Tel", "{{veli_ad}} / {{veli_tel}}"), ("IBAN", "{{ogrenci_iban}}")])
    doc.add_paragraph().add_run("B) İŞLETME BİLGİLERİ").bold = True
    _tablo(doc, [("İşletme Adı", "{{isletme_adi}}"), ("Adres", "{{isletme_adres}}"),
                 ("Vergi No / SGK Sicil No", "{{isletme_vergi_no}} / {{isletme_sgk_sicil}}"),
                 ("Yetkili / Telefon", "{{isletme_yetkili}} / {{isletme_tel}}"),
                 ("Usta Öğretici / Eğitici Personel", "{{usta_ogretici}}"),
                 ("Çalışan Personel Sayısı", "{{personel_sayisi}}")])
    doc.add_paragraph().add_run("C) SÖZLEŞME HÜKÜMLERİ").bold = True
    _maddeler(doc, [
        ("Konu", "Bu sözleşme, 3308 sayılı Mesleki Eğitim Kanunu ve Millî Eğitim Bakanlığı Ortaöğretim Kurumları "
                 "Yönetmeliği hükümleri çerçevesinde yukarıda bilgileri yazılı öğrencinin işletmede mesleki eğitim "
                 "görmesine ilişkin tarafların hak ve yükümlülüklerini düzenler."),
        ("Süre ve Eğitim Günleri", "Eğitim {{baslangic}} tarihinde başlar, {{bitis}} tarihinde sona erer. "
                                   "Öğrenci haftada {{egitim_gunleri}} günleri işletmede eğitim görür."),
        ("Ücret", "İşletme, öğrenciye aylık {{ucret}} TL'den (asgari ücretin net tutarı esas alınarak yasal "
                  "alt sınır: {{en_az_ucret}} TL) az olmamak üzere ücret öder. Ücret, izleyen ayın "
                  "{{dekont_son_gun}}. gününe kadar öğrencinin banka hesabına yatırılır ve dekontu okula teslim edilir."),
        ("Sigorta", "Öğrencinin 5510 sayılı Kanunun 5 inci maddesinin (b) bendi kapsamındaki sigorta primleri "
                    "3308 sayılı Kanunun 25 inci maddesi gereğince Millî Eğitim Bakanlığınca ödenir. İş kazası "
                    "hâlinde işletme durumu derhâl okul müdürlüğüne bildirir."),
        ("İşletmenin Yükümlülükleri", "Eğitimi usta öğretici/eğitici personel gözetiminde, çerçeve öğretim "
                                      "programına uygun yaptırmak; iş sağlığı ve güvenliği önlemlerini almak ve "
                                      "öğrenciye gerekli eğitimi vermek; öğrenciyi eğitim dışı işlerde çalıştırmamak; "
                                      "devamsızlıkları ve disiplin durumlarını okula bildirmek; değerlendirme "
                                      "formlarını zamanında doldurmak."),
        ("Öğrencinin Yükümlülükleri", "İşletmenin çalışma düzenine, iş sağlığı ve güvenliği kurallarına uymak; "
                                      "eğitime düzenli devam etmek; işletmeye ait araç-gereci özenle kullanmak."),
        ("Okulun Yükümlülükleri", "Koordinatör öğretmen görevlendirerek eğitimi izlemek, rehberlik etmek, "
                                  "öğrencinin değerlendirmesini mevzuata göre yapmak ve sigorta işlemlerini yürütmek."),
        ("Fesih", "Sözleşme; mevzuatta belirtilen hâllerde ve tarafların yazılı bildirimiyle, okul müdürlüğünün "
                  "bilgisi dâhilinde feshedilebilir."),
        ("Diğer Hükümler", "Bu sözleşmede hüküm bulunmayan hâllerde ilgili mevzuat hükümleri uygulanır. "
                           "Sözleşme üç nüsha olarak düzenlenmiş ve taraflarca imzalanmıştır."),
    ])
    _imza_tablosu(doc, [("Okul Müdürü", "{{mudur}}"), ("İşletme Yetkilisi", "{{isletme_yetkili}}"),
                        ("Öğrenci", "{{ogrenci_ad_soyad}}"), ("Veli", "{{veli_ad}}")])
    doc.save(str(yol))


def _sablon_ciraklik_sozlesmesi(yol: Path) -> None:
    doc = _yeni_belge()
    _kurum_basligi(doc)
    _ortali(doc, "ÇIRAKLIK SÖZLEŞMESİ (Mesleki Eğitim Merkezi Programı)", boyut=13)
    doc.add_paragraph("Sözleşme No: {{sozlesme_no}}        Sözleşme Tarihi: {{sozlesme_tarihi}}")
    _tablo(doc, [("Aday Çırak/Çırak Adı Soyadı", "{{ogrenci_ad_soyad}}"), ("T.C. Kimlik No", "{{ogrenci_tc}}"),
                 ("Doğum Tarihi", "{{ogrenci_dogum}}"), ("Alan / Dal", "{{alan}} / {{dal}}"),
                 ("Sınıf / Statü", "{{sinif}} / {{statu}}"), ("Veli / Vasi", "{{veli_ad}} ({{veli_tel}})"),
                 ("İşletme", "{{isletme_adi}}"), ("İşletme Adresi", "{{isletme_adres}}"),
                 ("İşveren / Yetkili", "{{isletme_yetkili}}"), ("Usta Öğretici", "{{usta_ogretici}}"),
                 ("Başlangıç - Bitiş", "{{baslangic}} - {{bitis}}"), ("Aylık Ücret", "{{ucret}} TL")])
    _maddeler(doc, [
        ("Dayanak", "Bu sözleşme 3308 sayılı Mesleki Eğitim Kanunu ve Mesleki eğitim merkezi programına ilişkin "
                    "mevzuat hükümlerine göre yazılı olarak düzenlenmiştir."),
        ("Eğitim", "Aday çırak/çırak haftada {{okul_gunleri}} günleri mesleki eğitim merkezinde teorik eğitim, "
                   "{{egitim_gunleri}} günleri işletmede pratik eğitim görür."),
        ("Ücret", "İşveren, aday çırak/çırağa yasal alt sınırın ({{en_az_ucret}} TL) altında olmamak üzere aylık "
                  "{{ucret}} TL ücret öder."),
        ("Sigorta", "Sigorta primleri 3308 sayılı Kanunun 25 inci maddesi uyarınca Millî Eğitim Bakanlığınca ödenir."),
        ("Yükümlülükler", "İşveren eğitimi usta öğretici gözetiminde yaptırır, iş sağlığı ve güvenliği önlemlerini "
                          "alır; çırak işyeri kurallarına ve eğitim programına uyar."),
        ("Fesih ve Diğer Hükümler", "Sözleşmenin feshi ve hüküm bulunmayan hâllerde 3308 sayılı Kanun ve ilgili "
                                    "mevzuat hükümleri uygulanır."),
    ])
    _imza_tablosu(doc, [("Merkez Müdürü", "{{mudur}}"), ("İşveren", "{{isletme_yetkili}}"),
                        ("Aday Çırak/Çırak", "{{ogrenci_ad_soyad}}"), ("Veli/Vasi", "{{veli_ad}}")])
    doc.save(str(yol))


def _sablon_gorevlendirme(yol: Path) -> None:
    doc = _yeni_belge()
    _kurum_basligi(doc)
    doc.add_paragraph()
    doc.add_paragraph("Sayı  : {{sayi}}\t\t\t\t\t\t\t{{bugun}}")
    doc.add_paragraph("Konu : İşletmelerde Mesleki Eğitim Koordinatörlük Görevi")
    doc.add_paragraph()
    _ortali(doc, "Sayın {{ogretmen_ad_soyad}}", kalin=False)
    _ortali(doc, "{{ogretmen_brans}} Öğretmeni", kalin=False)
    doc.add_paragraph()
    doc.add_paragraph(
        "Okulumuzda {{egitim_yili}} eğitim ve öğretim yılında işletmelerde mesleki eğitim gören öğrencilerin "
        "işyerindeki uygulamalı eğitimini izlemek, programa uygunluğunu ve sistemin işlerliğini sağlamak ve "
        "meslekî rehberlikte bulunmak üzere; Millî Eğitim Bakanlığı Ortaöğretim Kurumları Yönetmeliği ve Millî "
        "Eğitim Bakanlığı Yönetici ve Öğretmenlerinin Ders ve Ek Ders Saatlerine İlişkin Karar'ın 15 inci maddesi "
        "gereğince aşağıda belirtilen işletmelerde koordinatör öğretmen olarak görevlendirildiniz.")
    doc.add_paragraph("{{gorev_listesi}}")
    doc.add_paragraph("Haftalık toplam işletmelerde meslek eğitimi ek ders görevi: {{haftalik_toplam}} saat.")
    doc.add_paragraph("Ziyaretlerinize ait tespit ve önerileri ziyaret formuyla koordinatör müdür yardımcılığına "
                      "teslim etmeniz hususunda bilgilerinizi ve gereğini rica ederim.")
    doc.add_paragraph()
    p = doc.add_paragraph("{{mudur}}\nOkul Müdürü")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph("\nTEBELLÜĞ EDEN\nAdı Soyadı: {{ogretmen_ad_soyad}}\nTarih: ...../...../........\nİmza:")
    doc.save(str(yol))


def _sablon_ziyaret_formu(yol: Path) -> None:
    doc = _yeni_belge()
    _kurum_basligi(doc)
    _ortali(doc, "İŞLETME ZİYARET VE REHBERLİK FORMU", boyut=13)
    _tablo(doc, [("İşletme Adı", "{{isletme_adi}}"), ("Adres / Telefon", "{{isletme_adres}} / {{isletme_tel}}"),
                 ("Usta Öğretici", "{{usta_ogretici}}"), ("Ziyaret Tarihi", "{{ziyaret_tarihi}}"),
                 ("Koordinatör Öğretmen", "{{ogretmen_ad_soyad}}"), ("Öğrenciler", "{{ogrenci_listesi}}")])
    doc.add_paragraph()
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = "Kontrol Edilen Husus", "Evet", "Hayır"
    for h in ["Öğrenciler işletmede devam ediyor", "Eğitim çerçeve öğretim programına uygun yürütülüyor",
              "Usta öğretici/eğitici personel görevinin başında", "İş sağlığı ve güvenliği önlemleri alınmış",
              "Kişisel koruyucu donanım kullanılıyor", "Öğrenci eğitim dışı işlerde çalıştırılmıyor",
              "Ücret ödemeleri zamanında yapılıyor", "Devam-devamsızlık çizelgesi tutuluyor",
              "Öğrencilerin iş dosyaları kontrol edildi (md.89/1-ç)",
              "Devamsızlıklar iki iş günü içinde okula bildiriliyor (md.144/1-f)",
              "Öğrenciler günde 8 saatten fazla çalıştırılmıyor (md.144/2)",
              "Eğitim birimi / donanım eğitim için yeterli (md.142, md.148)",
              "Uyum sorunu veya özel eğitim ihtiyacı olan öğrenci yok (md.89/1-e)"]:
        c = t.add_row().cells
        c[0].text = h
    doc.add_paragraph()
    doc.add_paragraph().add_run("Tespitler:").bold = True
    doc.add_paragraph("{{tespitler}}")
    doc.add_paragraph().add_run("Öneriler / Alınan Tedbirler:").bold = True
    doc.add_paragraph("{{oneriler}}")
    _imza_tablosu(doc, [("Koordinatör Öğretmen", "{{ogretmen_ad_soyad}}"), ("İşletme Yetkilisi", "{{isletme_yetkili}}"),
                        ("Koordinatör Müdür Yrd.", "{{koordinator_mudur_yrd}}")])
    doc.save(str(yol))


def _sablon_devamsizlik(yol: Path) -> None:
    doc = _yeni_belge()
    _kurum_basligi(doc)
    doc.add_paragraph()
    doc.add_paragraph("Sayı  : {{sayi}}\t\t\t\t\t\t\t{{bugun}}")
    doc.add_paragraph("Konu : Öğrenci Devamsızlığı")
    doc.add_paragraph()
    _ortali(doc, "Sayın {{veli_ad}}", kalin=False)
    doc.add_paragraph()
    doc.add_paragraph(
        "Velisi bulunduğunuz okulumuz {{sinif}} sınıfı {{ogrenci_no}} numaralı öğrencisi {{ogrenci_ad_soyad}}'in "
        "{{bugun}} tarihi itibarıyla devamsızlık durumu aşağıda belirtilmiştir.")
    _tablo(doc, [("Okulda özürsüz", "{{okul_ozursuz}} gün"), ("Okulda özürlü", "{{okul_ozurlu}} gün"),
                 ("İşletmede özürsüz", "{{isletme_ozursuz}} gün"), ("İşletmede özürlü", "{{isletme_ozurlu}} gün"),
                 ("Toplam", "{{devamsizlik_toplam}} gün")])
    doc.add_paragraph()
    doc.add_paragraph(
        "Ortaöğretim Kurumları Yönetmeliği gereğince özürsüz devamsızlığı {{sinir_ozursuz}} günü veya toplam "
        "devamsızlığı {{sinir_toplam}} günü aşan öğrenci başarısız sayılmaktadır. Öğrencinizin devamını sağlamanız "
        "hususunda bilgilerinizi rica ederim.")
    p = doc.add_paragraph("\n{{mudur}}\nOkul Müdürü")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.save(str(yol))


def _sablon_beceri(yol: Path) -> None:
    doc = _yeni_belge()
    _kurum_basligi(doc)
    _ortali(doc, "İŞLETMEDE MESLEKİ EĞİTİM DEĞERLENDİRME FORMU", boyut=13)
    doc.add_paragraph("{{egitim_yili}} Eğitim-Öğretim Yılı     Dönem: ......")
    _tablo(doc, [("Öğrenci", "{{ogrenci_ad_soyad}} ({{ogrenci_no}})"), ("Sınıf / Alan / Dal", "{{sinif}} / {{alan}} / {{dal}}"),
                 ("İşletme", "{{isletme_adi}}"), ("Usta Öğretici", "{{usta_ogretici}}"),
                 ("Koordinatör Öğretmen", "{{koordinator}}")])
    doc.add_paragraph()
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = "Değerlendirme Ölçütü", "Puan Aralığı", "Puan"
    for h, a in [("İş sağlığı ve güvenliği kurallarına uyma", "0-15"), ("Araç-gereç ve ekipmanı doğru kullanma", "0-15"),
                 ("Mesleki bilgi ve beceriyi uygulama", "0-25"), ("İşi zamanında ve kurallara uygun tamamlama", "0-15"),
                 ("Çalışma düzenine ve disipline uyma", "0-10"), ("İletişim ve ekip çalışması", "0-10"),
                 ("Mesleğe ilgi ve sorumluluk", "0-10")]:
        c = t.add_row().cells
        c[0].text, c[1].text = h, a
    c = t.add_row().cells
    c[0].text, c[1].text = "TOPLAM", "0-100"
    _imza_tablosu(doc, [("Usta Öğretici", "{{usta_ogretici}}"), ("Koordinatör Öğretmen", "{{koordinator}}"),
                        ("Okul Müdürü", "{{mudur}}")])
    doc.save(str(yol))


def _sablon_ucret_bildirimi(yol: Path) -> None:
    doc = _yeni_belge()
    _kurum_basligi(doc)
    doc.add_paragraph()
    doc.add_paragraph("Sayı  : {{sayi}}\t\t\t\t\t\t\t{{bugun}}")
    doc.add_paragraph("Konu : {{donem}} Dönemi Öğrenci Ücretleri")
    doc.add_paragraph()
    _ortali(doc, "{{isletme_adi}}", kalin=False)
    _ortali(doc, "Yetkilisine", kalin=False)
    doc.add_paragraph()
    doc.add_paragraph(
        "İşletmenizde mesleki eğitim gören öğrencilerimize {{donem}} dönemi için 3308 sayılı Mesleki Eğitim "
        "Kanunu gereğince ödenmesi gereken asgari ücret tutarları aşağıda gösterilmiştir. Ücretlerin izleyen ayın "
        "{{dekont_son_gun}}. gününe kadar öğrencilerin banka hesaplarına yatırılarak dekontlarının okulumuza "
        "teslim edilmesini rica ederim.")
    doc.add_paragraph("{{ogrenci_listesi}}")
    doc.add_paragraph("Toplam ödenecek ücret: {{toplam_ucret}} TL — Tahmini devlet katkısı: {{toplam_dk}} TL")
    p = doc.add_paragraph("\n{{mudur}}\nOkul Müdürü")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.save(str(yol))


def _sablon_koordinator_bildirimi(yol: Path) -> None:
    doc = _yeni_belge()
    _kurum_basligi(doc)
    doc.add_paragraph()
    doc.add_paragraph("Sayı  : {{sayi}}\t\t\t\t\t\t\t{{bugun}}")
    doc.add_paragraph("Konu : Koordinatör Öğretmen Bildirimi")
    doc.add_paragraph()
    _ortali(doc, "{{isletme_adi}}", kalin=False)
    _ortali(doc, "Yetkilisine", kalin=False)
    doc.add_paragraph()
    doc.add_paragraph(
        "İşletmenizde mesleki eğitim gören öğrencilerimizin işyerindeki eğitimini izlemek, rehberlik etmek ve "
        "eğitimin programa uygunluğunu sağlamak üzere Millî Eğitim Bakanlığı Ortaöğretim Kurumları Yönetmeliğinin "
        "88 inci maddesi gereğince {{egitim_yili}} eğitim ve öğretim yılında aşağıda adı yazılı öğretmen(ler) "
        "koordinatör öğretmen olarak görevlendirilmiştir.")
    doc.add_paragraph().add_run("Koordinatör öğretmen(ler):").bold = True
    doc.add_paragraph("{{koordinator_listesi}}")
    doc.add_paragraph().add_run("İşletmenizdeki öğrencilerimiz:").bold = True
    doc.add_paragraph("{{ogrenci_listesi}}")
    doc.add_paragraph(
        "Öğrencilerin devamsızlıklarının iki iş günü içinde okulumuza bildirilmesi, puan çizelgelerinin dönem "
        "sonundan en az beş gün önce teslim edilmesi ve öğrencilerin günde sekiz saatten fazla çalıştırılmaması "
        "hususlarında bilgilerinizi ve gereğini rica ederim.")
    p = doc.add_paragraph("\n{{mudur}}\nOkul Müdürü")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.save(str(yol))


UYGUNLUK_OLCUTLERI = [
    ("Üretim / hizmet, okuldaki alan ve dalların çerçeve öğretim programına uygun", "{{kriter_program}}"),
    ("Makine, araç-gereç ve donanım, programın en az %80'ini uygulamaya yeterli", ""),
    ("{{egitim_birimi_esigi}} ve üzeri öğrenci için eğitim birimi kurulmuş", "{{kriter_egitim_birimi}}"),
    ("Usta öğretici / eğitici personel bulunuyor (ustalık belgesi veya iş pedagojisi)", "{{kriter_usta}}"),
    ("Öğrencilerin giyinme, yıkanma ve temizlenme yerleri uygun", ""),
    ("İl dışında ise öğrencinin konaklama ve yemek imkânı sağlanıyor", "{{kriter_il_disi}}"),
    ("Önceki yıllarda mesleki eğitim için uygun bulunmuş", "{{kriter_onceki}}"),
    ("İş sağlığı ve güvenliği önlemleri alınmış", ""),
]


def _sablon_isletme_uygunluk(yol: Path) -> None:
    doc = _yeni_belge()
    _kurum_basligi(doc)
    _ortali(doc, "İŞLETME UYGUNLUK TESPİT TUTANAĞI", boyut=13)
    _ortali(doc, "(Ortaöğretim Kurumları Yönetmeliği md.139-142)", kalin=False, boyut=10)
    _tablo(doc, [("İşletme Adı", "{{isletme_adi}}"), ("Adres / Telefon", "{{isletme_adres}} / {{isletme_tel}}"),
                 ("Yetkili", "{{isletme_yetkili}}"), ("Alan", "{{isletme_alan}}"),
                 ("Personel Sayısı", "{{personel_sayisi}}"),
                 ("Usta Öğretici / Eğitici Personel", "{{usta_ogretici}} ({{usta_ogretici_sayisi}} kişi)"),
                 ("Mevcut Öğrenci / Kontenjan", "{{ogrenci_sayisi}} / {{kontenjan}}")])
    doc.add_paragraph()
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for i, b in enumerate(["Ölçüt", "Kayıtlı Bilgi", "Uygun", "Uygun Değil"]):
        t.rows[0].cells[i].text = b
        t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for olcut, bilgi in UYGUNLUK_OLCUTLERI:
        c = t.add_row().cells
        c[0].text, c[1].text = olcut, bilgi
    doc.add_paragraph()
    doc.add_paragraph("Komisyonun görüşü: İşletme, mesleki eğitim yaptırmaya  [  ] UYGUNDUR   [  ] UYGUN DEĞİLDİR.")
    doc.add_paragraph("Uygun bulunan öğrenci kontenjanı: ............")
    doc.add_paragraph("Tespit Tarihi: {{uygunluk_tarihi}}")
    _imza_tablosu(doc, [("Komisyon Başkanı", "{{koordinator_mudur_yrd}}"), ("Üye (Alan/Bölüm Şefi)", ""),
                        ("Üye (Alan Öğretmeni)", ""), ("İşletme Yetkilisi", "{{isletme_yetkili}}")])
    doc.save(str(yol))


def _sablon_komisyon_karari(yol: Path) -> None:
    doc = _yeni_belge()
    _kurum_basligi(doc)
    _ortali(doc, "KOORDİNATÖRLÜK GÖREVİ PLANLAMA KOMİSYONU KARARI", boyut=13)
    doc.add_paragraph("Karar Tarihi: {{bugun}}        Eğitim-Öğretim Yılı: {{egitim_yili}}")
    doc.add_paragraph(
        "Ortaöğretim Kurumları Yönetmeliğinin 88 inci maddesi ile Millî Eğitim Bakanlığı Yönetici ve Öğretmenlerinin "
        "Ders ve Ek Ders Saatlerine İlişkin Kararın 15 inci maddesi gereğince; işletmelerin okula uzaklığı, ulaşım "
        "imkânları, işletme ve öğrenci sayıları dikkate alınarak koordinatörlük görevleri aşağıdaki şekilde "
        "planlanmıştır.")
    doc.add_paragraph().add_run("Planlama ölçütleri:").bold = True
    doc.add_paragraph("{{saat_olcutleri}}")
    doc.add_paragraph().add_run("Görev dağılımı (öğretmen - gün - işletme - öğrenci - haftalık saat):").bold = True
    doc.add_paragraph("{{plan_listesi}}")
    doc.add_paragraph("Toplam: {{isletme_sayisi}} işletme, {{ogrenci_sayisi}} öğrenci, {{ogretmen_sayisi}} "
                      "koordinatör öğretmen, haftalık {{toplam_saat}} saat.")
    doc.add_paragraph().add_run("Koordinatör görevlendirilmeyen işletmeler:").bold = True
    doc.add_paragraph("{{kapsam_disi}}")
    _imza_tablosu(doc, [("Komisyon Başkanı", "{{koordinator_mudur_yrd}}"), ("Üye", ""), ("Üye", ""), ("Üye", "")])
    doc.add_paragraph()
    p = doc.add_paragraph("UYGUNDUR\n{{bugun}}\n{{mudur}}\nOkul Müdürü")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(str(yol))


def _sablon_secim_tutanagi(yol: Path) -> None:
    doc = _yeni_belge()
    _kurum_basligi(doc)
    _ortali(doc, "İŞLETMELERDE MESLEKİ EĞİTİM ÖĞRENCİ SEÇİM KOMİSYONU TUTANAĞI", boyut=13)
    doc.add_paragraph("Tarih: {{bugun}}        Alan: {{alan}}        Kontenjan: {{kontenjan}}")
    doc.add_paragraph(
        "Ortaöğretim Kurumları Yönetmeliğinin 136 ncı maddesi gereğince işletmelere gönderilecek öğrenciler; alt "
        "sınıf yılsonu başarı ortalaması, kısa süreli uzaklaştırma cezasından daha ağır disiplin cezası almamış "
        "olma ve genel durumları dikkate alınarak sıralanmıştır. Şehit ve gazi çocukları ile koruma/eğitim tedbiri "
        "kararı bulunan öğrenciler sıralamaya alınmadan yerleştirilmiştir.")
    doc.add_paragraph().add_run("Sıralamaya girmeden yerleştirilen öğrenciler:").bold = True
    doc.add_paragraph("{{oncelikli_listesi}}")
    doc.add_paragraph().add_run("Başarı sırası (gönderilecek / yedek):").bold = True
    doc.add_paragraph("{{sira_listesi}}")
    doc.add_paragraph().add_run("Koşulu sağlamayan öğrenciler:").bold = True
    doc.add_paragraph("{{kosulsuz_listesi}}")
    _imza_tablosu(doc, [("Komisyon Başkanı", "{{koordinator_mudur_yrd}}"), ("Üye", ""), ("Üye", ""), ("Üye", "")])
    doc.save(str(yol))


# ------------------------------------------------------------------ bağlamlar

def kurum_baglam(db) -> dict:
    from ..marka import kurum_logosu
    k = db.kurum()
    return {"_logo": kurum_logosu(db), "kurum_adi": k["ad"], "kurum_il": k["il"], "kurum_ilce": (k["ilce"] or k["il"] or "").upper(),
            "kurum_kodu": k["kurum_kodu"], "mudur": k["mudur"], "koordinator_mudur_yrd": k["koordinator_mudur_yrd"],
            "egitim_yili": k["egitim_yili"], "kurum_adres": k["adres"], "kurum_tel": k["telefon"],
            "bugun": tr(date.today()), "sayi": "", "dekont_son_gun": db.ayar("dekont_teslim_son_gun"),
            "sinir_ozursuz": f"{db.ayar('devamsizlik_ozursuz_sinir'):g}",
            "sinir_toplam": f"{db.ayar('devamsizlik_toplam_sinir'):g}"}


def _gun_adlari(s: str | None) -> str:
    return ", ".join(GUNLER[g - 1] for g in gunler_listesi(s) if g <= 5)


def yerlestirme_baglam(db, yid: int) -> dict:
    r = db.q1("""SELECT y.*, o.tc, o.okul_no, o.ad, o.soyad, o.dogum_tarihi, o.adres AS ogr_adres,
                        o.telefon AS ogr_tel, o.veli_ad, o.veli_tel, o.iban AS ogr_iban, o.statu,
                        s.ad AS sinif, s.program, s.okul_gunleri, s.isletme_gunleri, a.ad AS alan, dl.ad AS dal,
                        i.ad AS isletme, i.adres AS isl_adres, i.yetkili, i.telefon AS isl_tel, i.vergi_no,
                        i.sgk_sicil, i.usta_ogretici, i.personel_sayisi, i.kamu_kurumu
                 FROM yerlestirmeler y JOIN ogrenciler o ON o.id = y.ogrenci_id
                 LEFT JOIN siniflar s ON s.id = o.sinif_id LEFT JOIN alanlar a ON a.id = s.alan_id
                 LEFT JOIN dallar dl ON dl.id = s.dal_id LEFT JOIN isletmeler i ON i.id = y.isletme_id
                 WHERE y.id = ?""", (yid,))
    if not r:
        raise ValueError("Yerleştirme kaydı bulunamadı")
    mem = ucret_srv.mem_mi(r["program"], r["statu"])
    oran, _ = ucret_srv.en_az_ucret_orani(db, mem, r["statu"], int(r["personel_sayisi"] or 0))
    en_az = db.ayar("asgari_ucret_net") * oran
    koord = db.q1("""SELECT o.ad || ' ' || o.soyad AS ad FROM koordinator_gorevleri kg
                     JOIN ogretmenler o ON o.id = kg.ogretmen_id WHERE kg.isletme_id = ? LIMIT 1""", (r["isletme_id"],))
    return {**kurum_baglam(db),
            "ogrenci_ad_soyad": f"{r['ad']} {r['soyad']}", "ogrenci_tc": r["tc"], "ogrenci_no": r["okul_no"],
            "ogrenci_dogum": tr(r["dogum_tarihi"]), "ogrenci_adres": r["ogr_adres"], "ogrenci_tel": r["ogr_tel"],
            "ogrenci_iban": r["ogr_iban"], "veli_ad": r["veli_ad"], "veli_tel": r["veli_tel"], "statu": r["statu"],
            "sinif": r["sinif"], "alan": r["alan"], "dal": r["dal"], "isletme_adi": r["isletme"],
            "isletme_adres": r["isl_adres"], "isletme_yetkili": r["yetkili"], "isletme_tel": r["isl_tel"],
            "isletme_vergi_no": r["vergi_no"], "isletme_sgk_sicil": r["sgk_sicil"],
            "usta_ogretici": r["usta_ogretici"], "personel_sayisi": r["personel_sayisi"],
            "baslangic": tr(r["baslangic"]), "bitis": tr(r["bitis"]), "sozlesme_no": r["sozlesme_no"],
            "sozlesme_tarihi": tr(r["sozlesme_tarihi"]), "tur": r["tur"],
            "ucret": para(max(float(r["ucret"] or 0), en_az)), "en_az_ucret": para(en_az),
            "egitim_gunleri": _gun_adlari(r["isletme_gunleri"]), "okul_gunleri": _gun_adlari(r["okul_gunleri"]),
            "koordinator": koord["ad"] if koord else ""}


def ogretmen_gorev_baglam(db, tid: int) -> dict:
    o = db.q1("SELECT * FROM ogretmenler WHERE id = ?", (tid,))
    rows = db.q("""SELECT kg.gun, kg.haftalik_saat, i.ad, i.adres FROM koordinator_gorevleri kg
                   JOIN isletmeler i ON i.id = kg.isletme_id WHERE kg.ogretmen_id = ? ORDER BY kg.gun, i.ad""", (tid,))
    satirlar = [f"{n}. {GUNLER[r['gun'] - 1]}: {r['ad']} ({r['haftalik_saat']} saat)" for n, r in enumerate(rows, 1)]
    return {**kurum_baglam(db), "ogretmen_ad_soyad": f"{o['ad']} {o['soyad']}", "ogretmen_brans": o["brans"] or "",
            "gorev_listesi": "\n".join(satirlar) or "-", "haftalik_toplam": sum(r["haftalik_saat"] for r in rows)}


def ziyaret_baglam(db, zid: int) -> dict:
    z = db.q1("""SELECT z.*, o.ad AS oad, o.soyad AS osoyad, i.ad AS isletme, i.adres, i.telefon, i.yetkili,
                        i.usta_ogretici FROM ziyaretler z LEFT JOIN ogretmenler o ON o.id = z.ogretmen_id
                 LEFT JOIN isletmeler i ON i.id = z.isletme_id WHERE z.id = ?""", (zid,))
    ogr = db.q("""SELECT o.ad, o.soyad, s.ad AS sinif FROM yerlestirmeler y JOIN ogrenciler o ON o.id = y.ogrenci_id
                  LEFT JOIN siniflar s ON s.id = o.sinif_id WHERE y.isletme_id = ? AND y.durum = 'Aktif'
                  AND (y.baslangic IS NULL OR y.baslangic <= ?) AND (y.bitis IS NULL OR y.bitis = '' OR y.bitis >= ?)""",
               (z["isletme_id"], z["tarih"], z["tarih"]))
    return {**kurum_baglam(db), "isletme_adi": z["isletme"], "isletme_adres": z["adres"], "isletme_tel": z["telefon"],
            "isletme_yetkili": z["yetkili"], "usta_ogretici": z["usta_ogretici"], "ziyaret_tarihi": tr(z["tarih"]),
            "ogretmen_ad_soyad": f"{z['oad'] or ''} {z['osoyad'] or ''}".strip(),
            "ogrenci_listesi": "\n".join(f"{r['ad']} {r['soyad']} ({r['sinif'] or ''})" for r in ogr) or "-",
            "tespitler": z["tespitler"] or "\n\n", "oneriler": z["oneriler"] or "\n\n"}


def devamsizlik_baglam(db, ogr_id: int) -> dict:
    from .devam import ozet
    o = db.q1("""SELECT o.*, s.ad AS sinif FROM ogrenciler o LEFT JOIN siniflar s ON s.id = o.sinif_id
                 WHERE o.id = ?""", (ogr_id,))
    x = next((d for d in ozet(db, sadece_devamsizlar=False) if d.ogrenci_id == ogr_id), None)
    return {**kurum_baglam(db), "ogrenci_ad_soyad": f"{o['ad']} {o['soyad']}", "ogrenci_no": o["okul_no"],
            "sinif": o["sinif"], "veli_ad": o["veli_ad"] or "Veli",
            "okul_ozursuz": f"{x.okul_ozursuz:g}" if x else "0", "okul_ozurlu": f"{x.okul_ozurlu:g}" if x else "0",
            "isletme_ozursuz": f"{x.isletme_ozursuz:g}" if x else "0",
            "isletme_ozurlu": f"{x.isletme_ozurlu:g}" if x else "0",
            "devamsizlik_toplam": f"{x.toplam:g}" if x else "0"}


def isletme_ucret_baglam(db, isletme_id: int, yil: int, ay: int) -> dict:
    i = db.q1("SELECT * FROM isletmeler WHERE id = ?", (isletme_id,))
    satirlar = [s for s in ucret_srv.aylik_hesap(db, yil, ay) if s.isletme == i["ad"]]
    liste = "\n".join(f"{n}. {s.ogrenci} ({s.sinif}) - {s.gun} gün - {para(s.odenecek_ucret)} TL"
                      for n, s in enumerate(satirlar, 1))
    return {**kurum_baglam(db), "isletme_adi": i["ad"], "donem": f"{AYLAR[ay - 1]} {yil}",
            "ogrenci_listesi": liste or "-", "toplam_ucret": para(sum(s.odenecek_ucret for s in satirlar)),
            "toplam_dk": para(sum(s.devlet_katkisi for s in satirlar))}


def isletme_baglam(db, isletme_id: int) -> dict:
    i = db.q1("""SELECT i.*, a.ad AS alan FROM isletmeler i LEFT JOIN alanlar a ON a.id = i.alan_id
                 WHERE i.id = ?""", (isletme_id,))
    if not i:
        raise ValueError("İşletme bulunamadı")
    ogr = db.q("""SELECT o.ad, o.soyad, o.okul_no, s.ad AS sinif FROM yerlestirmeler y
                  JOIN ogrenciler o ON o.id = y.ogrenci_id LEFT JOIN siniflar s ON s.id = o.sinif_id
                  WHERE y.isletme_id = ? AND y.durum = 'Aktif' ORDER BY s.ad, o.okul_no""", (isletme_id,))
    koord = db.q("""SELECT o.ad, o.soyad, o.brans, o.telefon, GROUP_CONCAT(DISTINCT kg.gun) AS gunler
                    FROM koordinator_gorevleri kg JOIN ogretmenler o ON o.id = kg.ogretmen_id
                    WHERE kg.isletme_id = ? GROUP BY o.id ORDER BY o.ad""", (isletme_id,))
    n = len(ogr)
    esik = int(db.ayar("egitim_birimi_esigi"))
    gecmis = db.q1("""SELECT COUNT(*) c FROM yerlestirmeler WHERE isletme_id = ? AND durum != 'Aktif'""",
                   (isletme_id,))["c"]
    return {**kurum_baglam(db), "isletme_adi": i["ad"], "isletme_adres": i["adres"], "isletme_tel": i["telefon"],
            "isletme_yetkili": i["yetkili"], "isletme_alan": i["alan"] or "", "personel_sayisi": i["personel_sayisi"],
            "usta_ogretici": i["usta_ogretici"], "usta_ogretici_sayisi": i["usta_ogretici_sayisi"] or 1,
            "kontenjan": i["kontenjan"] or "-", "ogrenci_sayisi": n, "egitim_birimi_esigi": esik,
            "uygunluk_tarihi": tr(i["uygunluk_tarihi"]) or "...../...../........",
            "ogrenci_listesi": "\n".join(f"{k}. {r['ad']} {r['soyad']} ({r['sinif'] or ''} / {r['okul_no'] or ''})"
                                         for k, r in enumerate(ogr, 1)) or "-",
            "koordinator_listesi": "\n".join(
                f"{r['ad']} {r['soyad']} ({r['brans'] or ''}) - ziyaret günü: {_gun_adlari(r['gunler'])}"
                + (f" - Tel: {r['telefon']}" if r["telefon"] else "") for r in koord) or "-",
            "kriter_program": f"Alan: {i['alan']}" if i["alan"] else "",
            "kriter_egitim_birimi": ("Var" if i["egitim_birimi"] else "Yok") + f" ({n} öğrenci)"
                                    if n >= esik else f"Gerekmiyor ({n} öğrenci)",
            "kriter_usta": (f"{i['usta_ogretici']} - " if i["usta_ogretici"] else "Tanımlı değil - ")
                           + ("belgesi var" if i["usta_belgesi"] else "belgesi işaretlenmemiş"),
            "kriter_il_disi": "İl dışında" if i["il_disi"] else "İl içinde (gerekmiyor)",
            "kriter_onceki": f"Önceki yerleştirme: {gecmis}" if gecmis else ""}


def komisyon_baglam(db) -> dict:
    from . import koordinator
    kapsam: list[tuple[str, str]] = []
    koordinator.isletme_ihtiyaclari(db, kapsam_disi=kapsam)
    rows = db.q("""SELECT kg.*, o.ad, o.soyad, i.ad AS isletme, i.guzergah, i.mesafe_km
                   FROM koordinator_gorevleri kg JOIN ogretmenler o ON o.id = kg.ogretmen_id
                   JOIN isletmeler i ON i.id = kg.isletme_id ORDER BY o.ad, o.soyad, kg.gun, i.ad""")
    ogr_say = {r["isletme_id"]: r["n"] for r in db.q(
        """SELECT isletme_id, COUNT(*) n FROM yerlestirmeler WHERE durum = 'Aktif' AND isletme_id IS NOT NULL
           GROUP BY isletme_id""")}
    satirlar = [f"{k}. {r['ad']} {r['soyad']} - {GUNLER[r['gun'] - 1] if 1 <= r['gun'] <= 5 else r['gun']} - "
                f"{r['isletme']} ({r['guzergah'] or '-'}, {float(r['mesafe_km'] or 0):g} km) - "
                f"{ogr_say.get(r['isletme_id'], 0)} öğrenci - {r['haftalik_saat']} saat"
                for k, r in enumerate(rows, 1)]
    isletmeler = {r["isletme_id"] for r in rows}
    return {**kurum_baglam(db), "saat_olcutleri": koordinator.saat_hesabi_aciklama(db),
            "plan_listesi": "\n".join(satirlar) or "Henüz görevlendirme yapılmadı.",
            "isletme_sayisi": len(isletmeler),
            "ogrenci_sayisi": sum(ogr_say.get(i, 0) for i in isletmeler),
            "ogretmen_sayisi": len({r["ogretmen_id"] for r in rows}),
            "toplam_saat": sum(r["haftalik_saat"] for r in rows),
            "kapsam_disi": "\n".join(f"{ad}: {neden}" for ad, neden in kapsam) or "-"}


# ------------------------------------------------------------------ kayıt defteri

@dataclass
class BelgeTuru:
    anahtar: str
    etiket: str
    kaynak: str  # yerlestirme | ogretmen | ziyaret | ogrenci | isletme_ay | isletme | isletme_koord | kurum | secim
    olusturucu: Callable[[Path], None]


BELGE_TURLERI: dict[str, BelgeTuru] = {b.anahtar: b for b in [
    BelgeTuru("isletme_sozlesmesi", "İşletmelerde Mesleki Eğitim Sözleşmesi", "yerlestirme", _sablon_isletme_sozlesmesi),
    BelgeTuru("ciraklik_sozlesmesi", "Çıraklık Sözleşmesi (MEM)", "yerlestirme", _sablon_ciraklik_sozlesmesi),
    BelgeTuru("beceri_degerlendirme", "İşletmede Mesleki Eğitim Değerlendirme Formu", "yerlestirme", _sablon_beceri),
    BelgeTuru("koordinator_gorevlendirme", "Koordinatör Öğretmen Görevlendirme Yazısı", "ogretmen", _sablon_gorevlendirme),
    BelgeTuru("ziyaret_formu", "İşletme Ziyaret ve Rehberlik Formu", "ziyaret", _sablon_ziyaret_formu),
    BelgeTuru("devamsizlik_bildirimi", "Veliye Devamsızlık Bildirimi", "ogrenci", _sablon_devamsizlik),
    BelgeTuru("isletme_ucret_bildirimi", "İşletmeye Aylık Ücret Bildirimi", "isletme_ay", _sablon_ucret_bildirimi),
    BelgeTuru("koordinator_bildirimi", "İşletmeye Koordinatör Öğretmen Bildirimi", "isletme_koord",
              _sablon_koordinator_bildirimi),
    BelgeTuru("isletme_uygunluk", "İşletme Uygunluk Tespit Tutanağı", "isletme", _sablon_isletme_uygunluk),
    BelgeTuru("koordinatorluk_komisyon_karari", "Koordinatörlük Planlama Komisyonu Kararı", "kurum",
              _sablon_komisyon_karari),
    BelgeTuru("ogrenci_secim_tutanagi", "Öğrenci Seçim Komisyonu Tutanağı", "secim", _sablon_secim_tutanagi),
]}


def sablon_yolu(anahtar: str) -> Path:
    yol = config.templates_dir() / f"{anahtar}.docx"
    if not yol.exists():
        BELGE_TURLERI[anahtar].olusturucu(yol)
    return yol


def tum_sablonlari_hazirla(yeniden: bool = False) -> None:
    for k, b in BELGE_TURLERI.items():
        yol = config.templates_dir() / f"{k}.docx"
        if yeniden or not yol.exists():
            b.olusturucu(yol)


def _guvenli(s: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "_", s).strip()[:80]


def uret(db, anahtar: str, kimlik: int, klasor: Path | None = None, yil: int | None = None,
         ay: int | None = None, baglam: dict | None = None) -> Path:
    """`secim` kaynaklı belgelerde bağlam hazır verilir (`baglam`); diğerleri kayıt kimliğinden üretilir."""
    b = BELGE_TURLERI[anahtar]
    if b.kaynak == "secim":
        baglam = {**kurum_baglam(db), **(baglam or {})}
        ad = f"{baglam.get('alan') or 'Tüm alanlar'}_{date.today():%Y-%m-%d}"
    elif b.kaynak in ("isletme", "isletme_koord"):
        baglam = isletme_baglam(db, kimlik)
        ad = baglam["isletme_adi"]
    elif b.kaynak == "kurum":
        baglam = komisyon_baglam(db)
        ad = f"{baglam['egitim_yili'] or ''}_{date.today():%Y-%m-%d}"
    elif b.kaynak == "yerlestirme":
        baglam = yerlestirme_baglam(db, kimlik)
        ad = f"{baglam['ogrenci_ad_soyad']}_{baglam['isletme_adi'] or ''}"
    elif b.kaynak == "ogretmen":
        baglam = ogretmen_gorev_baglam(db, kimlik)
        ad = baglam["ogretmen_ad_soyad"]
    elif b.kaynak == "ziyaret":
        baglam = ziyaret_baglam(db, kimlik)
        ad = f"{baglam['ziyaret_tarihi']}_{baglam['isletme_adi']}"
    elif b.kaynak == "ogrenci":
        baglam = devamsizlik_baglam(db, kimlik)
        ad = baglam["ogrenci_ad_soyad"]
    elif b.kaynak == "isletme_ay":
        today = date.today()
        baglam = isletme_ucret_baglam(db, kimlik, yil or today.year, ay or today.month)
        ad = f"{baglam['isletme_adi']}_{baglam['donem']}"
    else:
        raise ValueError(b.kaynak)
    klasor = klasor or config.output_dir() / b.etiket
    return doldur(sablon_yolu(anahtar), baglam, klasor / f"{_guvenli(ad)}.docx")
