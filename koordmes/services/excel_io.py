"""Excel dışa aktarma. İçe aktarma `aktarim` modülündedir; eski adlar buradan da kullanılabilir."""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import okuyucu
from .aktarim import TANIMLAR, IceAktarmaSonucu, IceAktarmaTanimi, baslik_bul, ice_aktar, sablon_olustur  # noqa: F401

# ------------------------------------------------------------------ dışa aktarma

_INCE = Side(style="thin", color="999999")
_KENAR = Border(left=_INCE, right=_INCE, top=_INCE, bottom=_INCE)


class ExcelYazici:
    def __init__(self):
        self.wb = Workbook()
        self._ilk = True

    def sayfa(self, ad: str, basliklar: list[str], satirlar: list[list[Any]], baslik: str | None = None,
              alt_bilgi: list[str] | None = None) -> None:
        if self._ilk:
            ws = self.wb.active
            ws.title = ad[:31]
            self._ilk = False
        else:
            ws = self.wb.create_sheet(ad[:31])
        r = 1
        if baslik:
            ws.cell(row=1, column=1, value=baslik).font = Font(bold=True, size=13)
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(basliklar)))
            r = 3
        for c, h in enumerate(basliklar, 1):
            cell = ws.cell(row=r, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F5597")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _KENAR
        ws.freeze_panes = ws.cell(row=r + 1, column=1)
        for i, satir in enumerate(satirlar, r + 1):
            for c, v in enumerate(satir, 1):
                if isinstance(v, date) and not isinstance(v, datetime):
                    v = datetime(v.year, v.month, v.day)
                cell = ws.cell(row=i, column=c, value=v)
                cell.border = _KENAR
                if isinstance(v, datetime):
                    cell.number_format = "DD.MM.YYYY"
                elif isinstance(v, float):
                    cell.number_format = "#,##0.00"
        son = r + len(satirlar) + 2
        for j, t in enumerate(alt_bilgi or []):
            ws.cell(row=son + j, column=1, value=t).font = Font(italic=True)
        for c in range(1, len(basliklar) + 1):
            uz = max([len(str(basliklar[c - 1]))] + [len(str(s[c - 1])) for s in satirlar[:500] if c - 1 < len(s)
                                                     and s[c - 1] is not None])
            ws.column_dimensions[get_column_letter(c)].width = min(45, max(8, uz + 2))

    def kaydet(self, yol: Path | str) -> Path:
        yol = Path(yol)
        yol.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(yol)
        return yol


def tablo_yaz(yol: Path | str, basliklar: list[str], satirlar: list[list[Any]], baslik: str | None = None,
              sayfa: str = "Liste", alt_bilgi: list[str] | None = None) -> Path:
    x = ExcelYazici()
    x.sayfa(sayfa, basliklar, satirlar, baslik, alt_bilgi)
    return x.kaydet(yol)


def dosya_oku(yol: Path | str) -> list[list[Any]]:
    """Dosyadaki ilk tablonun satırları."""
    return okuyucu.oku(yol)[0].satirlar
