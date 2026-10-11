"""Koordinatörlük formları (barkodlu) ve gelen belge takibi.

Form türleri:
- G: Günlük Rehberlik Görev Formu — koordinatör öğretmenin her görev günü ve işletmesi için.
- R: Aylık Rehberlik Raporu — öğretmen ve işletme başına ayda bir.
- A: Koordinatör Öğretmen Aylık Görev Formu — öğretmen başına ayda bir (ek ders ödemesine esas).

Her formun üzerinde 20 haneli bir barkod bulunur: 7 + tür kodu (1/2/3) + yyyymmdd + öğretmen (5) + işletme (5).
Form öğretmen tarafından doldurulup teslim edildiğinde barkod okuyucuyla okutulur; teslim tarihi işlenir ve
günlük formlarda ilgili ziyaret 'Yapıldı' olarak işaretlenir.
"""
from __future__ import annotations

import io
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .. import config
from ..mevzuat import AYLAR, GUNLER
from . import barkod as bk
from .belgeler import _guvenli, _ortali, _yeni_belge, kurum_baglam
from .sgk_dosya import tr_buyuk
from .utils import ay_araligi, tarih, tatil_gunleri, tr

FORM_TURLERI = {"G": "Günlük Rehberlik Görev Formu", "R": "Aylık Rehberlik Raporu",
                "A": "Koordinatör Öğretmen Aylık Görev Formu"}
_KOD = {"G": "1", "R": "2", "A": "3"}
_TUR = {v: k for k, v in _KOD.items()}


def barkod_metni(tur: str, t: date, ogretmen_id: int, isletme_id: int | None) -> str:
    return f"7{_KOD[tur]}{t:%Y%m%d}{ogretmen_id:05d}{(isletme_id or 0):05d}"


@dataclass
class BarkodBilgisi:
    tur: str
    tarih: date
    ogretmen_id: int
    isletme_id: int | None


def barkod_coz(kod: str) -> BarkodBilgisi | None:
    k = "".join(c for c in str(kod or "") if c.isdigit())
    if len(k) != 20 or k[0] != "7" or k[1] not in _TUR:
        return None
    try:
        t = date(int(k[2:6]), int(k[6:8]), int(k[8:10]))
    except ValueError:
        return None
    return BarkodBilgisi(_TUR[k[1]], t, int(k[10:15]), int(k[15:20]) or None)


@dataclass
class Form:
    tur: str
    tarih: date  # günlük formda görev günü, aylıklarda ayın ilk günü
    ogretmen_id: int
    ogretmen: str
    isletme_id: int | None = None
    isletme: str = ""
    gunler: tuple = ()  # aylık formlarda (tarih, işletme, saat) görev günleri

    @property
    def barkod(self) -> str:
        return barkod_metni(self.tur, self.tarih, self.ogretmen_id, self.isletme_id)

    @property
    def donem(self) -> str:
        return tr(self.tarih) if self.tur == "G" else f"{AYLAR[self.tarih.month - 1]} {self.tarih.year}"


def gorev_gunleri(db, bas: date, bit: date, ogretmen_id: int | None = None) -> list[tuple]:
    """(tarih, öğretmen_id, öğretmen, işletme_id, işletme, saat) — tatiller hariç koordinatörlük görev günleri."""
    tatiller = tatil_gunleri(db, bas, bit)
    kosul, params = ("WHERE g.ogretmen_id = ?", [ogretmen_id]) if ogretmen_id else ("", [])
    out = []
    for g in db.q(f"""SELECT g.*, o.ad || ' ' || o.soyad AS ogretmen, i.ad AS isletme FROM koordinator_gorevleri g
                      JOIN ogretmenler o ON o.id = g.ogretmen_id JOIN isletmeler i ON i.id = g.isletme_id {kosul}
                      ORDER BY ogretmen, i.ad""", params):
        gb, ge = tarih(g["baslangic"]), tarih(g["bitis"])
        d = bas
        while d <= bit:
            if d.isoweekday() == g["gun"] and d not in tatiller and (gb is None or d >= gb) and (ge is None or d <= ge):
                out.append((d, g["ogretmen_id"], g["ogretmen"], g["isletme_id"], g["isletme"], g["haftalik_saat"]))
            d += timedelta(days=1)
    return sorted(out, key=lambda x: (x[2], x[0], x[4]))


def formlar(db, turler: list[str], bas: date, bit: date, ogretmen_id: int | None = None) -> list[Form]:
    """Seçilen türlerde, aralıktaki görev günlerinden çıkan formlar (aylıklar aralıkla kesişen her ay için)."""
    out: list[Form] = []
    if "G" in turler:
        out += [Form("G", d, t, ta, i, ia) for d, t, ta, i, ia, _ in gorev_gunleri(db, bas, bit, ogretmen_id)]
    if "R" in turler or "A" in turler:
        ay = date(bas.year, bas.month, 1)
        while ay <= bit:
            ilk, son = ay_araligi(ay.year, ay.month)
            gunler = gorev_gunleri(db, ilk, son, ogretmen_id)
            if "R" in turler:
                gr: dict[tuple, list] = defaultdict(list)
                for g in gunler:
                    gr[(g[1], g[2], g[3], g[4])].append((g[0], g[4], g[5]))
                out += [Form("R", ilk, t, ta, i, ia, tuple(v)) for (t, ta, i, ia), v in gr.items()]
            if "A" in turler:
                go: dict[tuple, list] = defaultdict(list)
                for g in gunler:
                    go[(g[1], g[2])].append((g[0], g[4], g[5]))
                out += [Form("A", ilk, t, ta, None, "", tuple(v)) for (t, ta), v in go.items()]
            ay = son + timedelta(days=1)
    return sorted(out, key=lambda f: (f.ogretmen, f.tarih, "GRA".index(f.tur), f.isletme))


def kaydet(db, liste: list[Form]) -> int:
    """Formları takip listesine ekler (zaten olanlar atlanır); eklenen sayısı."""
    mevcut = {r["barkod"] for r in db.q("SELECT barkod FROM belge_takibi")}
    n = 0
    with db.conn:
        for f in liste:
            if f.barkod in mevcut:
                continue
            db.conn.execute("INSERT INTO belge_takibi (tur, tarih, ogretmen_id, isletme_id, barkod) "
                            "VALUES (?, ?, ?, ?, ?)", (f.tur, f.tarih.isoformat(), f.ogretmen_id, f.isletme_id,
                                                       f.barkod))
            mevcut.add(f.barkod)
            n += 1
    return n


# ------------------------------------------------------------------ Word

def _ogrenciler(db, isletme_id: int, ilk: date, son: date) -> list:
    return db.q("""SELECT o.ad || ' ' || o.soyad AS ad, o.okul_no, s.ad AS sinif,
                          (SELECT COALESCE(SUM(d.gun), 0) FROM devamsizlik d WHERE d.ogrenci_id = o.id
                           AND d.yer = 'İşletme' AND d.tarih BETWEEN ? AND ?) AS devamsiz
                   FROM yerlestirmeler y JOIN ogrenciler o ON o.id = y.ogrenci_id
                   LEFT JOIN siniflar s ON s.id = o.sinif_id
                   WHERE y.isletme_id = ? AND y.baslangic IS NOT NULL AND y.baslangic <= ?
                   AND (y.bitis IS NULL OR y.bitis = '' OR y.bitis >= ?) AND y.durum IN ('Aktif', 'Sona Erdi', 'Feshedildi')
                   ORDER BY s.ad, o.ad""", (ilk.isoformat(), son.isoformat(), isletme_id, son.isoformat(),
                                             ilk.isoformat()))


def _hucre_kalin(c, metin: str) -> None:
    c.text = metin
    for r in c.paragraphs[0].runs:
        r.bold = True


def _bilgi_tablosu(doc, satirlar: list[tuple[str, str]]):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for a, b in satirlar:
        c = t.add_row().cells
        _hucre_kalin(c[0], a)
        c[1].text = b
        c[0].width, c[1].width = Cm(5), Cm(12)
    return t


def _liste_tablosu(doc, basliklar: list[str], satirlar: list[list[str]], bos_satir: int = 0):
    t = doc.add_table(rows=1, cols=len(basliklar))
    t.style = "Table Grid"
    for i, b in enumerate(basliklar):
        _hucre_kalin(t.rows[0].cells[i], b)
    for s in satirlar + [[""] * len(basliklar)] * bos_satir:
        c = t.add_row().cells
        for i, v in enumerate(s):
            c[i].text = str(v)
    return t


def _kutu(doc, baslik: str, satir: int = 4):
    t = doc.add_table(rows=2, cols=1)
    t.style = "Table Grid"
    _hucre_kalin(t.cell(0, 0), baslik)
    t.cell(1, 0).text = "\n" * satir


def _imzalar(doc, unvanlar: list[tuple[str, str]]):
    doc.add_paragraph()
    t = doc.add_table(rows=3, cols=len(unvanlar))
    for i, (u, ad) in enumerate(unvanlar):
        t.cell(0, i).text, t.cell(1, i).text, t.cell(2, i).text = u, ad, "İmza:"
        for r in range(3):
            t.cell(r, i).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        t.cell(0, i).paragraphs[0].runs[0].bold = True


def _baslik(doc, k: dict, f: Form) -> None:
    ust = doc.add_table(rows=1, cols=2)
    sol, sag = ust.rows[0].cells
    sol.width, sag.width = Cm(10), Cm(7)
    p = sol.paragraphs[0]
    p.add_run(f"T.C.\n{k['kurum_ilce']} KAYMAKAMLIĞI\n{k['kurum_adi']} Müdürlüğü").bold = True
    p2 = sag.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2.add_run().add_picture(io.BytesIO(bk.png(f.barkod, 2, 60)), width=Cm(5.5))
    alt = sag.add_paragraph(f.barkod)
    alt.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    alt.runs[0].font.size = Pt(8)
    _ortali(doc, tr_buyuk(FORM_TURLERI[f.tur]), boyut=13)


def _form_yaz(db, doc, k: dict, f: Form) -> None:
    _baslik(doc, k, f)
    isl = db.q1("SELECT * FROM isletmeler WHERE id = ?", (f.isletme_id,)) if f.isletme_id else None
    imza_koord = ("Koordinatör Müdür Yardımcısı", k["koordinator_mudur_yrd"] or "")
    if f.tur == "G":
        _bilgi_tablosu(doc, [("Koordinatör Öğretmen", f.ogretmen), ("İşletme", f.isletme),
                             ("Adres / Telefon", f"{isl['adres'] or ''} {isl['telefon'] or ''}".strip() if isl else ""),
                             ("Usta Öğretici / Eğitici Personel", isl["usta_ogretici"] or "" if isl else ""),
                             ("Görev Tarihi", f"{tr(f.tarih)} {GUNLER[f.tarih.isoweekday() - 1]}"),
                             ("Giriş - Çıkış Saati", "....... : .......   /   ....... : .......")])
        doc.add_paragraph()
        ogr = _ogrenciler(db, f.isletme_id, f.tarih, f.tarih)
        _liste_tablosu(doc, ["Okul No", "Öğrenci", "Sınıf", "Devam (Var/Yok)", "Görüş ve Değerlendirme"],
                       [[o["okul_no"] or "", o["ad"], o["sinif"] or "", "", ""] for o in ogr], 0 if ogr else 3)
        doc.add_paragraph()
        _kutu(doc, "Yapılan Rehberlik ve Tespitler")
        _kutu(doc, "Öneriler / Alınması Gereken Önlemler", 3)
        _imzalar(doc, [("Koordinatör Öğretmen", f.ogretmen), ("İşletme Yetkilisi", isl["yetkili"] or "" if isl else ""),
                       imza_koord])
    elif f.tur == "R":
        ilk, son = ay_araligi(f.tarih.year, f.tarih.month)
        _bilgi_tablosu(doc, [("Koordinatör Öğretmen", f.ogretmen), ("İşletme", f.isletme), ("Dönem", f.donem),
                             ("Görev Günleri", ", ".join(tr(g[0]) for g in f.gunler))])
        doc.add_paragraph()
        ogr = _ogrenciler(db, f.isletme_id, ilk, son)
        _liste_tablosu(doc, ["Okul No", "Öğrenci", "Sınıf", "İşletmede Devamsızlık (gün)", "Başarı / Uyum Durumu"],
                       [[o["okul_no"] or "", o["ad"], o["sinif"] or "", f"{o['devamsiz']:g}", ""] for o in ogr],
                       0 if ogr else 3)
        doc.add_paragraph()
        _kutu(doc, "İşletmedeki Eğitimin Genel Değerlendirmesi")
        _kutu(doc, "Karşılaşılan Sorunlar ve Öneriler", 3)
        _imzalar(doc, [("Koordinatör Öğretmen", f.ogretmen), imza_koord])
    else:
        _bilgi_tablosu(doc, [("Koordinatör Öğretmen", f.ogretmen), ("Dönem", f.donem)])
        doc.add_paragraph()
        _liste_tablosu(doc, ["Tarih", "Gün", "İşletme", "Saat", "Yapıldı (İmza)"],
                       [[tr(g[0]), GUNLER[g[0].isoweekday() - 1], g[1], str(g[2]), ""] for g in f.gunler])
        p = doc.add_paragraph()
        p.add_run(f"Toplam görev: {len(f.gunler)} gün, {sum(g[2] for g in f.gunler)} saat").bold = True
        _kutu(doc, "Açıklama (yapılamayan görevler ve nedenleri)", 3)
        _imzalar(doc, [("Koordinatör Öğretmen", f.ogretmen), imza_koord, ("Okul Müdürü", k["mudur"] or "")])
    doc.add_paragraph().add_run("Bu form doldurulup imzalandıktan sonra koordinatör müdür yardımcısına teslim "
                                "edilir; teslimde barkod okutulur.").font.size = Pt(8)


def formlari_uret(db, liste: list[Form], klasor: Path | None = None) -> list[Path]:
    """Her öğretmen için tek Word dosyası (her form ayrı sayfa) üretir ve formları takibe ekler."""
    klasor = klasor or config.output_dir() / "Koordinatörlük Formları"
    k = kurum_baglam(db)
    ogretmene: dict[str, list[Form]] = defaultdict(list)
    for f in liste:
        ogretmene[f.ogretmen].append(f)
    yollar = []
    for ogretmen, fl in ogretmene.items():
        doc = _yeni_belge()
        for s in doc.sections:
            s.top_margin = s.bottom_margin = Cm(1.5)
        for i, f in enumerate(fl):
            if i:
                doc.add_page_break()
            _form_yaz(db, doc, k, f)
        ilk, son = min(f.tarih for f in fl), max(f.tarih for f in fl)
        yol = klasor / f"{_guvenli(ogretmen)}_{ilk:%Y-%m-%d}_{son:%Y-%m-%d}.docx"
        yol.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(yol))
        yollar.append(yol)
    kaydet(db, liste)
    return yollar


# ------------------------------------------------------------------ teslim

@dataclass
class TeslimSonucu:
    tamam: bool
    mesaj: str
    belge_id: int | None = None


def _etiket(db, b: BarkodBilgisi) -> str:
    o = db.q1("SELECT ad || ' ' || soyad AS ad FROM ogretmenler WHERE id = ?", (b.ogretmen_id,))
    i = db.q1("SELECT ad FROM isletmeler WHERE id = ?", (b.isletme_id,)) if b.isletme_id else None
    donem = tr(b.tarih) if b.tur == "G" else f"{AYLAR[b.tarih.month - 1]} {b.tarih.year}"
    return f"{FORM_TURLERI[b.tur]} - {o['ad'] if o else '?'}" + (f" - {i['ad']}" if i else "") + f" ({donem})"


def teslim_al(db, kod: str, teslim: date | None = None) -> TeslimSonucu:
    b = barkod_coz(kod)
    if b is None:
        return TeslimSonucu(False, f"Barkod tanınmadı: {kod}")
    if not db.q1("SELECT 1 FROM ogretmenler WHERE id = ?", (b.ogretmen_id,)):
        return TeslimSonucu(False, f"Barkoddaki öğretmen kayıtlı değil: {kod}")
    if b.isletme_id and not db.q1("SELECT 1 FROM isletmeler WHERE id = ?", (b.isletme_id,)):
        b.isletme_id = None
    teslim = teslim or date.today()
    barkod = barkod_metni(b.tur, b.tarih, b.ogretmen_id, b.isletme_id)
    etiket = _etiket(db, b)
    r = db.q1("SELECT id, teslim_tarihi FROM belge_takibi WHERE barkod = ?", (barkod,))
    if r and r["teslim_tarihi"]:
        return TeslimSonucu(False, f"Zaten teslim alınmış ({tr(r['teslim_tarihi'])}): {etiket}", r["id"])
    with db.conn:
        if r:
            bid = r["id"]
            db.conn.execute("UPDATE belge_takibi SET teslim_tarihi = ? WHERE id = ?", (teslim.isoformat(), bid))
        else:
            bid = db.conn.execute("INSERT INTO belge_takibi (tur, tarih, ogretmen_id, isletme_id, barkod, "
                                  "teslim_tarihi) VALUES (?, ?, ?, ?, ?, ?)",
                                  (b.tur, b.tarih.isoformat(), b.ogretmen_id, b.isletme_id, barkod,
                                   teslim.isoformat())).lastrowid
        if b.tur == "G" and b.isletme_id:
            z = db.conn.execute("SELECT id FROM ziyaretler WHERE ogretmen_id = ? AND isletme_id = ? AND tarih = ?",
                                (b.ogretmen_id, b.isletme_id, b.tarih.isoformat())).fetchone()
            if z:
                db.conn.execute("UPDATE ziyaretler SET durum = 'Yapıldı' WHERE id = ?", (z[0],))
            else:
                db.conn.execute("INSERT INTO ziyaretler (ogretmen_id, isletme_id, tarih, durum) "
                                "VALUES (?, ?, ?, 'Yapıldı')", (b.ogretmen_id, b.isletme_id, b.tarih.isoformat()))
    return TeslimSonucu(True, f"Teslim alındı: {etiket}", bid)


def teslimi_geri_al(db, belge_id: int) -> None:
    db.update("belge_takibi", belge_id, {"teslim_tarihi": None})


def son_teslim(db, tur: str, t: date) -> date:
    gun = db.ayar("koord_form_teslim_gun")
    if tur == "G":
        return t + timedelta(days=gun)
    return ay_araligi(t.year, t.month)[1] + timedelta(days=gun)


@dataclass
class TakipSatiri:
    id: int
    tur: str
    tarih: date
    ogretmen: str
    isletme: str
    barkod: str
    son_teslim: date
    teslim: date | None
    aciklama: str

    @property
    def durum(self) -> str:
        if self.teslim:
            return "Teslim alındı" + (" (gecikmeli)" if self.teslim > self.son_teslim else "")
        return "Gecikti" if date.today() > self.son_teslim else "Bekliyor"


def takip_listesi(db, bas: date, bit: date, tur: str | None = None, ogretmen_id: int | None = None,
                  sadece_bekleyen: bool = False) -> list[TakipSatiri]:
    kosul, params = ["b.tarih BETWEEN ? AND ?"], [bas.isoformat(), bit.isoformat()]
    if tur:
        kosul.append("b.tur = ?")
        params.append(tur)
    if ogretmen_id:
        kosul.append("b.ogretmen_id = ?")
        params.append(ogretmen_id)
    if sadece_bekleyen:
        kosul.append("(b.teslim_tarihi IS NULL OR b.teslim_tarihi = '')")
    out = []
    for r in db.q(f"""SELECT b.*, o.ad || ' ' || o.soyad AS ogretmen, i.ad AS isletme FROM belge_takibi b
                      LEFT JOIN ogretmenler o ON o.id = b.ogretmen_id LEFT JOIN isletmeler i ON i.id = b.isletme_id
                      WHERE {' AND '.join(kosul)} ORDER BY b.tarih, ogretmen, isletme""", params):
        t = tarih(r["tarih"])
        out.append(TakipSatiri(r["id"], r["tur"], t, r["ogretmen"] or "", r["isletme"] or "", r["barkod"] or "",
                               son_teslim(db, r["tur"], t), tarih(r["teslim_tarihi"]), r["aciklama"] or ""))
    return out


def ozet(satirlar: list[TakipSatiri]) -> list[list]:
    """Öğretmen bazında [öğretmen, toplam, teslim, bekleyen, geciken]."""
    g: dict[str, list[int]] = defaultdict(lambda: [0, 0, 0, 0])
    for s in satirlar:
        x = g[s.ogretmen]
        x[0] += 1
        d = s.durum
        x[1] += d.startswith("Teslim")
        x[2] += d == "Bekliyor"
        x[3] += d == "Gecikti"
    return [[o, *v] for o, v in sorted(g.items())]
