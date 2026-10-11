"""Sınav planlama çıktıları: sınav takvimi, öğrenci giriş belgesi, öğretmen görevlendirme yazısı, görev ve ücret
çizelgesi, sınav onayı, öğretmen imza cetveli, yoklama listeleri, sınav / girmeyen öğrenci / kâğıt teslim / soru ve
cevap anahtarı tutanakları ve soru zarfı kapakları (Word / Excel)."""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any

from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .. import config
from . import sinav
from .belgeler import _guvenli, _ortali, _yeni_belge
from .excel_io import ExcelYazici
from .utils import tr


def klasor(db, donem_id: int, hedef: Path | None = None) -> Path:
    if hedef is not None:
        hedef.mkdir(parents=True, exist_ok=True)
        return hedef
    d = sinav.donem(db, donem_id)
    if d is None:
        raise ValueError("Sınav dönemi bulunamadı.")
    yol = config.output_dir() / "Sınav Planlama" / (_guvenli(d["ad"]) or f"Dönem {donem_id}")
    yol.mkdir(parents=True, exist_ok=True)
    return yol


def _buyuk(s: str | None) -> str:
    return (s or "").replace("i", "İ").replace("ı", "I").upper()


def _saat(r: Any, d: Any = None) -> str:
    ozel = sinav.ozel_saat(d, r["sekil"]) if d is not None and "sekil" in r.keys() else ""
    if ozel:
        return ozel
    return f"{r['baslangic']}-{r['bitis']}" if r["baslangic"] else (f"{r['sira']}. oturum" if r["sira"] else "")


def _donem_basligi(db, donem_id: int) -> str:
    d = sinav.donem(db, donem_id)
    if d["baslik_aciklama"]:
        return d["baslik_aciklama"].strip()
    yil = db.kurum()["egitim_yili"] or ""
    return f"{yil} Eğitim-Öğretim Yılı {d['ad']}".strip()


def _onaylayan(db, d: Any) -> tuple[str, str]:
    return d["onaylayan_unvan"] or "Okul Müdürü", d["onaylayan"] or db.kurum()["mudur"] or ""


def _imzacilar(db, d: Any) -> list[tuple[str, str]]:
    """Dönemin 'Başlık ve Onay Bilgisi': düzenleyen (varsa) ve onaylayan (yoksa okul müdürü)."""
    out = [(d["duzenleyen_unvan"] or "Düzenleyen", d["duzenleyen"])] if d["duzenleyen"] else []
    return out + [_onaylayan(db, d)]


def _onay_satiri(doc, db, d: Any) -> None:
    unvan, ad = _onaylayan(db, d)
    imza = doc.add_paragraph(f"{ad}\n{unvan}")
    imza.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _kurum_basligi(doc, db) -> None:
    from ..marka import kurum_logosu
    k = db.kurum()
    logo = kurum_logosu(db)
    if logo and Path(logo).exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(logo), height=Cm(1.8))
    _ortali(doc, "T.C.")
    if k["ilce"] or k["il"]:
        _ortali(doc, f"{_buyuk(k['ilce'] or k['il'])} {'KAYMAKAMLIĞI' if k['ilce'] else 'VALİLİĞİ'}")
    if k["ad"]:
        _ortali(doc, f"{_buyuk(k['ad'])} MÜDÜRLÜĞÜ")


def _tablo(doc, basliklar: list[str], satirlar: list[list[Any]], genislik: list[float] | None = None,
           boyut: int = 9):
    t = doc.add_table(rows=1, cols=len(basliklar))
    t.style = "Table Grid"
    for i, h in enumerate(basliklar):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(boyut)
    for s in satirlar:
        hucreler = t.add_row().cells
        for i, v in enumerate(s):
            hucreler[i].text = ""
            hucreler[i].paragraphs[0].add_run("" if v is None else str(v)).font.size = Pt(boyut)
    if genislik:
        for row in t.rows:
            for i, g in enumerate(genislik):
                row.cells[i].width = Cm(g)
    return t


def _imzalar(doc, kisiler: list[tuple[str, str]]) -> None:
    if not kisiler:
        return
    doc.add_paragraph()
    sutun = min(4, len(kisiler))
    t = doc.add_table(rows=0, cols=sutun)
    for i in range(0, len(kisiler), sutun):
        ust, ad, imza = t.add_row().cells, t.add_row().cells, t.add_row().cells
        for j, (unvan, kisi) in enumerate(kisiler[i:i + sutun]):
            ust[j].text, ad[j].text, imza[j].text = unvan, kisi, "İmza: ...................."
            for c in (ust[j], ad[j], imza[j]):
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            ust[j].paragraphs[0].runs[0].bold = True


def _yatay(doc) -> None:
    s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = s.page_height, s.page_width
    s.left_margin = s.right_margin = Cm(1.5)


def _planli(db, donem_id: int) -> list[dict]:
    rows = [r for r in sinav.sinav_listesi(db, donem_id) if r["seans_id"] and r["ogrenci_sayisi"]]
    if not rows:
        raise ValueError("Seansa yerleştirilmiş sınav yok. Önce 'Planlama' sekmesinden sınavları yerleştirin.")
    return rows


def _takvim_satiri(r: dict, d: Any = None) -> list[Any]:
    return [tr(r["tarih"]), sinav.gun_adi(r["tarih"]), _saat(r, d), r["ders"], r["seviye"] or "", r["grup"] or "",
            r["sekil"], r["sure_dk"], r["ogrenci_sayisi"], r["salonlar"], ", ".join(r["komisyon"]),
            ", ".join(r["gozculer"])]


TAKVIM_BASLIKLARI = ["Tarih", "Gün", "Saat", "Ders", "Seviye", "Grup", "Şekil", "Süre (dk)", "Öğrenci",
                     "Salon(lar)", "Komisyon", "Gözcüler"]


# ------------------------------------------------------------------ sınav takvimi

def sinav_takvimi_excel(db, donem_id: int, hedef: Path | None = None) -> Path:
    rows = _planli(db, donem_id)
    d = sinav.donem(db, donem_id)
    baslik = f"{_donem_basligi(db, donem_id)} Sınav Takvimi"
    x = ExcelYazici()
    x.sayfa("Sınav Takvimi", TAKVIM_BASLIKLARI, [_takvim_satiri(r, d) for r in rows], baslik=baslik,
            alt_bilgi=[f"{db.kurum()['ad'] or ''} - {tr(date.today())}"])
    program = sinav.ogrenci_programi(db, donem_id)
    ogr = [[r["okul_no"], f"{r['ad']} {r['soyad']}", r["sinif"] or "", r["ders"], r["seviye"] or "", r["grup"] or "",
            r["tur"] or "", tr(r["tarih"]) if r["tarih"] else "Yerleşmedi", _saat(r, d) if r["tarih"] else "",
            r["salon"] or "", r["sira_no"] or ""] for r in program]
    x.sayfa("Öğrenci Sınavları", ["Okul No", "Adı Soyadı", "Sınıf", "Ders", "Seviye", "Grup", "Türü", "Tarih", "Saat",
                                  "Salon", "Sıra No"], ogr, baslik=f"{baslik} - Öğrenci Listesi")
    seanslar: dict[tuple, list[dict]] = defaultdict(list)
    for r in rows:
        seanslar[(r["tarih"], r["sira"])].append(r)
    x.sayfa("Seans Özeti", ["Tarih", "Gün", "Saat", "Sınav", "Öğrenci", "Dersler", "Salon(lar)"],
            [[tr(t), sinav.gun_adi(t), _saat(liste[0]), len(liste), sum(r["ogrenci_sayisi"] for r in liste),
              ", ".join(sinav.sinav_adi(r) for r in liste), ", ".join(r["salonlar"] for r in liste if r["salonlar"])]
             for (t, _), liste in seanslar.items()], baslik=f"{baslik} - Seans Özeti")
    yerlesmeyen = [r for r in sinav.sinav_listesi(db, donem_id) if not r["seans_id"] and r["ogrenci_sayisi"]]
    if yerlesmeyen:
        x.sayfa("Yerleşmeyen Sınavlar", ["Ders", "Seviye", "Grup", "Öğrenci"],
                [[r["ders"], r["seviye"] or "", r["grup"] or "", r["ogrenci_sayisi"]] for r in yerlesmeyen])
    salonlar: dict[str, list] = defaultdict(list)
    for r in sorted((r for r in program if r["salon"] and r["tarih"]),
                    key=lambda r: (r["salon"], r["tarih"], r["sira"], r["sira_no"] or 0)):
        salonlar[r["salon"]].append(r)
    adlar: set[str] = set()
    for salon, liste in salonlar.items():
        ad = re.sub(r"[\[\]:*?/\\]", "_", f"Salon - {salon}")[:31]
        while ad in adlar:
            ad = ad[:28] + f" {len(adlar)}"
        adlar.add(ad)
        x.sayfa(ad, ["Tarih", "Saat", "Sıra No", "Okul No", "Adı Soyadı", "Sınıf", "Ders"],
                [[tr(r["tarih"]), _saat(r, d), r["sira_no"] or "", r["okul_no"] or "", f"{r['ad']} {r['soyad']}",
                  r["sinif"] or "", sinav.sinav_adi(r)] for r in liste], baslik=f"{baslik} - {salon} Oturma Çizelgesi")
    return x.kaydet(klasor(db, donem_id, hedef) / "Sınav Takvimi.xlsx")


def sinav_takvimi_word(db, donem_id: int, hedef: Path | None = None) -> Path:
    rows = _planli(db, donem_id)
    d = sinav.donem(db, donem_id)
    nobet = sinav.nobetciler(db, donem_id)
    doc = _yeni_belge()
    _yatay(doc)
    _kurum_basligi(doc, db)
    _ortali(doc, f"{_donem_basligi(db, donem_id)} Sınav Takvimi".upper(), boyut=13)
    gunler: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        gunler[r["tarih"]].append(r)
    for t, liste in gunler.items():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.add_run(f"{tr(t)} {sinav.gun_adi(t)}").bold = True
        if t in nobet:
            p.add_run(f"    Nöbetçi Müdür Yardımcısı: {nobet[t]['ad']}")
        _tablo(doc, ["Saat", "Ders", "Seviye", "Grup", "Şekil", "Süre", "Öğrenci", "Salon(lar)", "Komisyon",
                     "Gözcüler"],
               [[_saat(r, d), r["ders"], r["seviye"] or "", r["grup"] or "", r["sekil"], f"{r['sure_dk']} dk",
                 r["ogrenci_sayisi"], r["salonlar"], ", ".join(r["komisyon"]), ", ".join(r["gozculer"])]
                for r in liste], [2.2, 4.5, 1.4, 1.2, 2.2, 1.4, 1.5, 3.5, 4.5, 4.5], boyut=8)
    _imzalar(doc, _imzacilar(db, d))
    yol = klasor(db, donem_id, hedef) / "Sınav Takvimi.docx"
    doc.save(str(yol))
    return yol


# ------------------------------------------------------------------ öğrenci giriş belgesi

def giris_belgeleri(db, donem_id: int, hedef: Path | None = None) -> Path:
    """Her öğrenci için sınav giriş belgesi (sayfada iki öğrenci)."""
    rows = sinav.ogrenci_programi(db, donem_id)
    ogrenciler: dict[int, list] = defaultdict(list)
    for r in rows:
        ogrenciler[r["ogrenci_id"]].append(r)
    if not ogrenciler:
        raise ValueError("Sınavlara kayıtlı öğrenci yok.")
    doc = _yeni_belge()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.2)
    kurum = db.kurum()
    d = sinav.donem(db, donem_id)
    baslik = _donem_basligi(db, donem_id)
    for i, liste in enumerate(ogrenciler.values()):
        if i and i % 2 == 0:
            doc.add_page_break()
        o = liste[0]
        _ortali(doc, kurum["ad"] or "", boyut=11)
        _ortali(doc, f"{baslik} - SINAV GİRİŞ BELGESİ", boyut=11)
        p = doc.add_paragraph()
        p.add_run("Öğrenci: ").bold = True
        p.add_run(f"{o['ad']} {o['soyad']}    ")
        p.add_run("Okul No: ").bold = True
        p.add_run(f"{o['okul_no'] or ''}    ")
        p.add_run("Sınıf: ").bold = True
        p.add_run(o["sinif"] or "")
        _tablo(doc, ["Tarih", "Gün", "Saat", "Ders", "Seviye", "Türü", "Şekil", "Salon", "Sıra No"],
               [[tr(r["tarih"]) if r["tarih"] else "Belirlenmedi", sinav.gun_adi(r["tarih"]) if r["tarih"] else "",
                 _saat(r, d) if r["tarih"] else "", r["ders"] + (f" ({r['grup']})" if r["grup"] else ""),
                 r["seviye"] or "", r["tur"] or "", r["sekil"], r["salon"] or "", r["sira_no"] or ""] for r in liste],
               [2.2, 2.0, 2.0, 4.4, 1.3, 1.8, 2.2, 2.0, 1.3])
        n = doc.add_paragraph("Sınava kimlik belgenizle, sınav saatinden en az 15 dakika önce geliniz.")
        n.runs[0].italic = True
        n.runs[0].font.size = Pt(9)
        _onay_satiri(doc, db, d)
        if i % 2 == 0:
            ayrac = doc.add_paragraph("- " * 60)
            ayrac.alignment = WD_ALIGN_PARAGRAPH.CENTER
    yol = klasor(db, donem_id, hedef) / "Öğrenci Sınav Giriş Belgeleri.docx"
    doc.save(str(yol))
    return yol


# ------------------------------------------------------------------ öğretmen görevleri

def _ogretmen_gorevleri(db, donem_id: int) -> dict[int, list]:
    out: dict[int, list] = defaultdict(list)
    for r in db.q("""SELECT g.ogretmen_id, g.gorev, o.ad, o.soyad, o.brans, o.unvan, s.ders, s.seviye, s.grup,
                            s.sekil, s.sure_dk, se.tarih, se.sira, se.baslangic, se.bitis, sl.ad AS salon
                     FROM sinav_gorevlileri g JOIN sinavlar s ON s.id = g.sinav_id
                     JOIN ogretmenler o ON o.id = g.ogretmen_id
                     JOIN sinav_seanslari se ON se.id = s.seans_id
                     LEFT JOIN sinav_salonlari sl ON sl.id = g.salon_id
                     WHERE s.donem_id = ? ORDER BY o.ad, o.soyad, o.id, se.tarih, se.sira""", (donem_id,)):
        out[r["ogretmen_id"]].append(r)
    if not out:
        raise ValueError("Görevlendirme yok. Önce 'Planlama' sekmesinden görevlileri atayın.")
    return out


def gorevlendirme_yazilari(db, donem_id: int, hedef: Path | None = None) -> Path:
    """Her öğretmen için bir sayfa: bütün sınav görevleri ve tebliğ imzası."""
    gorevler = _ogretmen_gorevleri(db, donem_id)
    d = sinav.donem(db, donem_id)
    kurum = db.kurum()
    doc = _yeni_belge()
    for i, liste in enumerate(gorevler.values()):
        if i:
            doc.add_page_break()
        o = liste[0]
        _kurum_basligi(doc, db)
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.add_run("Sayı  : ").bold = True
        p.add_run("\t\t\t\t\t\t\t\t" + tr(date.today()))
        p = doc.add_paragraph()
        p.add_run("Konu : ").bold = True
        p.add_run("Sınav Görevlendirmesi")
        doc.add_paragraph()
        _ortali(doc, f"Sayın {o['ad']} {o['soyad']}", boyut=12)
        _ortali(doc, f"{o['brans'] or o['unvan'] or 'Öğretmen'}", kalin=False, boyut=11)
        metin = (f"{kurum['egitim_yili'] or ''} eğitim-öğretim yılı {d['ad']} ({d['tur']}) sınavlarında aşağıda "
                 f"belirtilen tarih ve saatlerde görevlendirilmiş bulunmaktasınız. Sınav başlamadan en az 30 dakika "
                 f"önce okulda bulunarak görevinizi yerine getirmenizi rica ederim.")
        par = doc.add_paragraph(metin.strip())
        par.paragraph_format.first_line_indent = Cm(1.25)
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _tablo(doc, ["Tarih", "Gün", "Saat", "Ders", "Seviye", "Görevi", "Salon"],
               [[tr(r["tarih"]), sinav.gun_adi(r["tarih"]), _saat(r, d),
                 r["ders"] + (f" ({r['grup']})" if r["grup"] else ""), r["seviye"] or "", r["gorev"], r["salon"] or ""]
                for r in liste], [2.3, 2.2, 2.4, 5.0, 1.4, 3.2, 2.3])
        doc.add_paragraph()
        _onay_satiri(doc, db, d)
        doc.add_paragraph()
        doc.add_paragraph("Tebliğ eden: ........................................\t\tTebellüğ eden: "
                          f"{o['ad']} {o['soyad']}")
        doc.add_paragraph("Tarih: ....../....../..........\t\t\t\tİmza: ............................")
    yol = klasor(db, donem_id, hedef) / "Öğretmen Görevlendirme Yazıları.docx"
    doc.save(str(yol))
    return yol


def gorev_cizelgesi_excel(db, donem_id: int, hedef: Path | None = None) -> Path:
    gorevler = _ogretmen_gorevleri(db, donem_id)
    ozet, ayrinti = [], []
    gunler = sorted({r["tarih"] for liste in gorevler.values() for r in liste})
    for liste in gorevler.values():
        o = liste[0]
        say = Counter(r["gorev"] for r in liste)
        gun = Counter(r["tarih"] for r in liste)
        ozet.append([f"{o['ad']} {o['soyad']}", o["brans"] or "", say[sinav.KOMISYON_BASKANI],
                     say[sinav.KOMISYON_UYESI], say[sinav.GOZCU], len(liste)] + [gun[t] or "" for t in gunler])
        for r in liste:
            ayrinti.append([f"{o['ad']} {o['soyad']}", tr(r["tarih"]), sinav.gun_adi(r["tarih"]), _saat(r), r["ders"],
                            r["seviye"] or "", r["grup"] or "", r["gorev"], r["salon"] or ""])
    for r in sinav.ogretmen_listesi(db, donem_id):
        if r["id"] not in gorevler and r["gorev_alir"]:
            ozet.append([f"{r['ad']} {r['soyad']}", r["brans"] or "", 0, 0, 0, 0] + ["" for _ in gunler])
    baslik = f"{_donem_basligi(db, donem_id)} Öğretmen Görev Çizelgesi"
    x = ExcelYazici()
    x.sayfa("Görev Sayıları", ["Öğretmen", "Branş", "Komisyon Başkanı", "Komisyon Üyesi", "Gözcü", "Toplam"]
            + [tr(t) for t in gunler], ozet, baslik=baslik)
    x.sayfa("Görev Listesi", ["Öğretmen", "Tarih", "Gün", "Saat", "Ders", "Seviye", "Grup", "Görev", "Salon"],
            ayrinti, baslik=baslik)
    return x.kaydet(klasor(db, donem_id, hedef) / "Öğretmen Görev Çizelgesi.xlsx")


def gorev_ucret_cizelgesi_excel(db, donem_id: int, hedef: Path | None = None) -> Path:
    """Sınav görevi ek ders ücreti: öğretmen başına ücretli görev ve saat, ayrıntıda her görevin ücret durumu."""
    liste = sinav.gorev_ucretleri(db, donem_id)
    if not liste:
        raise ValueError("Görevlendirme yok. Önce 'Planlama' sekmesinden görevlileri atayın.")
    d = sinav.donem(db, donem_id)
    baslik = f"{_donem_basligi(db, donem_id)} Sınav Görev Ücret Çizelgesi"
    ozet = [[k["ad"], k["brans"], k["onceki"], k["komisyon"], k["gozcu"], k["gelmedi"], k["ucretsiz"], k["sinir_disi"],
             k["ucretli_komisyon"], k["ucretli_gozcu"], float(k["saat"])] for k in liste]
    ozet.append(["TOPLAM", "", "", *[sum(s[i] for s in ozet) for i in range(3, 10)], float(sum(k["saat"] for k in liste))])
    ayrinti = [[k["ad"], tr(g["tarih"]), sinav.gun_adi(g["tarih"]), _saat(g), g["ders"], g["seviye"] or "", g["gorev"],
                g["ucret_durumu"], float(g["saat"])] for k in liste for g in k["gorevler"]]
    alt = [f"Komisyon görevi {d['komisyon_ucret_saat'] or 0:g} saat, gözcülük {d['gozcu_ucret_saat'] or 0:g} saat ek "
           f"ders olarak hesaplanmıştır."]
    if d["yillik_azami_gorev"]:
        alt.append(f"Yıllık ücretli görev sınırı: {d['yillik_azami_gorev']} (önceki dönemler dahil).")
    x = ExcelYazici()
    x.sayfa("Ücret Özeti", ["Öğretmen", "Branş", "Önceki Görev", "Komisyon", "Gözcülük", "Gelmedi", "Ücretsiz",
                            "Sınır Dışı", "Ücretli Komisyon", "Ücretli Gözcülük", "Ek Ders Saati"], ozet,
            baslik=baslik, alt_bilgi=alt)
    x.sayfa("Görev Ayrıntısı", ["Öğretmen", "Tarih", "Gün", "Saat", "Ders", "Seviye", "Görev", "Ücret Durumu",
                                "Ek Ders Saati"], ayrinti, baslik=baslik)
    return x.kaydet(klasor(db, donem_id, hedef) / "Sınav Görev Ücret Çizelgesi.xlsx")


# ------------------------------------------------------------------ sınav evrakı

def _sinav_bilgisi(doc, r: dict, salon: str | None = None, d: Any = None) -> None:
    _tablo(doc, ["Ders", "Seviye", "Grup", "Şekil", "Tarih", "Saat", "Süre", "Salon"],
           [[r["ders"], r["seviye"] or "", r["grup"] or "", r["sekil"], f"{tr(r['tarih'])} {sinav.gun_adi(r['tarih'])}",
             _saat(r, d), f"{r['sure_dk']} dk", salon if salon is not None else (r["salonlar"] or "")]],
           [4.0, 1.4, 1.2, 2.2, 3.2, 2.2, 1.4, 2.6])


def _katilim(r: dict) -> tuple[str, str]:
    """Yoklama girildiyse giren / girmeyen sayısı, girilmediyse doldurulacak boşluk."""
    if not (r["girdi"] or r["girmedi"]):
        return "........", "........"
    return str(r["ogrenci_sayisi"] - r["girmedi"]), str(r["girmedi"])


def _komisyon_imzalari(r: dict, gozcu: bool = True) -> list[tuple[str, str]]:
    kisiler = [(g["gorev"], g["ad"]) for g in sorted(r["gorevliler"], key=lambda g: g["gorev"] != sinav.KOMISYON_BASKANI)
               if g["gorev"] != sinav.GOZCU]
    if gozcu:
        kisiler += [(sinav.GOZCU, g) for g in r["gozculer"]]
    return kisiler or [(sinav.KOMISYON_BASKANI, ""), (sinav.KOMISYON_UYESI, "")]


def _paragraf(doc, metin: str) -> None:
    par = doc.add_paragraph(metin)
    par.paragraph_format.first_line_indent = Cm(1.25)
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def yoklama_listeleri(db, donem_id: int, hedef: Path | None = None) -> Path:
    """Her sınavın her salonu için yoklama / imza listesi."""
    rows = _planli(db, donem_id)
    d = sinav.donem(db, donem_id)
    baslik = _donem_basligi(db, donem_id)
    doc = _yeni_belge()
    ilk = True
    for r in rows:
        ogrenciler = sinav.sinav_ogrencileri(db, r["id"])
        salonlar: dict[str, list] = defaultdict(list)
        for o in ogrenciler:
            salonlar[o["salon"] or ""].append(o)
        gozcu = {g["salon"] or "": g["ad"] for g in r["gorevliler"] if g["gorev"] == sinav.GOZCU}
        for salon, liste in salonlar.items():
            if not ilk:
                doc.add_page_break()
            ilk = False
            _ortali(doc, db.kurum()["ad"] or "", boyut=11)
            _ortali(doc, f"{baslik} - SINAV YOKLAMA LİSTESİ", boyut=11)
            _sinav_bilgisi(doc, r, salon or "—", d)
            doc.add_paragraph()
            _tablo(doc, ["S.No", "Okul No", "Adı Soyadı", "Sınıfı", "Türü", "Sıra No", "İmza"],
                   [[i, o["okul_no"] or "", f"{o['ad']} {o['soyad']}", o["sinif"] or "", o["tur"] or "",
                     o["sira_no"] or "", sinav.GIRMEDI if o["katilim"] == sinav.GIRMEDI else ""]
                    for i, o in enumerate(liste, 1)], [1.2, 2.0, 5.4, 1.8, 1.6, 1.4, 3.4], boyut=10)
            if any(o["katilim"] for o in liste):
                girmeyen = sum(o["katilim"] == sinav.GIRMEDI for o in liste)
                doc.add_paragraph(f"Sınava giren: {len(liste) - girmeyen}   Girmeyen: {girmeyen}   "
                                  f"Toplam: {len(liste)}")
            else:
                doc.add_paragraph(f"Sınava giren: ........   Girmeyen: ........   Toplam: {len(liste)}")
            kisiler = [(g["gorev"], g["ad"]) for g in r["gorevliler"] if g["gorev"] != sinav.GOZCU]
            if salon in gozcu:
                kisiler.append((sinav.GOZCU, gozcu[salon]))
            _imzalar(doc, kisiler)
    yol = klasor(db, donem_id, hedef) / "Sınav Yoklama Listeleri.docx"
    doc.save(str(yol))
    return yol


def sinav_tutanaklari(db, donem_id: int, hedef: Path | None = None) -> Path:
    rows = _planli(db, donem_id)
    d = sinav.donem(db, donem_id)
    baslik = _donem_basligi(db, donem_id)
    doc = _yeni_belge()
    for i, r in enumerate(rows):
        if i:
            doc.add_page_break()
        _kurum_basligi(doc, db)
        _ortali(doc, f"{baslik}", boyut=11)
        _ortali(doc, "SINAV TUTANAĞI", boyut=13)
        doc.add_paragraph()
        _sinav_bilgisi(doc, r, d=d)
        doc.add_paragraph()
        giren, girmeyen = _katilim(r)
        _paragraf(doc, f"Okulumuz {r['seviye'] or ''}{'. sınıf ' if r['seviye'] else ''}{r['ders']} dersi "
                       f"{r['sekil'].lower()} sınavı yukarıda belirtilen tarih ve saatte komisyonumuzca yapılmıştır. "
                       f"Sınava kayıtlı {r['ogrenci_sayisi']} öğrenciden {giren} öğrenci sınava girmiş, {girmeyen} "
                       f"öğrenci girmemiştir. Sınav sırasında ......................................................"
                       f"...................................................................................... İş bu "
                       f"tutanak tarafımızca düzenlenip imza altına alınmıştır.")
        doc.add_paragraph("Sınav başlama saati: ............   Bitiş saati: ............   "
                          "Soru kâğıdı sayısı: ............   Cevap kâğıdı sayısı: ............")
        _imzalar(doc, _komisyon_imzalari(r))
    yol = klasor(db, donem_id, hedef) / "Sınav Tutanakları.docx"
    doc.save(str(yol))
    return yol


def girmeyen_ogrenci_tutanaklari(db, donem_id: int, hedef: Path | None = None) -> Path:
    """Her sınav için sınava girmeyen öğrenci tutanağı. Yoklama girilmemiş sınavlarda elle doldurulacak boş satırlar."""
    rows = _planli(db, donem_id)
    d = sinav.donem(db, donem_id)
    baslik = _donem_basligi(db, donem_id)
    doc = _yeni_belge()
    for i, r in enumerate(rows):
        if i:
            doc.add_page_break()
        _kurum_basligi(doc, db)
        _ortali(doc, baslik, boyut=11)
        _ortali(doc, "SINAVA GİRMEYEN ÖĞRENCİ TUTANAĞI", boyut=13)
        doc.add_paragraph()
        _sinav_bilgisi(doc, r, d=d)
        doc.add_paragraph()
        girmeyen = [o for o in sinav.sinav_ogrencileri(db, r["id"]) if o["katilim"] == sinav.GIRMEDI]
        if r["girdi"] or girmeyen:
            _paragraf(doc, f"Yukarıda bilgileri yazılı sınava kayıtlı {r['ogrenci_sayisi']} öğrenciden aşağıda adı "
                           f"yazılı {len(girmeyen)} öğrenci sınava girmemiştir.")
            satirlar = [[j, o["okul_no"] or "", f"{o['ad']} {o['soyad']}", o["sinif"] or "", o["tur"] or "",
                         o["salon"] or "", ""] for j, o in enumerate(girmeyen, 1)]
        else:
            _paragraf(doc, f"Yukarıda bilgileri yazılı sınava kayıtlı {r['ogrenci_sayisi']} öğrenciden aşağıda adı "
                           f"yazılı ........ öğrenci sınava girmemiştir.")
            satirlar = [[j, "", "", "", "", "", ""] for j in range(1, 6)]
        _tablo(doc, ["S.No", "Okul No", "Adı Soyadı", "Sınıfı", "Türü", "Salon", "Açıklama"],
               satirlar or [["", "", "Sınava girmeyen öğrenci yoktur.", "", "", "", ""]],
               [1.2, 2.0, 5.2, 1.8, 1.6, 2.0, 3.0], boyut=10)
        _imzalar(doc, _komisyon_imzalari(r))
    yol = klasor(db, donem_id, hedef) / "Sınava Girmeyen Öğrenci Tutanakları.docx"
    doc.save(str(yol))
    return yol


def sinav_onayi(db, donem_id: int, hedef: Path | None = None) -> Path:
    """Sınav programı ve komisyonlarının müdürlük onayı (Olur yazısı)."""
    rows = _planli(db, donem_id)
    d = sinav.donem(db, donem_id)
    doc = _yeni_belge()
    _yatay(doc)
    _kurum_basligi(doc, db)
    p = doc.add_paragraph()
    p.add_run("Sayı  : ").bold = True
    p.add_run("\t\t\t\t\t\t\t\t\t\t\t" + tr(date.today()))
    p = doc.add_paragraph()
    p.add_run("Konu : ").bold = True
    p.add_run(f"{d['ad']} Sınav Programı")
    _ortali(doc, "MÜDÜRLÜK MAKAMINA", boyut=12)
    _paragraf(doc, f"{_donem_basligi(db, donem_id)} sınavlarının {tr(rows[0]['tarih'])} - {tr(rows[-1]['tarih'])} "
                   f"tarihleri arasında aşağıdaki program, komisyon ve gözcülerle yapılması hususunu olurlarınıza "
                   f"arz ederim.")
    _tablo(doc, ["Tarih", "Saat", "Ders", "Seviye", "Şekil", "Öğrenci", "Salon(lar)", "Komisyon", "Gözcüler"],
           [[f"{tr(r['tarih'])} {sinav.gun_adi(r['tarih'])}", _saat(r, d), r["ders"] + (f" ({r['grup']})" if r["grup"]
             else ""), r["seviye"] or "", r["sekil"], r["ogrenci_sayisi"], r["salonlar"], ", ".join(r["komisyon"]),
             ", ".join(r["gozculer"])] for r in rows], [3.2, 2.0, 4.5, 1.3, 2.0, 1.4, 3.2, 4.5, 4.0], boyut=8)
    doc.add_paragraph()
    t = doc.add_table(rows=4, cols=2)
    duzenleyen = (d["duzenleyen_unvan"] or "Müdür Yardımcısı", d["duzenleyen"] or "")
    for i, metin in enumerate(["Arz ederim.", duzenleyen[1], duzenleyen[0], ""]):
        t.cell(i, 0).text = metin
    unvan, ad = _onaylayan(db, d)
    for i, metin in enumerate(["OLUR", f"..../..../{date.today().year}", ad, unvan]):
        t.cell(i, 1).text = metin
    for row in t.rows:
        for c in row.cells:
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.cell(0, 1).paragraphs[0].runs[0].bold = True
    yol = klasor(db, donem_id, hedef) / "Sınav Onayı.docx"
    doc.save(str(yol))
    return yol


def ogretmen_imza_cetveli(db, donem_id: int, hedef: Path | None = None) -> Path:
    """Her sınav günü için görevli öğretmenlerin imza cetveli (nöbetçi müdür yardımcısı onaylı)."""
    gorevler = _ogretmen_gorevleri(db, donem_id)
    d = sinav.donem(db, donem_id)
    nobet = sinav.nobetciler(db, donem_id)
    baslik = _donem_basligi(db, donem_id)
    gunler: dict[str, list] = defaultdict(list)
    for liste in gorevler.values():
        for r in liste:
            gunler[r["tarih"]].append(r)
    doc = _yeni_belge()
    for i, t in enumerate(sorted(gunler)):
        if i:
            doc.add_page_break()
        _ortali(doc, db.kurum()["ad"] or "", boyut=11)
        _ortali(doc, f"{baslik} - ÖĞRETMEN İMZA CETVELİ", boyut=11)
        _ortali(doc, f"{tr(t)} {sinav.gun_adi(t)}", kalin=False, boyut=11)
        liste = sorted(gunler[t], key=lambda r: (r["sira"] or 0, r["ders"], r["gorev"] != sinav.KOMISYON_BASKANI,
                                                  r["gorev"] == sinav.GOZCU, r["ad"]))
        _tablo(doc, ["S.No", "Adı Soyadı", "Saat", "Ders", "Görevi", "Salon", "Geliş Saati", "İmza"],
               [[j, f"{r['ad']} {r['soyad']}", _saat(r, d), r["ders"] + (f" ({r['seviye']}. Sınıf)" if r["seviye"]
                 else ""), r["gorev"], r["salon"] or "", "", ""] for j, r in enumerate(liste, 1)],
               [1.0, 4.0, 2.0, 4.0, 2.6, 1.8, 1.6, 2.6], boyut=9)
        n = nobet.get(t)
        _imzalar(doc, [((n["unvan"] or "Nöbetçi Müdür Yardımcısı") if n else "Nöbetçi Müdür Yardımcısı",
                        n["ad"] if n else "")])
    yol = klasor(db, donem_id, hedef) / "Öğretmen İmza Cetveli.docx"
    doc.save(str(yol))
    return yol


def kagit_teslim_tutanaklari(db, donem_id: int, hedef: Path | None = None) -> Path:
    """Sınav sonunda cevap kâğıtlarının komisyondan okul idaresine teslim tutanağı."""
    rows = _planli(db, donem_id)
    d = sinav.donem(db, donem_id)
    nobet = sinav.nobetciler(db, donem_id)
    baslik = _donem_basligi(db, donem_id)
    doc = _yeni_belge()
    for i, r in enumerate(rows):
        if i:
            doc.add_page_break()
        _kurum_basligi(doc, db)
        _ortali(doc, baslik, boyut=11)
        _ortali(doc, "SINAV KÂĞITLARI TESLİM TUTANAĞI", boyut=13)
        doc.add_paragraph()
        _sinav_bilgisi(doc, r, d=d)
        doc.add_paragraph()
        giren, girmeyen = _katilim(r)
        _paragraf(doc, f"Yukarıda bilgileri yazılı sınava kayıtlı {r['ogrenci_sayisi']} öğrenciden {giren} öğrenci "
                       f"sınava girmiş, {girmeyen} öğrenci girmemiştir. Sınava giren öğrencilere ait ........ adet "
                       f"cevap kâğıdı, ........ adet artan soru kitapçığı, yoklama listesi ve sınav tutanağı zarf "
                       f"içinde kapalı ve mühürlü olarak teslim edilmiştir.")
        teslim_eden = [(f"Teslim Eden ({g})", a) for g, a in _komisyon_imzalari(r, gozcu=False)[:1]]
        n = nobet.get(r["tarih"])
        teslim_alan = ("Teslim Alan (Nöbetçi Müdür Yardımcısı)", n["ad"]) if n else \
            ("Teslim Alan", d["duzenleyen"] or "")
        _imzalar(doc, teslim_eden + [teslim_alan])
    yol = klasor(db, donem_id, hedef) / "Sınav Kâğıdı Teslim Tutanakları.docx"
    doc.save(str(yol))
    return yol


def soru_cevap_anahtari_tutanaklari(db, donem_id: int, hedef: Path | None = None) -> Path:
    """Komisyonun hazırladığı soru ve cevap anahtarının okul idaresine teslim tutanağı."""
    rows = _planli(db, donem_id)
    d = sinav.donem(db, donem_id)
    baslik = _donem_basligi(db, donem_id)
    doc = _yeni_belge()
    for i, r in enumerate(rows):
        if i:
            doc.add_page_break()
        _kurum_basligi(doc, db)
        _ortali(doc, baslik, boyut=11)
        _ortali(doc, "SORU VE CEVAP ANAHTARI TESLİM TUTANAĞI", boyut=13)
        doc.add_paragraph()
        _sinav_bilgisi(doc, r, d=d)
        doc.add_paragraph()
        _paragraf(doc, f"{r['ders']} dersi {r['sekil'].lower()} sınavına ait sorular ve cevap anahtarı (puanlama "
                       f"ölçütleriyle birlikte) komisyonumuzca hazırlanmış; ........ sayfa soru metni ve ........ sayfa "
                       f"cevap anahtarı kapalı zarf içinde okul idaresine teslim edilmiştir.")
        _imzalar(doc, _komisyon_imzalari(r, gozcu=False) + [("Teslim Alan", d["duzenleyen"] or "")])
    yol = klasor(db, donem_id, hedef) / "Soru ve Cevap Anahtarı Teslim Tutanakları.docx"
    doc.save(str(yol))
    return yol


def zarf_kapaklari(db, donem_id: int, hedef: Path | None = None) -> Path:
    rows = _planli(db, donem_id)
    d = sinav.donem(db, donem_id)
    baslik = _donem_basligi(db, donem_id)
    doc = _yeni_belge()
    for i, r in enumerate(rows):
        if i:
            doc.add_page_break()
        _ortali(doc, db.kurum()["ad"] or "", boyut=16)
        _ortali(doc, baslik, boyut=14)
        doc.add_paragraph()
        _ortali(doc, "SINAV SORU ZARFI", boyut=26)
        doc.add_paragraph()
        _ortali(doc, _buyuk(r["ders"]), boyut=30)
        if r["seviye"]:
            _ortali(doc, f"{r['seviye']}. SINIF" + (f" - {r['grup']} GRUBU" if r["grup"] else ""), boyut=24)
        doc.add_paragraph()
        for etiket, deger in [("Tarih", f"{tr(r['tarih'])} {sinav.gun_adi(r['tarih'])}"), ("Saat", _saat(r, d)),
                              ("Süre", f"{r['sure_dk']} dakika"), ("Sınav şekli", r["sekil"]),
                              ("Salon", r["salonlar"] or "—"), ("Öğrenci sayısı", str(r["ogrenci_sayisi"])),
                              ("Komisyon", ", ".join(r["komisyon"]) or "—")]:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            a = p.add_run(f"{etiket}: ")
            a.bold = True
            a.font.size = Pt(16)
            p.add_run(deger).font.size = Pt(16)
        doc.add_paragraph()
        _ortali(doc, "Zarf, sınav salonunda öğrencilerin önünde açılacaktır.", kalin=False, boyut=12)
    yol = klasor(db, donem_id, hedef) / "Soru Zarfı Kapakları.docx"
    doc.save(str(yol))
    return yol


CIKTILAR = [
    ("Sınav Takvimi (Excel)", sinav_takvimi_excel),
    ("Sınav Takvimi / İlan (Word)", sinav_takvimi_word),
    ("Öğrenci Sınav Giriş Belgeleri (Word)", giris_belgeleri),
    ("Öğretmen Görevlendirme Yazıları (Word)", gorevlendirme_yazilari),
    ("Öğretmen Görev Çizelgesi (Excel)", gorev_cizelgesi_excel),
    ("Sınav Görev Ücret Çizelgesi (Excel)", gorev_ucret_cizelgesi_excel),
    ("Sınav Onayı (Word)", sinav_onayi),
    ("Öğretmen İmza Cetveli (Word)", ogretmen_imza_cetveli),
    ("Sınav Yoklama / İmza Listeleri (Word)", yoklama_listeleri),
    ("Sınav Tutanakları (Word)", sinav_tutanaklari),
    ("Sınava Girmeyen Öğrenci Tutanakları (Word)", girmeyen_ogrenci_tutanaklari),
    ("Sınav Kâğıdı Teslim Tutanakları (Word)", kagit_teslim_tutanaklari),
    ("Soru ve Cevap Anahtarı Tutanakları (Word)", soru_cevap_anahtari_tutanaklari),
    ("Soru Zarfı Kapakları (Word)", zarf_kapaklari),
]


def hepsini_uret(db, donem_id: int, hedef: Path | None = None) -> tuple[list[Path], list[str]]:
    """Üretilebilen bütün çıktılar ve üretilemeyenlerin nedenleri."""
    yollar, hatalar = [], []
    for ad, fn in CIKTILAR:
        try:
            yollar.append(fn(db, donem_id, hedef))
        except ValueError as e:
            hatalar.append(f"{ad}: {e}")
    return yollar, hatalar
