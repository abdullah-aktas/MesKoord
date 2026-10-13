import os
import random
import time
from collections import Counter, defaultdict
from datetime import date, timedelta

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

PAZARTESI = date(2026, 11, 2)


def _okul(db, ogrenci_basina: int = 8) -> dict:
    """Küçük bir okul: üç sınıf (9, 10, 12), kültür ve meslek dersleri, öğretmenler ve ders atamaları."""
    bt = db.insert("alanlar", {"ad": "Bilişim Teknolojileri"})
    ders = {ad: db.insert("dersler", {"ad": ad, "tur": tur, "alan_id": bt if tur != "Kültür" else None})
            for ad, tur in [("Matematik", "Kültür"), ("Fizik", "Kültür"), ("Kimya", "Kültür"), ("Tarih", "Kültür"),
                            ("Web Tasarımı", "Atölye"), ("Programlama Temelleri", "Meslek")]}
    sinif = {ad: db.insert("siniflar", {"ad": ad, "seviye": sv, "alan_id": bt if sv > 9 else None})
             for ad, sv in [("9-A", 9), ("10-A", 10), ("12-A", 12)]}
    ogretmen = {}
    for ad, soyad, brans, alan, unvan in [
            ("Cem", "Aksoy", "Matematik", None, "Öğretmen"), ("Deniz", "Tunç", "Matematik", None, "Öğretmen"),
            ("Oğuz", "Bayram", "Fizik", None, "Öğretmen"), ("Pınar", "Sezer", "Kimya", None, "Öğretmen"),
            ("Kemal", "Ateş", "Tarih", None, "Öğretmen"), ("Serdar", "Altun", "Bilişim Teknolojileri", bt, "Öğretmen"),
            ("Hande", "Işık", "Bilişim Teknolojileri", bt, "Öğretmen"), ("Seda", "Bulut", "İngilizce", None, "Öğretmen"),
            ("Gül", "Eren", "Biyoloji", None, "Öğretmen"), ("Yakup", "Karaca", "Din Kültürü", None, "Öğretmen"),
            ("Selma", "Uçar", "Matematik", None, "Müdür Yardımcısı")]:
        ogretmen[f"{ad} {soyad}"] = db.insert("ogretmenler", {"ad": ad, "soyad": soyad, "brans": brans,
                                                               "alan_id": alan, "unvan": unvan})
    db.insert("ders_atamalari", {"sinif_id": sinif["9-A"], "ders_id": ders["Matematik"],
                                 "ogretmen_id": ogretmen["Cem Aksoy"], "haftalik_saat": 4})
    db.insert("ders_atamalari", {"sinif_id": sinif["10-A"], "ders_id": ders["Web Tasarımı"],
                                 "ogretmen_id": ogretmen["Serdar Altun"], "haftalik_saat": 4})
    ogrenci = defaultdict(list)
    no = 100
    for ad_, sid in sinif.items():
        for i in range(ogrenci_basina):
            no += 1
            ogrenci[ad_].append(db.insert("ogrenciler", {"ad": f"Öğrenci{no}", "soyad": f"Soyad{i}",
                                                          "okul_no": str(no), "sinif_id": sid}))
    return {"ders": ders, "sinif": sinif, "ogretmen": ogretmen, "ogrenci": ogrenci}


def _donem(db, gun: int = 5, saatler=None) -> int:
    from koordmes.services import sinav
    did = sinav.donem_ekle(db, "2026-2027 Kasım Sorumluluk Sınavları")
    sinav.seanslari_olustur(db, did, sinav.tarihleri_uret(PAZARTESI, PAZARTESI + timedelta(days=gun - 1)),
                            saatler or sinav.VARSAYILAN_SAATLER[:4])
    return did


def _seans_bilgisi(db, did):
    from koordmes.services import sinav
    return {r["id"]: r for r in sinav.seanslar(db, did)}


# ------------------------------------------------------------------ seanslar

def test_seans_cizelgesi(db):
    from koordmes.services import sinav

    did = sinav.donem_ekle(db, "Deneme")
    tarihler = sinav.tarihleri_uret(PAZARTESI, PAZARTESI + timedelta(days=6), haric=[PAZARTESI + timedelta(days=2)])
    assert [t.weekday() for t in tarihler] == [0, 1, 3, 4]
    assert len(sinav.tarihleri_uret(PAZARTESI, PAZARTESI + timedelta(days=6), hafta_sonu=True)) == 7
    saatler = sinav.saatleri_coz("9:00-9.40, 10.00 - 10:40; 13:00-13:40")
    assert saatler == [("09:00", "09:40"), ("10:00", "10:40"), ("13:00", "13:40")]
    assert sinav.seanslari_olustur(db, did, tarihler, saatler) == 12
    assert sinav.seanslari_olustur(db, did, tarihler, saatler, sil=False) == 0
    d = sinav.donem(db, did)
    assert (d["baslangic"], d["bitis"]) == (PAZARTESI.isoformat(), (PAZARTESI + timedelta(days=4)).isoformat())

    ilk_gun = [r for r in sinav.seanslar(db, did) if r["tarih"] == PAZARTESI.isoformat()]
    sinav.seans_acik_yap(db, ilk_gun[2]["id"], False)
    assert sinav.gun_duzenini_uygula(db, did, PAZARTESI) == 3
    kapali = [r for r in sinav.seanslar(db, did) if not r["acik"]]
    assert len(kapali) == 4 and {r["sira"] for r in kapali} == {3}
    assert len(sinav.seanslar(db, did, sadece_acik=True)) == 8
    assert "Pazartesi 1. oturum (09:00-09:40)" in sinav.seans_etiketi(sinav.seanslar(db, did)[0])

    s = sinav.sinav_ekle(db, did, "Matematik", 9)
    sinav.seans_ata(db, s, ilk_gun[0]["id"])
    assert sinav.seanslari_olustur(db, did, tarihler[:1], saatler) == 3
    r = db.q1("SELECT seans_id, cakili FROM sinavlar WHERE id = ?", (s,))
    assert r["seans_id"] is None and r["cakili"] == 0


# ------------------------------------------------------------------ e-Okul aktarımı

def _xlsx(yol, satirlar):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for s in satirlar:
        ws.append(s)
    wb.save(yol)
    return yol


def test_ders_ve_seviye_cozumleme():
    from koordmes.services import sinav

    assert sinav.ders_ve_seviye("MATEMATİK (9. Sınıf)") == ("Matematik", 9)
    assert sinav.ders_ve_seviye("9. SINIF FİZİK") == ("Fizik", 9)
    assert sinav.ders_ve_seviye("Kimya - 10") == ("Kimya", 10)
    assert sinav.ders_ve_seviye("TÜRK DİLİ VE EDEBİYATI 11. SINIF") == ("Türk Dili ve Edebiyatı", 11)
    assert sinav.ders_ve_seviye("T.C. İNKILAP TARİHİ VE ATATÜRKÇÜLÜK") == ("T.C. İnkılap Tarihi ve Atatürkçülük", None)
    assert sinav.seviye_coz("10. Sınıf") == 10 and sinav.seviye_coz(11.0) == 11


def test_eokul_duz_tablo_aktarimi_tekrarlanabilir(db, tmp_path):
    from koordmes.services import sinav

    okul = _okul(db, 3)
    did = sinav.donem_ekle(db, "Sorumluluk")
    yol = _xlsx(tmp_path / "OOK12001R012.xlsx", [
        ["T.C. MİLLÎ EĞİTİM BAKANLIĞI"],
        ["SORUMLU DERSİ OLAN ÖĞRENCİLER VE DERSLERİ"],
        [],
        ["S.No", "Öğrenci No", "Adı Soyadı", "Sınıfı", "Sorumlu Olduğu Ders", "Dersin Sınıf Seviyesi", "Grup"],
        [1, 101, "ÖĞRENCİ101 SOYAD0", "9/A Şubesi", "MATEMATİK", 9, None],
        [2, 104, "ÖĞRENCİ104 SOYAD0", "10-A", "Matematik", "9. Sınıf", None],
        [3, 104, "ÖĞRENCİ104 SOYAD0", "10-A", "WEB TASARIMI", 10, None],
        [4, 107, "ÖĞRENCİ107 SOYAD0", "12-A", "FİZİK", 10, "B"],
        [5, 107, "ÖĞRENCİ107 SOYAD0", "12-A", "Fizik", 10, None],
        [6, 999, "OLMAYAN KİŞİ", "11-C", "KİMYA", 11, None],
        ["Toplam: 6 kayıt"],
    ])
    s = sinav.eokul_sorumlu_aktar(db, did, yol)
    assert s.kayit == 6 and s.sinav_eklenen == 4 and s.ogrenci_eklenen == 5
    assert s.eslesmeyen == ["Olmayan Kişi (No: 999, 11-C)"]
    mat = db.q1("SELECT * FROM sinavlar WHERE donem_id = ? AND ders = 'Matematik' AND seviye = 9", (did,))
    assert {r["ogrenci_id"] for r in sinav.sinav_ogrencileri(db, mat["id"])} == \
        {okul["ogrenci"]["9-A"][0], okul["ogrenci"]["10-A"][0]}
    web = db.q1("SELECT * FROM sinavlar WHERE donem_id = ? AND ders = 'Web Tasarımı'", (did,))
    assert web["ders_id"] == okul["ders"]["Web Tasarımı"] and web["alan_id"] is not None
    assert mat["alan_id"] is None and mat["ders_id"] == okul["ders"]["Matematik"]
    assert db.q1("SELECT COUNT(*) n FROM sinavlar WHERE ders = 'Fizik'")["n"] == 2  # grup ayrı sınav

    tekrar = sinav.eokul_sorumlu_aktar(db, did, yol)
    assert (tekrar.sinav_eklenen, tekrar.ogrenci_eklenen, tekrar.sinav_mevcut, tekrar.ogrenci_mevcut) == (0, 0, 4, 5)

    olustur = sinav.eokul_sorumlu_aktar(db, did, yol, olustur=True)
    assert olustur.yeni_ogrenci == 1 and olustur.yeni_sinif == 1 and olustur.sinav_eklenen == 1
    yeni = db.q1("SELECT o.*, si.ad sinif, si.seviye FROM ogrenciler o JOIN siniflar si ON si.id = o.sinif_id "
                 "WHERE okul_no = '999'")
    assert (yeni["ad"], yeni["soyad"], yeni["sinif"], yeni["seviye"], yeni["aktif"]) == ("Olmayan", "Kişi", "11-C", 11, 1)
    assert sinav.eokul_sorumlu_aktar(db, did, yol, olustur=True).yeni_ogrenci == 0


def test_eokul_ogrenci_satiri_altinda_dersler(db, tmp_path):
    from koordmes.services import sinav

    okul = _okul(db, 3)
    did = sinav.donem_ekle(db, "Sorumluluk")
    yol = _xlsx(tmp_path / "devam.xlsx", [
        ["Okul No", "Adı", "Soyadı", "Sınıf/Şube", "Ders Adı"],
        [102, "Öğrenci102", "Soyad1", "9-A", None],
        [None, None, None, None, "MATEMATİK (9. Sınıf)"],
        [None, None, None, None, "9. SINIF FİZİK"],
        [105, "Öğrenci105", "Soyad1", "10-A", "KİMYA - 10"],
        [None, None, None, None, "Tarih 9, Matematik (10)"],
    ])
    s = sinav.eokul_sorumlu_aktar(db, did, yol)
    assert s.kayit == 5 and s.ogrenci_eklenen == 5 and not s.eslesmeyen
    sinavlar = {(r["ders"], r["seviye"]) for r in db.q("SELECT ders, seviye FROM sinavlar WHERE donem_id = ?", (did,))}
    assert sinavlar == {("Matematik", 9), ("Fizik", 9), ("Kimya", 10), ("Tarih", 9), ("Matematik", 10)}
    ali = okul["ogrenci"]["10-A"][1]
    assert db.q1("SELECT COUNT(*) n FROM sinav_ogrencileri WHERE ogrenci_id = ?", (ali,))["n"] == 3

    yol2 = _xlsx(tmp_path / "etiketli.xlsx", [
        ["SORUMLU DERSİ OLAN ÖĞRENCİLER VE DERSLERİ"],
        ["Öğrenci No: 108", None, "Adı Soyadı:", "ÖĞRENCİ108 SOYAD1", "Sınıfı: 12-A"],
        ["Ders Adı", "Sınıf Seviyesi"],
        ["Matematik", 9],
        ["Kimya", 10],
        ["Öğrenci No: 109", None, "Adı Soyadı:", "ÖĞRENCİ109 SOYAD2", "Sınıfı: 12-A"],
        ["Ders Adı", "Sınıf Seviyesi"],
        ["Kimya", 10],
    ])
    s2 = sinav.eokul_sorumlu_aktar(db, did, yol2)
    assert s2.kayit == 3 and s2.ogrenci_eklenen == 3 and not s2.eslesmeyen
    kimya = db.q1("SELECT id FROM sinavlar WHERE ders = 'Kimya' AND seviye = 10")["id"]
    assert {r["ogrenci_id"] for r in sinav.sinav_ogrencileri(db, kimya)} == \
        {okul["ogrenci"]["10-A"][1], okul["ogrenci"]["12-A"][1], okul["ogrenci"]["12-A"][2]}


def test_elle_sinav_ve_birlestirme(db):
    from koordmes.services import sinav

    okul = _okul(db, 3)
    did = sinav.donem_ekle(db, "Beceri", tur="Beceri")
    a = sinav.sinav_ekle(db, did, "MATEMATİK", 9)
    assert sinav.sinav_ekle(db, did, "Matematik", 9) == a
    b = sinav.sinav_ekle(db, did, "Matematik", 9, grup="B")
    assert sinav.sinav_ogrenci_ekle(db, a, okul["ogrenci"]["9-A"][:2]) == 2
    assert sinav.sinav_ogrenci_ekle(db, a, okul["ogrenci"]["9-A"][:2]) == 0
    sinav.sinav_ogrenci_ekle(db, b, okul["ogrenci"]["9-A"][1:3])
    assert sinav.sinavlari_birlestir(db, [a, b]) == 1
    assert len(sinav.sinav_ogrencileri(db, a)) == 3 and db.q1("SELECT 1 FROM sinavlar WHERE id = ?", (b,)) is None
    sinav.sinav_ogrenci_cikar(db, a, okul["ogrenci"]["9-A"][:1])
    assert len(sinav.sinav_ogrencileri(db, a)) == 2
    assert sinav.kural_ekle(db, did, "Matematik", "Fizik") == sinav.kural_ekle(db, did, "fizik", "MATEMATİK")
    with pytest.raises(ValueError):
        sinav.kural_ekle(db, did, "Matematik", "matematik")


# ------------------------------------------------------------------ planlama

def _senaryo(db):
    """Çakışan öğrencili sınavlar, uygulamalı sınav, kural, kapalı seans ve çakılı sınav."""
    from koordmes.services import sinav

    okul = _okul(db, 8)
    did = _donem(db, gun=5)
    o = okul["ogrenci"]
    tum = o["9-A"] + o["10-A"] + o["12-A"]
    rng = random.Random(5)
    sinavlar = {}
    for ders, sv, kitle in [("Matematik", 9, tum[:14]), ("Fizik", 9, tum[4:18]), ("Kimya", 10, tum[10:22]),
                            ("Tarih", 9, tum[::2]), ("Matematik", 10, tum[1::3]), ("Fizik", 10, tum[12:24]),
                            ("Programlama Temelleri", 10, o["10-A"] + o["12-A"][:3]),
                            ("Web Tasarımı", 10, o["10-A"][2:] + o["12-A"][:4]),
                            ("Kimya", 9, rng.sample(tum, 9)), ("Tarih", 10, rng.sample(tum, 7))]:
        sid = sinav.sinav_ekle(db, did, ders, sv)
        sinav.sinav_ogrenci_ekle(db, sid, kitle)
        sinavlar[(ders, sv)] = sid
    sinav.sinav_guncelle(db, sinavlar[("Web Tasarımı", 10)], sekil="Uygulama", sure_dk=80)
    sinav.kural_ekle(db, did, "Matematik", "Fizik")
    seans = sinav.seanslar(db, did)
    sinav.seans_ata(db, sinavlar[("Tarih", 9)], seans[5]["id"])
    kapali = [r["id"] for r in seans if r["tarih"] == PAZARTESI.isoformat()]
    sinav.ogrenci_kapali_ayarla(db, did, tum[0], kapali)
    return did, okul, sinavlar


def _dogrula(db, did):
    from koordmes.services import sinav

    seans = _seans_bilgisi(db, did)
    d = sinav.donem(db, did)
    kurallar = [(r["ders1"], r["ders2"]) for r in sinav.kurallar(db, did)]
    rows = db.q("SELECT so.ogrenci_id, s.id, s.ders, s.sekil, s.seans_id FROM sinav_ogrencileri so "
                "JOIN sinavlar s ON s.id = so.sinav_id WHERE s.donem_id = ? AND s.seans_id IS NOT NULL", (did,))
    kapali = {(r["ogrenci_id"], r["seans_id"]) for r in db.q("SELECT * FROM sinav_ogrenci_kapali")}
    ogr = defaultdict(list)
    for r in rows:
        ogr[r["ogrenci_id"]].append(r)
        assert (r["ogrenci_id"], r["seans_id"]) not in kapali
        assert seans[r["seans_id"]]["acik"]
    for liste in ogr.values():
        assert len({r["seans_id"] for r in liste}) == len(liste)
        gun = defaultdict(list)
        for r in liste:
            gun[seans[r["seans_id"]]["tarih"]].append(r)
        for g in gun.values():
            assert len(g) <= d["ogrenci_gunluk_azami"]
            if len(g) > 1:
                assert not any("Uygulama" in r["sekil"] for r in g)
                dersler = {r["ders"] for r in g}
                assert not any(a in dersler and b in dersler for a, b in kurallar)
    assert not [b for b in sinav.denetle(db, did) if b.seviye == sinav.KRITIK]


def test_planlama_kisitlari_saglar(db):
    from koordmes.services import sinav

    did, okul, sinavlar = _senaryo(db)
    cakili_seans = db.q1("SELECT seans_id FROM sinavlar WHERE id = ?", (sinavlar[("Tarih", 9)],))["seans_id"]
    s = sinav.planla(db, did, sure_limit=5, seed=3)
    assert s.yerlesen == s.toplam == len(sinavlar) and not s.yerlesemeyen and s.cakili == 1
    assert db.q1("SELECT seans_id, cakili FROM sinavlar WHERE id = ?", (sinavlar[("Tarih", 9)],))["seans_id"] == cakili_seans
    _dogrula(db, did)
    ilk = [(r["id"], r["seans_id"]) for r in db.q("SELECT id, seans_id FROM sinavlar ORDER BY id")]
    sinav.planla(db, did, sure_limit=5, seed=3)
    assert [(r["id"], r["seans_id"]) for r in db.q("SELECT id, seans_id FROM sinavlar ORDER BY id")] == ilk


def test_yerlesemeyen_sinav_nedeniyle_raporlanir(db):
    from koordmes.services import sinav

    okul = _okul(db, 4)
    did = _donem(db, gun=1, saatler=sinav.VARSAYILAN_SAATLER[:2])
    ogr = okul["ogrenci"]["9-A"]
    for ders in ("Matematik", "Fizik", "Kimya"):
        sinav.sinav_ogrenci_ekle(db, sinav.sinav_ekle(db, did, ders, 9), ogr)
    s = sinav.planla(db, did, sure_limit=2, seed=1)
    assert s.yerlesen == 2 and len(s.yerlesemeyen) == 1
    assert "günlük sınav sınırı" in s.yerlesemeyen[0][2] or "aynı seansta" in s.yerlesemeyen[0][2]
    assert any("yerleştirilmemiş" in b.mesaj for b in sinav.denetle(db, did))


def test_12_siniflar_one_alinir(db):
    from koordmes.services import sinav

    okul = _okul(db, 8)
    did = _donem(db, gun=5)
    on2, dokuz = okul["ogrenci"]["12-A"], okul["ogrenci"]["9-A"]
    son12, son9 = [], []
    for i, ders in enumerate(["Matematik", "Fizik", "Kimya", "Tarih"]):
        a = sinav.sinav_ekle(db, did, ders, 10)
        sinav.sinav_ogrenci_ekle(db, a, on2)
        b = sinav.sinav_ekle(db, did, ders, 9)
        sinav.sinav_ogrenci_ekle(db, b, dokuz)
        son12.append(a)
        son9.append(b)
    sinav.planla(db, did, sure_limit=5, seed=7)
    sira = {r["id"]: i for i, r in enumerate(sinav.seanslar(db, did))}
    yer = {r["id"]: sira[r["seans_id"]] for r in db.q("SELECT id, seans_id FROM sinavlar")}
    assert sum(yer[x] for x in son12) / 4 < sum(yer[x] for x in son9) / 4
    _dogrula(db, did)


def test_buyuk_olcekli_planlama_hizli(db):
    from koordmes.services import sinav

    rng = random.Random(11)
    did = sinav.donem_ekle(db, "Büyük")
    sinav.seanslari_olustur(db, did, sinav.tarihleri_uret(PAZARTESI, PAZARTESI + timedelta(days=13)),
                            sinav.VARSAYILAN_SAATLER)
    c = db.conn
    with c:
        sinif = c.execute("INSERT INTO siniflar (ad, seviye) VALUES ('12-A', 12)").lastrowid
        diger = c.execute("INSERT INTO siniflar (ad, seviye) VALUES ('10-A', 10)").lastrowid
        c.executemany("INSERT INTO ogrenciler (ad, soyad, okul_no, sinif_id) VALUES (?, ?, ?, ?)",
                      [(f"Ö{i}", "S", str(i), sinif if i % 5 == 0 else diger) for i in range(1500)])
        c.executemany("INSERT INTO ogretmenler (ad, soyad, brans) VALUES (?, ?, ?)",
                      [(f"T{i}", "Ö", f"Ders {i % 30}") for i in range(45)])
        c.executemany("INSERT INTO sinavlar (donem_id, ders, seviye) VALUES (?, ?, ?)",
                      [(did, f"Ders {i % 100}", 9 + i // 100) for i in range(300)])
    ogrenciler = [r["id"] for r in db.q("SELECT id FROM ogrenciler")]
    sinavlar = [r["id"] for r in db.q("SELECT id FROM sinavlar")]
    kayit = {(s, o) for o in ogrenciler for s in rng.sample(sinavlar, rng.choice([1, 2, 2, 3, 4]))}
    with c:
        c.executemany("INSERT INTO sinav_ogrencileri (sinav_id, ogrenci_id) VALUES (?, ?)", sorted(kayit))
    bas = time.monotonic()
    s = sinav.planla(db, did, sure_limit=10, seed=1)
    assert time.monotonic() - bas < 25
    assert s.yerlesen == 300, s.yerlesemeyen[:5]
    _dogrula(db, did)


# ------------------------------------------------------------------ görevliler, salonlar, denetim, çıktılar

def test_gorevli_atamasi(db):
    from koordmes.services import sinav

    did, okul, sinavlar = _senaryo(db)
    sinav.planla(db, did, sure_limit=5, seed=3)
    g = sinav.gorevlileri_ata(db, did, seed=2)
    assert not g.eksik and g.sinav == len(sinavlar)
    t = okul["ogretmen"]
    assert db.q1("SELECT COUNT(*) n FROM sinav_gorevlileri WHERE ogretmen_id = ?", (t["Selma Uçar"],))["n"] == 0
    rows = db.q("SELECT g.*, s.seans_id, s.ders, s.seviye FROM sinav_gorevlileri g JOIN sinavlar s ON s.id = g.sinav_id")
    assert len({(r["ogretmen_id"], r["seans_id"]) for r in rows}) == len(rows)
    kom = defaultdict(list)
    for r in rows:
        if r["gorev"] != sinav.GOZCU:
            kom[r["sinav_id"]].append(r)
    for sid, liste in kom.items():
        assert len(liste) == 2 and sum(r["gorev"] == sinav.KOMISYON_BASKANI for r in liste) == 1
    mat9 = {r["ogretmen_id"] for r in kom[sinavlar[("Matematik", 9)]]}
    assert mat9 == {t["Cem Aksoy"], t["Deniz Tunç"]}
    baskan = next(r["ogretmen_id"] for r in kom[sinavlar[("Matematik", 9)]] if r["gorev"] == sinav.KOMISYON_BASKANI)
    assert baskan == t["Cem Aksoy"]  # dersi okutan öğretmen önce gelir
    web = {r["ogretmen_id"] for r in kom[sinavlar[("Web Tasarımı", 10)]]}
    assert web == {t["Serdar Altun"], t["Hande Işık"]}
    gun = Counter((r["ogretmen_id"], r["seans_id"]) for r in rows)
    assert max(gun.values()) == 1
    assert not [b for b in sinav.denetle(db, did) if b.seviye == sinav.KRITIK]

    kilitli = rows[0]
    db.execute("UPDATE sinav_gorevlileri SET kilitli = 1 WHERE id = ?", (kilitli["id"],))
    g2 = sinav.gorevlileri_ata(db, did, seed=9)
    assert g2.kilitli == 1 and db.q1("SELECT kilitli FROM sinav_gorevlileri WHERE id = ?", (kilitli["id"],))


def test_gorevler_esit_dagitilir(db):
    from koordmes.services import sinav

    did = _donem(db, gun=5)
    for i in range(8):
        db.insert("ogretmenler", {"ad": f"Öğretmen{i}", "soyad": "X", "brans": "Görsel Sanatlar"})
    sinif = db.insert("siniflar", {"ad": "9-A", "seviye": 9})
    ogr = [db.insert("ogrenciler", {"ad": f"Ö{i}", "soyad": "Y", "sinif_id": sinif}) for i in range(30)]
    for i in range(12):
        sinav.sinav_ogrenci_ekle(db, sinav.sinav_ekle(db, did, f"Seçmeli Ders {i}", 9), ogr[i * 2:i * 2 + 10])
    sinav.planla(db, did, sure_limit=3, seed=1)
    g = sinav.gorevlileri_ata(db, did, seed=4)
    assert not g.eksik and len(g.uyarilar) == 12  # alan dışı komisyon uyarısı
    assert max(g.yuk.values()) - min(g.yuk.values()) <= 1


def test_salonlara_yerlestirme_ve_kapasite(db):
    from koordmes.services import sinav

    okul = _okul(db, 8)
    did = _donem(db, gun=1)
    s1, s2 = db.insert("sinav_salonlari", {"ad": "Salon 1", "kapasite": 10, "sira": 1}), \
        db.insert("sinav_salonlari", {"ad": "Salon 2", "kapasite": 10, "sira": 2})
    db.insert("sinav_salonlari", {"ad": "Pasif", "kapasite": 50, "sira": 0, "aktif": 0})
    seans = sinav.seanslar(db, did)[0]["id"]
    a = sinav.sinav_ekle(db, did, "Matematik", 9)
    b = sinav.sinav_ekle(db, did, "Fizik", 10)
    sinav.sinav_ogrenci_ekle(db, a, okul["ogrenci"]["9-A"] + okul["ogrenci"]["12-A"][:7])
    sinav.sinav_ogrenci_ekle(db, b, okul["ogrenci"]["10-A"])
    sinav.seans_ata(db, a, seans)
    sinav.seans_ata(db, b, seans)
    m = sinav.salonlara_yerlestir(db, did)
    assert "3 öğrenci salonlara sığmadı" in m
    rows = db.q("SELECT so.*, s.id sinav FROM sinav_ogrencileri so JOIN sinavlar s ON s.id = so.sinav_id")
    dolu = Counter(r["salon_id"] for r in rows)
    assert dolu[s1] == 10 and dolu[s2] == 10 and dolu[None] == 3
    for salon in (s1, s2):
        koltuk = sorted((r["sira_no"], r["sinav"]) for r in rows if r["salon_id"] == salon)
        assert [k for k, _ in koltuk] == list(range(1, 11))
        sinavlar = [x for _, x in koltuk]
        if len(set(sinavlar)) > 1:
            azinlik = min(Counter(sinavlar).values())
            assert sum(x != y for x, y in zip(sinavlar, sinavlar[1:])) >= 2 * azinlik - 1
    assert any("kapasitesi" in b.mesaj and b.seviye == sinav.KRITIK for b in sinav.denetle(db, did))

    db.update("sinav_salonlari", s2, {"kapasite": 20})
    sinav.salonlara_yerlestir(db, did)
    assert db.q1("SELECT COUNT(*) n FROM sinav_ogrencileri WHERE salon_id IS NULL")["n"] == 0
    g = sinav.gorevlileri_ata(db, did, seed=1)
    gozcu = db.q("SELECT * FROM sinav_gorevlileri WHERE gorev = ?", (sinav.GOZCU,))
    assert len(gozcu) == 3 and all(r["salon_id"] for r in gozcu) and not g.eksik


def test_denetim_elle_yapilan_catismayi_bulur(db):
    from koordmes.services import sinav

    okul = _okul(db, 4)
    did = _donem(db, gun=2)
    ogr = okul["ogrenci"]["9-A"]
    a = sinav.sinav_ekle(db, did, "Matematik", 9)
    b = sinav.sinav_ekle(db, did, "Fizik", 9, sekil="Uygulama")
    sinav.sinav_ogrenci_ekle(db, a, ogr)
    sinav.sinav_ogrenci_ekle(db, b, ogr[:1])
    seans = sinav.seanslar(db, did)
    sinav.seans_ata(db, a, seans[0]["id"])
    sinav.seans_ata(db, b, seans[0]["id"])
    bulgular = sinav.denetle(db, did)
    assert any(x.seviye == sinav.KRITIK and "aynı seansta" in x.mesaj for x in bulgular)
    assert any(x.seviye == sinav.KRITIK and "uygulamalı" in x.mesaj for x in bulgular)
    sinav.seans_ata(db, b, seans[1]["id"])
    assert any("uygulamalı" in x.mesaj for x in sinav.denetle(db, did))
    sinav.seans_ata(db, b, seans[-1]["id"])
    assert not [x for x in sinav.denetle(db, did) if x.seviye == sinav.KRITIK]


def test_butun_ciktilar_uretilir(db, tmp_path):
    from docx import Document
    from openpyxl import load_workbook

    from koordmes.services import sinav, sinav_belge

    db.execute("UPDATE kurum SET ad = 'Atatürk Mesleki ve Teknik Anadolu Lisesi', il = 'Ankara', "
               "ilce = 'Yenimahalle', mudur = 'Hasan Tekin' WHERE id = 1")
    did, okul, sinavlar = _senaryo(db)
    for i in range(3):
        db.insert("sinav_salonlari", {"ad": f"Salon {i + 1}", "kapasite": 12, "sira": i})
    sinav.plani_temizle(db, did, cakili_dahil=True)
    with pytest.raises(ValueError):
        sinav_belge.sinav_takvimi_excel(db, did, tmp_path)
    assert sinav.planla(db, did, sure_limit=5, seed=2).yerlesen == len(sinavlar)
    sinav.salonlara_yerlestir(db, did)
    sinav.gorevlileri_ata(db, did, seed=2)
    yollar, hatalar = sinav_belge.hepsini_uret(db, did, tmp_path / "cikti")
    assert not hatalar and len(yollar) == len(sinav_belge.CIKTILAR)
    for y in yollar:
        assert y.exists() and y.stat().st_size > 3000
    takvim = Document(str(tmp_path / "cikti" / "Sınav Takvimi.docx"))
    metin = "\n".join(c.text for t in takvim.tables for r in t.rows for c in r.cells)
    assert "Matematik" in metin and "Web Tasarımı" in metin
    assert any("YENİMAHALLE KAYMAKAMLIĞI" in p.text for p in takvim.paragraphs)
    wb = load_workbook(tmp_path / "cikti" / "Sınav Takvimi.xlsx")
    assert wb.sheetnames[:2] == ["Sınav Takvimi", "Öğrenci Sınavları"]
    assert wb["Sınav Takvimi"].max_row >= 3 + len(sinavlar)
    yazilar = Document(str(tmp_path / "cikti" / "Öğretmen Görevlendirme Yazıları.docx"))
    assert sum("Sayın" in p.text for p in yazilar.paragraphs) == \
        db.q1("SELECT COUNT(DISTINCT ogretmen_id) n FROM sinav_gorevlileri")["n"]
    varsayilan = sinav_belge.sinav_takvimi_excel(db, did)
    assert varsayilan.parent.parent.name == "Sınav Planlama" and varsayilan.exists()


# ------------------------------------------------------------------ seans saatleri, toplu düzenleme, aktarım türleri

def test_seans_saati_ozeti_ve_ozel_saat(db):
    from koordmes.services import sinav, sinav_belge

    okul = _okul(db, 3)
    did = _donem(db, gun=2)
    seans = sinav.seanslar(db, did)
    assert sinav.seans_saati_ayarla(db, [seans[0]["id"], seans[4]["id"]], "8.30", "9:10") == 2
    yeni = _seans_bilgisi(db, did)
    assert (yeni[seans[0]["id"]]["baslangic"], yeni[seans[4]["id"]]["bitis"]) == ("08:30", "09:10")
    assert yeni[seans[1]["id"]]["baslangic"] == "10:00"
    with pytest.raises(ValueError):
        sinav.seans_saati_ayarla(db, [seans[0]["id"]], "10:00", "09:00")
    a = sinav.sinav_ekle(db, did, "Matematik", 9)
    b = sinav.sinav_ekle(db, did, "Web Tasarımı", 10, sekil="Uygulama")
    sinav.sinav_ogrenci_ekle(db, a, okul["ogrenci"]["9-A"])
    sinav.sinav_ogrenci_ekle(db, b, okul["ogrenci"]["10-A"][:2])
    sinav.seans_ata(db, a, seans[0]["id"])
    sinav.seans_ata(db, b, seans[0]["id"])
    assert sinav.seans_ozeti(db, did) == {seans[0]["id"]: (2, 5)}
    db.update("sinav_donemleri", did, {"uygulama_saati": "13:30"})
    d = sinav.donem(db, did)
    assert sinav.ozel_saat(d, "Uygulama") == "13:30" and sinav.ozel_saat(d, "Yazılı") == ""
    saat = {r["id"]: sinav_belge._saat(r, d) for r in sinav.sinav_listesi(db, did)}
    assert saat == {a: "08:30-09:10", b: "13:30"}


def test_toplu_duzenleme_ve_grup_kodu_birlestirme(db):
    from koordmes.services import sinav

    okul = _okul(db, 3)
    did = _donem(db, gun=1)
    ogr = okul["ogrenci"]["9-A"]
    a = sinav.sinav_ekle(db, did, "Matematik", 9, grup="EML")
    b = sinav.sinav_ekle(db, did, "Matematik", 9, grup="ATL")
    c = sinav.sinav_ekle(db, did, "Matematik", 9, grup="TL")
    sinav.sinav_ogrenci_ekle(db, a, ogr[:1], tur=sinav.SORUMLULUK)
    sinav.sinav_ogrenci_ekle(db, b, ogr[1:2], tur=sinav.OYS)
    sinav.sinav_ogrenci_ekle(db, c, ogr[2:])
    sinav.seans_ata(db, a, sinav.seanslar(db, did)[0]["id"], cakili=False)
    assert sinav.sinavlari_toplu_guncelle(db, [a, b], sekil="Sözlü", sure_dk=30, komisyon_sayisi=3, gozcu_sayisi=0,
                                          ucretsiz=1) == 2
    sec = "SELECT sekil, sure_dk, komisyon_sayisi, gozcu_sayisi, cakili, ucretsiz, grup, aciklama FROM sinavlar WHERE id = ?"
    assert tuple(db.q1(sec, (b,)))[:6] == ("Sözlü", 30, 3, 0, 0, 1)
    sinav.sinavlari_toplu_guncelle(db, [a, b], cakili=1, komisyon_sayisi=None)
    assert db.q1(sec, (a,))["cakili"] == 1 and db.q1(sec, (b,))["cakili"] == 0  # seansı olmayan çakılamaz
    assert db.q1(sec, (a,))["komisyon_sayisi"] is None
    with pytest.raises(ValueError):
        sinav.sinavlari_toplu_guncelle(db, [a], ders="Fizik")
    assert sinav.sinavlari_birlestir(db, [a, b, c]) == 2
    r = db.q1(sec, (a,))
    assert r["grup"] == "EML-ATL-TL" and "Birleştirilen: Matematik (9. Sınıf) - ATL grubu" in r["aciklama"]
    tur = {x["ogrenci_id"]: x["tur"] for x in sinav.sinav_ogrencileri(db, a)}
    assert len(tur) == 3 and tur[ogr[0]] == sinav.SORUMLULUK and tur[ogr[1]] == sinav.OYS


def test_eokul_oys_grup_kodu_ve_tek_ders(db, tmp_path):
    from koordmes.services import sinav

    _okul(db, 3)
    did = sinav.donem_ekle(db, "OYS", tur="Ortalama Yükseltme")
    yol = _xlsx(tmp_path / "liste.xlsx", [
        ["Öğrenci No", "Adı Soyadı", "Sınıfı", "Sorumlu Olduğu Ders", "Dersin Sınıf Seviyesi"],
        [101, "ÖĞRENCİ101 SOYAD0", "9-A", "MATEMATİK", 9],
        [107, "ÖĞRENCİ107 SOYAD0", "12-A", "FİZİK", 10],
        [108, "ÖĞRENCİ108 SOYAD1", "12-A", "FİZİK", 10],
        [108, "ÖĞRENCİ108 SOYAD1", "12-A", "KİMYA", 11],
    ])
    s = sinav.eokul_sorumlu_aktar(db, did, yol, tur=sinav.OYS, grup_kodu="EML")
    assert s.sinav_eklenen == 3 and s.ogrenci_eklenen == 4
    grup = {(r["ders"], r["seviye"]): r["grup"] for r in db.q("SELECT * FROM sinavlar WHERE donem_id = ?", (did,))}
    assert grup == {("Matematik", 9): "EML-OYS", ("Fizik", 10): "EML-OYS", ("Kimya", 11): "EML-OYS"}
    assert {r["tur"] for r in db.q("SELECT tur FROM sinav_ogrencileri")} == {sinav.OYS}
    assert {r["tur"] for r in sinav.ogrenci_programi(db, did)} == {sinav.OYS}
    assert sinav.eokul_sorumlu_aktar(db, did, yol, tur=sinav.OYS, ayri=False).sinav_eklenen == 3  # grupsuz ortak sınav

    tek = sinav.donem_ekle(db, "Tek Ders", tur="Tek Ders")
    t = sinav.eokul_sorumlu_aktar(db, tek, yol, tur=sinav.TEK_DERS, tek_ders=True)
    assert t.elenen == 2 and t.ogrenci_eklenen == 1  # 9. sınıf öğrencisi ve iki dersi olan öğrenci elenir
    assert [r["ders"] for r in db.q("SELECT ders FROM sinavlar WHERE donem_id = ?", (tek,))] == ["Fizik"]


# ------------------------------------------------------------------ kurallar, bağlı sınav, en az gün

def test_kural_turleri_ve_gun_kurallari(db):
    from koordmes.services import sinav

    okul = _okul(db, 6)
    did = _donem(db, gun=4)
    o = okul["ogrenci"]
    tarihler = sorted({r["tarih"] for r in sinav.seanslar(db, did)})
    s = {}
    for ders, sv, kitle in [("Tarih", 9, o["9-A"]), ("Kimya", 10, o["10-A"]), ("Matematik", 10, o["10-A"][:3]),
                            ("Fizik", 10, o["10-A"][3:]), ("Programlama Temelleri", 9, o["9-A"][:2]),
                            ("Web Tasarımı", 10, o["12-A"])]:
        s[ders] = sinav.sinav_ekle(db, did, ders, sv)
        sinav.sinav_ogrenci_ekle(db, s[ders], kitle)
    sinav.kural_ekle(db, did, "Tarih", "Kimya", "kesin")
    sinav.kural_ekle(db, did, "Matematik", "Fizik", "ayni_seans")
    assert sinav.kural_ekle(db, did, "Kimya", "Tarih", "kesin") == sinav.kural_ekle(db, did, "Tarih", "Kimya", "kesin")
    with pytest.raises(ValueError):
        sinav.kural_ekle(db, did, "Tarih", "Fizik", "yok")
    sinav.gun_kurali_ekle(db, did, [tarihler[2]], ders="Programlama Temelleri")
    sinav.gun_kurali_ekle(db, did, [tarihler[3]], seviye=12, kapsam="ogrenci")
    with pytest.raises(ValueError):
        sinav.gun_kurali_ekle(db, did, [tarihler[0]])
    assert "12. sınıf öğrencileri: yalnız" in sinav.gun_kurali_etiketi(sinav.gun_kurallari(db, did)[1])

    p = sinav.planla(db, did, sure_limit=3, seed=1)
    assert p.yerlesen == p.toplam == len(s)
    seans = _seans_bilgisi(db, did)
    yer = {r["id"]: seans[r["seans_id"]] for r in db.q("SELECT id, seans_id FROM sinavlar")}
    assert yer[s["Tarih"]]["tarih"] != yer[s["Kimya"]]["tarih"]
    assert yer[s["Matematik"]]["id"] == yer[s["Fizik"]]["id"]
    assert yer[s["Programlama Temelleri"]]["tarih"] == tarihler[2]
    assert yer[s["Web Tasarımı"]]["tarih"] == tarihler[3]
    assert not [b for b in sinav.denetle(db, did) if b.seviye == sinav.KRITIK]

    ilk = [r["id"] for r in sinav.seanslar(db, did) if r["tarih"] == tarihler[0]]
    sinav.seans_ata(db, s["Fizik"], sinav.seanslar(db, did)[-1]["id"])
    sinav.seans_ata(db, s["Tarih"], ilk[0])
    sinav.seans_ata(db, s["Kimya"], ilk[1])
    sinav.seans_ata(db, s["Matematik"], ilk[2])
    sinav.seans_ata(db, s["Web Tasarımı"], ilk[3])
    mesaj = " ".join(b.mesaj for b in sinav.denetle(db, did) if b.seviye == sinav.KRITIK)
    assert "kesinlikle aynı güne gelmesin" in mesaj and "aynı seansta yapılsın" in mesaj
    assert "sınav günü kuralına aykırı" in mesaj


def test_bagli_sozlu_sinav_ertesi_gunun_ilk_seansina(db):
    from koordmes.services import sinav

    okul = _okul(db, 4)
    did = _donem(db, gun=3)
    seans = sinav.seanslar(db, did)
    sinav.seans_acik_yap(db, seans[4]["id"], False)  # 2. günün 1. oturumu kapalı
    yazili = sinav.sinav_ekle(db, did, "Matematik", 9)
    sozlu = sinav.sinav_ekle(db, did, "Matematik", 9, grup="Sözlü", sekil="Sözlü")
    sinav.sinav_guncelle(db, sozlu, bagli_sinav_id=yazili)
    sinav.sinav_ogrenci_ekle(db, yazili, okul["ogrenci"]["9-A"])
    sinav.sinav_ogrenci_ekle(db, sozlu, okul["ogrenci"]["9-A"])
    sinav.seans_ata(db, yazili, seans[1]["id"])
    p = sinav.planla(db, did, sure_limit=2, seed=1)
    assert p.yerlesen == 2
    assert db.q1("SELECT seans_id FROM sinavlar WHERE id = ?", (sozlu,))["seans_id"] == seans[5]["id"]
    assert not [b for b in sinav.denetle(db, did) if b.seviye == sinav.KRITIK]
    sinav.seans_ata(db, sozlu, seans[6]["id"])
    assert any("bağlı olduğu" in b.mesaj for b in sinav.denetle(db, did) if b.seviye == sinav.KRITIK)


def test_yerlesmislari_koruma_ve_en_az_gun(db):
    from koordmes.services import sinav

    okul = _okul(db, 4)
    did = _donem(db, gun=5)
    ogr = okul["ogrenci"]["9-A"]
    ids = []
    for ders in ("Matematik", "Fizik", "Kimya", "Tarih"):
        ids.append(sinav.sinav_ekle(db, did, ders, 9))
        sinav.sinav_ogrenci_ekle(db, ids[-1], ogr)
    sinav.planla(db, did, sure_limit=2, seed=4)
    once = {r["id"]: r["seans_id"] for r in db.q("SELECT id, seans_id FROM sinavlar")}
    yeni = sinav.sinav_ekle(db, did, "Coğrafya", 9)
    sinav.sinav_ogrenci_ekle(db, yeni, ogr)
    p = sinav.planla(db, did, sure_limit=2, seed=99, koru=True)
    sonra = {r["id"]: r["seans_id"] for r in db.q("SELECT id, seans_id FROM sinavlar")}
    assert all(sonra[i] == once[i] for i in ids) and sonra[yeni] and p.korunan == 4

    e = sinav.en_az_gun(db, did, sure_limit=4, seed=1)
    assert e.gun == 3 and e.toplam_gun == 5 and "en az 3 sınav gününde" in e.ozet()  # günde en çok 2 sınav
    assert {r["id"]: r["seans_id"] for r in db.q("SELECT id, seans_id FROM sinavlar")} == sonra


# ------------------------------------------------------------------ öğretmenler, elle görevlendirme, ücret

def test_nobet_ve_musaitlik_kapali_seanslari_doldurur(db):
    from koordmes.services import sinav

    okul = _okul(db, 2)
    t = okul["ogretmen"]
    db.update("ogretmenler", t["Cem Aksoy"], {"nobet_gun": 1})
    did = _donem(db, gun=3)
    pazartesi = {r["id"] for r in sinav.seanslar(db, did) if r["tarih"] == PAZARTESI.isoformat()}
    sali = {r["id"] for r in sinav.seanslar(db, did) if r["tarih"] == (PAZARTESI + timedelta(days=1)).isoformat()}
    sinav.ogretmenleri_hazirla(db, did)
    kapali = {r["id"]: sinav.kume_coz(r["kapali_seanslar"]) for r in sinav.ogretmen_listesi(db, did)}
    assert kapali[t["Cem Aksoy"]] == pazartesi and not kapali[t["Deniz Tunç"]]
    db.update("ogretmenler", t["Deniz Tunç"], {"musait_olmayan": "2-*,3-1"})
    assert sinav.nobetten_doldur(db, did) == 1 and sinav.nobetten_doldur(db, did) == 0
    kapali = {r["id"]: sinav.kume_coz(r["kapali_seanslar"]) for r in sinav.ogretmen_listesi(db, did)}
    assert kapali[t["Deniz Tunç"]] == sali


def test_elle_gorevlendirme_ve_komisyon_sayisi(db):
    from koordmes.services import sinav

    okul = _okul(db, 4)
    t = okul["ogretmen"]
    did = _donem(db, gun=2)
    seans = sinav.seanslar(db, did)
    a = sinav.sinav_ekle(db, did, "Matematik", 9)
    b = sinav.sinav_ekle(db, did, "Fizik", 10)
    sinav.sinav_ogrenci_ekle(db, a, okul["ogrenci"]["9-A"])
    sinav.sinav_ogrenci_ekle(db, b, okul["ogrenci"]["10-A"])
    sinav.seans_ata(db, a, seans[0]["id"])
    sinav.seans_ata(db, b, seans[0]["id"])
    sinav.sinav_guncelle(db, a, komisyon_sayisi=3, gozcu_sayisi=0)
    g = sinav.gorevlileri_ata(db, did, seed=1)
    say = Counter((r["sinav_id"], r["gorev"] == sinav.GOZCU) for r in db.q("SELECT * FROM sinav_gorevlileri"))
    assert say[(a, False)] == 3 and say[(a, True)] == 0 and say[(b, False)] == 2 and not g.eksik

    adaylar = sinav.gorevli_adaylari(db, a)
    durum = {x["id"]: x["durum"] for x in adaylar}
    assert durum[t["Selma Uçar"]] == "Görev almaz" and adaylar[0]["uygun"]
    mesgul = next(r["ogretmen_id"] for r in db.q("SELECT * FROM sinav_gorevlileri WHERE sinav_id = ?", (b,)))
    assert durum[mesgul] == "Aynı seansta görevli"
    uyari = sinav.gorevli_ekle(db, a, mesgul, sinav.KOMISYON_BASKANI)
    assert uyari and any("seans" in u for u in uyari)
    baskan = db.q("SELECT * FROM sinav_gorevlileri WHERE sinav_id = ? AND gorev = ?", (a, sinav.KOMISYON_BASKANI))
    assert [r["ogretmen_id"] for r in baskan] == [mesgul] and baskan[0]["kilitli"] == 1
    sinav.gorevli_sil(db, [baskan[0]["id"]])

    ga = db.q("SELECT * FROM sinav_gorevlileri WHERE sinav_id = ? ORDER BY id", (a,))
    gb = db.q("SELECT * FROM sinav_gorevlileri WHERE sinav_id = ? ORDER BY id", (b,))
    sinav.gorevli_degistir(db, ga[0]["id"], diger_id=gb[0]["id"])
    assert db.q1("SELECT ogretmen_id FROM sinav_gorevlileri WHERE id = ?", (ga[0]["id"],))["ogretmen_id"] == \
        gb[0]["ogretmen_id"]
    serbest = next(x["id"] for x in sinav.gorevli_adaylari(db, a) if x["uygun"])
    assert not [u for u in sinav.gorevli_degistir(db, ga[1]["id"], ogretmen_id=serbest) if "seans" in u]
    sinav.gorevli_guncelle(db, [ga[1]["id"]], durum=sinav.GELMEDI, kilitli=0)
    r = db.q1("SELECT * FROM sinav_gorevlileri WHERE id = ?", (ga[1]["id"],))
    assert (r["ogretmen_id"], r["durum"], r["kilitli"]) == (serbest, sinav.GELMEDI, 0)
    with pytest.raises(ValueError):
        sinav.gorevli_guncelle(db, [r["id"]], durum="Belki")
    sinav.gorevlileri_ata(db, did, seed=2)
    assert db.q1("SELECT kilitli FROM sinav_gorevlileri WHERE id = ?", (ga[0]["id"],))["kilitli"] == 1


def test_gorev_ucretleri(db):
    from koordmes.services import sinav

    did = _donem(db, gun=2)
    t1 = db.insert("ogretmenler", {"ad": "Ali", "soyad": "Kaya", "brans": "Matematik"})
    t2 = db.insert("ogretmenler", {"ad": "Ayşe", "soyad": "Kara", "brans": "Matematik"})
    db.update("sinav_donemleri", did, {"komisyon_ucret_saat": 2, "gozcu_ucret_saat": 1, "yillik_azami_gorev": 3})
    sinav.ogretmenleri_hazirla(db, did)
    sinav.ogretmen_ayarla(db, did, t1, onceki_gorev=1)
    seans = sinav.seanslar(db, did)
    ids = []
    for i in range(4):
        ids.append(sinav.sinav_ekle(db, did, "Matematik", 9 + i))
        sinav.seans_ata(db, ids[-1], seans[i]["id"])
        sinav.gorevli_ekle(db, ids[-1], t1, sinav.KOMISYON_BASKANI)
    sinav.gorevli_ekle(db, ids[0], t2, sinav.GOZCU)
    sinav.sinav_guncelle(db, ids[1], ucretsiz=1)
    g = db.q1("SELECT id FROM sinav_gorevlileri WHERE sinav_id = ? AND ogretmen_id = ?", (ids[2], t1))["id"]
    sinav.gorevli_guncelle(db, [g], durum=sinav.GELMEDI)
    u = {k["ogretmen_id"]: k for k in sinav.gorev_ucretleri(db, did)}
    k = u[t1]
    assert (k["komisyon"], k["ucretsiz"], k["gelmedi"], k["ucretli_komisyon"], k["sinir_disi"], k["saat"]) == \
        (4, 1, 1, 2, 0, 4.0)
    assert [x["ucret_durumu"] for x in k["gorevler"]] == ["Ücretli", "Ücretsiz", "Gelmedi", "Ücretli"]
    sinav.ogretmen_ayarla(db, did, t1, onceki_gorev=2)
    k = {k["ogretmen_id"]: k for k in sinav.gorev_ucretleri(db, did)}[t1]
    assert (k["ucretli_komisyon"], k["sinir_disi"], k["saat"]) == (1, 1, 2.0)
    assert (u[t2]["ucretli_gozcu"], u[t2]["saat"]) == (1, 1.0)


# ------------------------------------------------------------------ salon ayırma, elle salon, oturma düzeni, yoklama

def test_sinava_salon_ayirma_elle_salon_ve_oturma_duzeni(db):
    from koordmes.services import sinav

    okul = _okul(db, 8)
    did = _donem(db, gun=1)
    s1 = db.insert("sinav_salonlari", {"ad": "Salon 1", "kapasite": 10, "sira": 1})
    s2 = db.insert("sinav_salonlari", {"ad": "Salon 2", "kapasite": 10, "sira": 2})
    atl = db.insert("sinav_salonlari", {"ad": "Bilişim Atölyesi", "kapasite": 6, "sira": 3})
    seans = sinav.seanslar(db, did)[0]["id"]
    a = sinav.sinav_ekle(db, did, "Matematik", 9)
    b = sinav.sinav_ekle(db, did, "Web Tasarımı", 10, sekil="Uygulama")
    sinav.sinav_ogrenci_ekle(db, a, okul["ogrenci"]["9-A"])
    sinav.sinav_ogrenci_ekle(db, b, okul["ogrenci"]["10-A"])
    sinav.seans_ata(db, a, seans)
    sinav.seans_ata(db, b, seans)
    sinav.sinav_salonlari_ayarla(db, b, [atl])
    assert sinav.sinav_salonlari(db, b) == [atl]
    kilit = okul["ogrenci"]["9-A"][0]
    sinav.ogrenci_salon_ata(db, a, [kilit], s2)
    db.update("sinav_donemleri", did, {"oturma_duzeni": "Okul numarasına göre"})
    m = sinav.salonlara_yerlestir(db, did)
    assert "2 öğrenci salonlara sığmadı" in m  # uygulama sınavına ayrılan 6 kişilik atölyeden taşanlar
    yer = {(r["sinav_id"], r["ogrenci_id"]): r for r in db.q("SELECT * FROM sinav_ogrencileri")}
    assert (yer[(a, kilit)]["salon_id"], yer[(a, kilit)]["salon_kilitli"]) == (s2, 1)
    assert Counter(yer[(b, o)]["salon_id"] for o in okul["ogrenci"]["10-A"]) == {atl: 6, None: 2}
    salon1 = sorted((r["sira_no"], r["ogrenci_id"]) for r in yer.values() if r["salon_id"] == s1)
    assert [o for _, o in salon1] == okul["ogrenci"]["9-A"][1:]  # okul numarası sırası
    assert all(r["salon_id"] != atl for (sid, _), r in yer.items() if sid == a)

    db.update("sinav_salonlari", atl, {"kapasite": 8})
    sinav.salonlara_yerlestir(db, did, duzen="Karışık")
    assert db.q1("SELECT salon_id FROM sinav_ogrencileri WHERE sinav_id = ? AND ogrenci_id = ?", (a, kilit))["salon_id"] == s2
    sinav.ogrenci_salon_ata(db, a, [kilit], None)
    r = db.q1("SELECT * FROM sinav_ogrencileri WHERE sinav_id = ? AND ogrenci_id = ?", (a, kilit))
    assert (r["salon_id"], r["salon_kilitli"]) == (None, 0)


def test_yoklama_nobetci_ve_yeni_belgeler(db, tmp_path):
    from docx import Document
    from openpyxl import load_workbook

    from koordmes.services import sinav, sinav_belge

    db.execute("UPDATE kurum SET ad = 'Atatürk MTAL', mudur = 'Hasan Tekin' WHERE id = 1")
    did, okul, sinavlar = _senaryo(db)
    for i in range(3):
        db.insert("sinav_salonlari", {"ad": f"Salon {i + 1}", "kapasite": 12, "sira": i})
    sinav.planla(db, did, sure_limit=5, seed=2)
    sinav.salonlara_yerlestir(db, did)
    sinav.gorevlileri_ata(db, did, seed=2)
    db.update("sinav_donemleri", did, {"baslik_aciklama": "2026-2027 Kasım Dönemi Sorumluluk Sınavları",
                                       "duzenleyen": "Selma Uçar", "duzenleyen_unvan": "Müdür Yardımcısı"})
    t = okul["ogretmen"]
    ilk_gun = min(r["tarih"] for r in sinav.sinav_listesi(db, did) if r["tarih"])
    sinav.nobetci_ayarla(db, did, ilk_gun, t["Selma Uçar"])
    assert sinav.nobetciler(db, did)[ilk_gun]["ad"] == "Selma Uçar"
    mat = sinavlar[("Matematik", 9)]
    ogr = [r["ogrenci_id"] for r in sinav.sinav_ogrencileri(db, mat)]
    sinav.ogrenci_katilim_ayarla(db, mat, ogr, sinav.GIRDI)
    sinav.ogrenci_katilim_ayarla(db, mat, ogr[:2], sinav.GIRMEDI)
    with pytest.raises(ValueError):
        sinav.ogrenci_katilim_ayarla(db, mat, ogr, "Belki")
    r = next(x for x in sinav.sinav_listesi(db, did) if x["id"] == mat)
    assert (r["girdi"], r["girmedi"]) == (len(ogr) - 2, 2)

    hedef = tmp_path / "cikti"
    yollar, hatalar = sinav_belge.hepsini_uret(db, did, hedef)
    assert not hatalar and len(yollar) == len(sinav_belge.CIKTILAR)

    def metin(ad):
        d = Document(str(hedef / ad))
        return "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])

    girmeyen = metin("Sınava Girmeyen Öğrenci Tutanakları.docx")
    adlar = {f"{r['ad']} {r['soyad']}" for r in sinav.sinav_ogrencileri(db, mat) if r["ogrenci_id"] in ogr[:2]}
    assert all(a in girmeyen for a in adlar) and "yazılı 2 öğrenci sınava girmemiştir" in girmeyen
    assert f"kayıtlı {len(ogr)} öğrenciden {len(ogr) - 2} öğrenci sınava girmiş, 2" in metin("Sınav Tutanakları.docx")
    takvim = metin("Sınav Takvimi.docx")
    assert "2026-2027 KASIM" in takvim and "Nöbetçi Müdür Yardımcısı: Selma Uçar" in takvim
    assert "Selma Uçar" in takvim and "Hasan Tekin" in takvim
    onay = metin("Sınav Onayı.docx")
    assert "MÜDÜRLÜK MAKAMINA" in onay and "OLUR" in onay and "Hasan Tekin" in onay
    assert "Selma Uçar" in metin("Öğretmen İmza Cetveli.docx")
    assert "TESLİM TUTANAĞI" in metin("Sınav Kâğıdı Teslim Tutanakları.docx")
    assert "CEVAP ANAHTARI" in metin("Soru ve Cevap Anahtarı Teslim Tutanakları.docx")
    wb = load_workbook(hedef / "Sınav Takvimi.xlsx")
    assert wb.sheetnames[:3] == ["Sınav Takvimi", "Öğrenci Sınavları", "Seans Özeti"]
    assert {"Salon - Salon 1", "Salon - Salon 2"} <= set(wb.sheetnames)
    ozet = wb["Seans Özeti"]
    assert sum(ozet.cell(row=i, column=5).value or 0 for i in range(4, ozet.max_row + 1)) == \
        db.q1("SELECT COUNT(*) n FROM sinav_ogrencileri")["n"]
    ucret = load_workbook(hedef / "Sınav Görev Ücret Çizelgesi.xlsx")
    assert ucret.sheetnames == ["Ücret Özeti", "Görev Ayrıntısı"]
    assert "TOPLAM" in [c.value for c in ucret["Ücret Özeti"]["A"]]


# ------------------------------------------------------------------ arayüz

@pytest.fixture(scope="module")
def app():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


@pytest.fixture
def mesajlar(monkeypatch):
    from koordmes.ui import ortak

    kayit = []
    monkeypatch.setattr(ortak, "bilgi", lambda _p, m, *a, **k: kayit.append(("bilgi", m)))
    monkeypatch.setattr(ortak, "hata", lambda _p, m, *a, **k: kayit.append(("hata", m)))
    monkeypatch.setattr(ortak, "onay", lambda *a, **k: True)
    monkeypatch.setattr(ortak, "kaydedildi", lambda *a, **k: kayit.append(("kaydedildi", a[1])))
    return kayit


def _sekmeleri_gez(app, sayfa) -> None:
    for i in range(sayfa.count()):
        sayfa.setCurrentIndex(i)
        app.processEvents()
        sayfa.currentWidget().yenile()
    sayfa.yenile()


def test_arayuz_bos_veritabaninda_acilir(app, db, mesajlar):
    from koordmes.ui.sayfa_sinav import SinavSayfasi

    w = SinavSayfasi(db)
    w.show()
    assert w.donem_id is None
    assert [w.tabText(i) for i in range(w.count())] == ["İş Adımları", "Seans Çizelgesi", "Sınavlar ve Öğrenciler",
                                                        "Öğretmenler", "Salonlar", "Planlama", "Çıktılar"]
    _sekmeleri_gez(app, w)
    w.seans.olustur()
    w.planlama.planla(True)
    w.planlama.denetle()
    w.ciktilar.hepsi()
    w.donem_sil()
    assert len(mesajlar) == 4 and all(t == "bilgi" for t, _ in mesajlar)
    w.close()


def test_arayuz_donemle_butun_sekmeler(app, ornek_db, mesajlar, tmp_path):
    from PySide6.QtCore import Qt

    from koordmes.services import sinav
    from koordmes.ui.sayfa_sinav import SinavDialog, SinavSayfasi

    db = ornek_db
    did = sinav.donem_ekle(db, "2026-2027 Kasım Sorumluluk Sınavları")
    sinav.seanslari_olustur(db, did, sinav.tarihleri_uret(PAZARTESI, PAZARTESI + timedelta(days=4)),
                            sinav.VARSAYILAN_SAATLER[:3])
    ogr = defaultdict(list)
    for r in db.q("SELECT o.id, si.seviye FROM ogrenciler o JOIN siniflar si ON si.id = o.sinif_id ORDER BY o.id"):
        ogr[r["seviye"]].append(r["id"])
    for i, ders in enumerate(["Matematik", "Fizik", "Kimya", "Tarih", "Coğrafya", "Biyoloji"]):
        for sv, liste in ogr.items():
            sinav.sinav_ogrenci_ekle(db, sinav.sinav_ekle(db, did, ders, sv), liste[i::4])
    db.insert("sinav_salonlari", {"ad": "Salon 1", "kapasite": 30, "sira": 1})
    toplam = db.q1("SELECT COUNT(*) n FROM sinavlar WHERE donem_id = ?", (did,))["n"]

    w = SinavSayfasi(db)
    w.show()
    app.processEvents()
    assert w.donem_id == did
    _sekmeleri_gez(app, w)

    w.setCurrentWidget(w.seans)
    app.processEvents()
    assert w.seans.cizelge.rowCount() == 5 and w.seans.cizelge.columnCount() == 3
    hucre = w.seans.cizelge.item(0, 0)
    hucre.setCheckState(Qt.Unchecked)
    assert db.q1("SELECT acik FROM sinav_seanslari WHERE id = ?", (hucre.data(Qt.UserRole),))["acik"] == 0
    w.seans.cizelge.setCurrentCell(0, 0)
    w.seans.duzeni_uygula()
    assert sum(1 for r in sinav.seanslar(db, did) if not r["acik"]) == 5
    w.seans.tumunu_ac()
    assert all(r["acik"] for r in sinav.seanslar(db, did))

    w.setCurrentWidget(w.sinavlar)
    app.processEvents()
    assert w.sinavlar.model.rowCount() == toplam
    w.sinavlar.view.selectRow(0)
    app.processEvents()
    assert w.sinavlar.omodel.rowCount() > 0
    kayit = w.sinavlar._secili_sinav()
    d = SinavDialog(db, did, kayit, w)
    d.sekil.setCurrentText("Uygulama")
    d.sure.setValue(60)
    d._kaydet()
    r = db.q1("SELECT sekil, sure_dk FROM sinavlar WHERE id = ?", (kayit["id"],))
    assert (r["sekil"], r["sure_dk"]) == ("Uygulama", 60)
    d = SinavDialog(db, did, None, w)
    d.ders.setText("ALMANCA")
    d.seviye.setValue(10)
    d._kaydet()
    assert db.q1("SELECT ders FROM sinavlar WHERE id = ?", (d.sinav_id,))["ders"] == "Almanca"
    d = SinavDialog(db, did, None, w)
    d.ders.setText("Almanca")
    d.seviye.setValue(10)
    d._kaydet()
    assert mesajlar[-1][0] == "hata" and "zaten var" in mesajlar[-1][1]

    w.setCurrentWidget(w.ogretmenler)
    app.processEvents()
    assert w.ogretmenler.model.rowCount() == db.q1("SELECT COUNT(*) n FROM ogretmenler WHERE aktif = 1")["n"]
    w.ogretmenler.ders1.setCurrentText("Matematik")
    w.ogretmenler.ders2.setCurrentText("Fizik")
    w.ogretmenler.kural_ekle()
    assert w.ogretmenler.kurallar.count() == 1

    w.setCurrentWidget(w.planlama)
    app.processEvents()
    w.planlama.sure.setValue(3)
    w.planlama.planla(True)
    bas = time.monotonic()
    while w.planlama._calisiyor and time.monotonic() - bas < 60:
        app.processEvents()
        time.sleep(0.02)
    assert not w.planlama._calisiyor
    assert "yerleştirildi" in w.planlama.sonuc.toPlainText()
    planli = db.q1("SELECT COUNT(*) n FROM sinavlar WHERE donem_id = ? AND seans_id IS NOT NULL", (did,))["n"]
    assert planli == toplam
    assert "Öğrencisi olmayan 1 sınav planlanmadı" in w.planlama.sonuc.toPlainText()
    assert db.q1("SELECT COUNT(*) n FROM sinav_gorevlileri g JOIN sinavlar s ON s.id = g.sinav_id "
                 "WHERE s.donem_id = ?", (did,))["n"] > 0
    w.planlama.denetle()
    assert w.planlama.sonuc.toPlainText()
    assert w.planlama.model.rowCount() == toplam + 1

    w.setCurrentWidget(w.ciktilar)
    app.processEvents()
    w.ciktilar.hepsi()
    assert mesajlar[-1][0] == "kaydedildi"
    _sekmeleri_gez(app, w)
    w.donem_sil()
    assert w.donem_id is None and not sinav.donemler(db)
    _sekmeleri_gez(app, w)
    w.close()


def _sec(sekme, kimlik) -> None:
    for i in range(sekme.proxy.rowCount()):
        if sekme.model.kimlik(sekme.proxy.mapToSource(sekme.proxy.index(i, 0)).row()) == kimlik:
            sekme.view.selectRow(i)
            return


def test_arayuz_yeni_islemler(app, db, mesajlar, monkeypatch):
    from PySide6.QtWidgets import QInputDialog

    from koordmes.services import sinav
    from koordmes.ui import sayfa_sinav
    from koordmes.ui.sayfa_sinav import (GorevlilerDialog, GunKuraliDialog, SinavDialog, SinavSayfasi,
                                         TopluDuzenleDialog)

    okul = _okul(db, 4)
    did = _donem(db, gun=3)
    seans = sinav.seanslar(db, did)
    a = sinav.sinav_ekle(db, did, "Matematik", 9)
    b = sinav.sinav_ekle(db, did, "Fizik", 10)
    sinav.sinav_ogrenci_ekle(db, a, okul["ogrenci"]["9-A"])
    sinav.sinav_ogrenci_ekle(db, b, okul["ogrenci"]["10-A"])
    sinav.seans_ata(db, a, seans[0]["id"])
    salon = db.insert("sinav_salonlari", {"ad": "Salon 1", "kapasite": 30, "sira": 1})
    monkeypatch.setattr(QInputDialog, "getItem", lambda _p, _t, _l, ogeler, *a, **k: (ogeler[min(1, len(ogeler) - 1)],
                                                                                     True))
    monkeypatch.setattr(QInputDialog, "getText", lambda *a, **k: ("08:15-08:55", True))
    monkeypatch.setattr(QInputDialog, "getInt", lambda *a, **k: (2, True))

    w = SinavSayfasi(db)
    w.show()
    w.setCurrentWidget(w.seans)
    app.processEvents()
    assert "1 sınav / 4 öğr." in w.seans.cizelge.item(0, 0).text()
    w.seans.cizelge.setCurrentCell(0, 1)
    w.seans.saat_degistir()
    assert "08:15-08:55" in w.seans.cizelge.item(0, 1).text()

    w.setCurrentWidget(w.sinavlar)
    app.processEvents()
    d = SinavDialog(db, did, next(r for r in sinav.sinav_listesi(db, did) if r["id"] == a), w)
    d.komisyon.setValue(3)
    d.gozcu.setValue(0)
    d.bagli.setCurrentIndex(d.bagli.findData(b))
    d.ucretsiz.setChecked(True)
    d._kaydet()
    r = db.q1("SELECT * FROM sinavlar WHERE id = ?", (a,))
    assert (r["komisyon_sayisi"], r["gozcu_sayisi"], r["bagli_sinav_id"], r["ucretsiz"]) == (3, 0, b, 1)

    def toplu_exec(dlg):
        dlg.isaret["sure_dk"].setChecked(True)
        dlg.sure.setValue(70)
        dlg.isaret["gozcu_sayisi"].setChecked(True)
        return True

    monkeypatch.setattr(TopluDuzenleDialog, "exec", toplu_exec)
    w.sinavlar.yenile()
    w.sinavlar.view.selectAll()
    w.sinavlar.toplu()
    assert {(r["sure_dk"], r["gozcu_sayisi"]) for r in db.q("SELECT * FROM sinavlar")} == {(70, None)}
    _sec(w.sinavlar, a)
    app.processEvents()
    w.sinavlar.oview.selectRow(0)
    w.sinavlar.salona_ata()
    assert db.q1("SELECT salon_id FROM sinav_ogrencileri WHERE sinav_id = ? AND salon_kilitli = 1", (a,))["salon_id"] == salon
    w.sinavlar.oview.selectRow(0)
    w.sinavlar.katilim(sinav.GIRMEDI)
    girmeyen = db.q("SELECT ogrenci_id FROM sinav_ogrencileri WHERE sinav_id = ? AND katilim = ?", (a, sinav.GIRMEDI))
    assert [r["ogrenci_id"] for r in girmeyen] == [w.sinavlar.omodel.kimlik(0)]

    w.setCurrentWidget(w.ogretmenler)
    app.processEvents()
    o = w.ogretmenler
    o.view.selectAll()
    o.onceki_gorev()
    assert {r["onceki_gorev"] for r in sinav.ogretmen_listesi(db, did)} == {2}
    o.nobetten_doldur()
    assert mesajlar[-1][0] == "bilgi"
    o.ders1.setCurrentText("Matematik")
    o.ders2.setCurrentText("Fizik")
    o.kural_turu.setCurrentIndex(o.kural_turu.findData("kesin"))
    o.kural_ekle()
    assert [r["tur"] for r in sinav.kurallar(db, did)] == ["kesin"]

    def gun_exec(dlg):
        dlg.gunler.item(1).setCheckState(sayfa_sinav.Qt.Checked)
        dlg.seviye.setValue(12)
        return True

    monkeypatch.setattr(GunKuraliDialog, "exec", gun_exec)
    o.gun_kurali_ekle()
    assert o.gun_kurallari.count() == 1
    o.nobetciler.setCurrentRow(0)
    o.nobetci_sec()
    assert len(sinav.nobetciler(db, did)) == 1

    sayfa_sinav._salonlari_al(w.salonlar, [])
    assert db.q1("SELECT COUNT(*) n FROM sinav_salonlari")["n"] == 4

    sinav.gorevlileri_ata(db, did, seed=1)
    g = GorevlilerDialog(db, did, a, w)
    n = g.model.rowCount()
    assert sum(1 for r in g.model.rows if r[0] != sinav.GOZCU) == 3
    g.view.selectRow(0)
    g.guncelle(durum=sinav.GELMEDI)
    assert db.q1("SELECT COUNT(*) n FROM sinav_gorevlileri WHERE durum = ?", (sinav.GELMEDI,))["n"] == 1
    g.ekle(sinav.GOZCU)
    assert g.model.rowCount() == n + 1
    g.close()

    w.setCurrentWidget(w.planlama)
    app.processEvents()
    w.planlama.en_az_gun()
    bas = time.monotonic()
    while w.planlama._calisiyor and time.monotonic() - bas < 60:
        app.processEvents()
        time.sleep(0.02)
    assert "en az" in w.planlama.sonuc.toPlainText() or "sığdırılamadı" in w.planlama.sonuc.toPlainText()
    w.close()
