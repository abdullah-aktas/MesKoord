from datetime import date

from koordmes.services import belgeler, devam, ekders, excel_io, koordinator, program, takvim


def test_program_catismasiz_ve_tam(ornek_db):
    assert program.catismalar(ornek_db) == []
    assert program.eksik_saatler(ornek_db) == []


def test_program_okul_gunlerine_uyar(ornek_db):
    rows = ornek_db.q("""SELECT ph.gun FROM program_hucreleri ph JOIN ders_atamalari da ON da.id = ph.atama_id
                         JOIN siniflar s ON s.id = da.sinif_id WHERE s.ad = '12-A'""")
    assert rows and {r["gun"] for r in rows} <= {1, 2}


def test_atolye_blok_bolunmez(ornek_db):
    rows = ornek_db.q("""SELECT ph.gun, ph.saat FROM program_hucreleri ph JOIN ders_atamalari da ON da.id = ph.atama_id
                         JOIN dersler d ON d.id = da.ders_id JOIN siniflar s ON s.id = da.sinif_id
                         WHERE s.ad = '12-A' AND d.kisa_ad = 'WEB' ORDER BY ph.gun, ph.saat""")
    by_day = {}
    for r in rows:
        by_day.setdefault(r["gun"], []).append(r["saat"])
    for saatler in by_day.values():
        assert len(saatler) == 4 and saatler == list(range(saatler[0], saatler[0] + 4))


def test_kilitli_hucre_korunur(ornek_db):
    h = ornek_db.q1("SELECT * FROM program_hucreleri LIMIT 1")
    ornek_db.execute("UPDATE program_hucreleri SET kilitli = 1 WHERE id = ?", (h["id"],))
    program.programi_olustur(ornek_db, sure_limit=10, seed=7)
    assert ornek_db.q1("SELECT * FROM program_hucreleri WHERE atama_id = ? AND gun = ? AND saat = ? AND kilitli = 1",
                       (h["atama_id"], h["gun"], h["saat"]))
    assert program.catismalar(ornek_db) == []


def test_koordinator_sinirlari(ornek_db):
    sonuc = koordinator.otomatik_gorevlendir(ornek_db)
    assert sonuc.gorevler
    azami = koordinator.haftalik_azami(ornek_db)
    assert azami == 20
    for t, yuk in sonuc.yukler.items():
        assert yuk <= azami
    gun_yuk = {}
    for g in sonuc.gorevler:
        gun_yuk[(g["ogretmen_id"], g["gun"])] = gun_yuk.get((g["ogretmen_id"], g["gun"]), 0) + g["haftalik_saat"]
    assert max(gun_yuk.values()) <= 8
    dolu = {(r["ogretmen_id"], r["gun"]) for r in ornek_db.q(
        "SELECT DISTINCT da.ogretmen_id, ph.gun FROM program_hucreleri ph JOIN ders_atamalari da ON da.id = ph.atama_id")}
    assert not any((g["ogretmen_id"], g["gun"]) in dolu for g in sonuc.gorevler)


def test_ziyaret_plani_tatil_atlar(ornek_db):
    n = koordinator.ziyaret_plani_olustur(ornek_db, 2026, 10)
    assert n > 0
    assert not ornek_db.q("SELECT 1 FROM ziyaretler WHERE tarih = '2026-10-29'")


def test_ek_ders_puantaj(ornek_db):
    p = ekders.puantaj(ornek_db, 2026, 10)
    assert p
    for x in p:
        assert x.toplam >= 0
        assert all(v >= 0 for v in x.gunluk.values())
        assert 29 not in x.gunluk or x.gunluk[29] == 0
    bt2 = next(x for x in p if x.ad == "Hande Işık")
    assert bt2.ders > 0


def test_ek_ders_haftalik_hesap(db):
    t = db.insert("ogretmenler", {"ad": "A", "soyad": "B", "unvan": "Öğretmen"})
    s = db.insert("siniflar", {"ad": "9-A"})
    d = db.insert("dersler", {"ad": "Mat"})
    a = db.insert("ders_atamalari", {"sinif_id": s, "ders_id": d, "ogretmen_id": t, "haftalik_saat": 20})
    for gun in range(1, 6):
        for saat in range(1, 5):
            db.insert("program_hucreleri", {"atama_id": a, "gun": gun, "saat": saat})
    [p] = ekders.puantaj(db, 2026, 10)
    # Haftada 20 saat - 15 aylık karşılık = 5 ek ders; karşılık pazartesiden itibaren dolar
    assert p.gunluk[1] == 1 and p.gunluk[2] == 4 and p.gunluk[5] == 0 and p.gunluk[9] == 4


def test_ek_ders_karsilik_once_derslerle_dolar(db):
    t = db.insert("ogretmenler", {"ad": "A", "soyad": "B", "unvan": "Öğretmen"})
    s = db.insert("siniflar", {"ad": "12-A"})
    d = db.insert("dersler", {"ad": "Atölye"})
    a = db.insert("ders_atamalari", {"sinif_id": s, "ders_id": d, "ogretmen_id": t, "haftalik_saat": 16})
    for gun, saatler in ((1, 8), (4, 8)):
        for saat in range(1, saatler + 1):
            db.insert("program_hucreleri", {"atama_id": a, "gun": gun, "saat": saat})
    i = db.insert("isletmeler", {"ad": "X"})
    db.insert("koordinator_gorevleri", {"ogretmen_id": t, "isletme_id": i, "gun": 3, "haftalik_saat": 5})
    [p] = ekders.puantaj(db, 2026, 10)
    # 16 ders saati karşılığı (15) doldurur; çarşamba koordinatörlüğün tamamı ek derstir
    assert p.gunluk[7] == 5 and p.gunluk[8] == 1
    assert p.isletme == 20
    assert p.brut > 0


def test_devamsizlik_ozeti(ornek_db):
    oz = devam.ozet(ornek_db)
    en = max(oz, key=lambda x: x.toplam_ozursuz)
    assert en.toplam_ozursuz == 7 and en.durum == "Veli bildirimi"


def test_uyarilar(ornek_db):
    u = takvim.uyarilar(ornek_db, date(2026, 10, 1))
    konular = {x.konu for x in u}
    assert "SGK" in konular and "Sözleşme" in konular


def test_excel_disa_ve_ice_aktar(db, tmp_path):
    yol = tmp_path / "ogr.xlsx"
    excel_io.tablo_yaz(yol, ["S.No", "Öğrenci No", "T.C. Kimlik No", "Adı Soyadı", "Sınıfı"],
                       [[1, 501, "10000000146", "Ayşe Nur Kaya", "12-A"], [2, 502, "", "Can Demir", "12-B"]],
                       baslik="12. Sınıf Listesi")
    s = excel_io.ice_aktar(db, "ogrenciler", yol)
    assert s.eklenen == 2 and not s.hatalar
    r = db.q1("SELECT o.*, s.ad AS sinif FROM ogrenciler o JOIN siniflar s ON s.id = o.sinif_id WHERE okul_no = '501'")
    assert r["ad"] == "Ayşe Nur" and r["soyad"] == "Kaya" and r["sinif"] == "12-A"
    s = excel_io.ice_aktar(db, "ogrenciler", yol)
    assert s.guncellenen == 2


def test_belge_uretimi(ornek_db, tmp_path):
    from docx import Document
    y = ornek_db.q1("SELECT id FROM yerlestirmeler WHERE isletme_id IS NOT NULL LIMIT 1")["id"]
    yol = belgeler.uret(ornek_db, "isletme_sozlesmesi", y, tmp_path)
    metin = "\n".join(p.text for p in Document(str(yol)).paragraphs)
    tablo = " ".join(c.text for t in Document(str(yol)).tables for row in t.rows for c in row.cells)
    assert "{{" not in metin + tablo
    assert "Atatürk Mesleki" in metin
    t = ornek_db.q1("SELECT ogretmen_id FROM koordinator_gorevleri LIMIT 1")["ogretmen_id"]
    assert belgeler.uret(ornek_db, "koordinator_gorevlendirme", t, tmp_path).exists()
